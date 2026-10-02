# -*- coding: utf-8 -*-
"""Controleert of opdracht, startbestand, kaarten en sleutel bij elkaar passen.

Alle posities worden afgeleid uit gegevens.py: er staan geen vaste rijnummers
in. Voeg je daar een boeking of een medewerker toe, dan schuift deze controle
vanzelf mee.
"""
import importlib.util
import os
import re
import sys
from openpyxl import load_workbook

sys.dont_write_bytecode = True   # geen __pycache__ naast het lesmateriaal

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    spec = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')
u = laad('u', 'uitleg.py')
sl = laad('sl', 'sleutel.py')

wb = load_workbook(os.path.join(HIER, 'halloween-startbestand.xlsx'))
fouten = []

FORMULE = re.compile(r"=\s*[A-Z$'(]|[A-Z][A-Z.]+\(")


def cellen_van(verwijzing):
    """'E4:E15' -> alle cellen daartussen. 'C17' -> die ene cel."""
    if ':' not in verwijzing:
        return [verwijzing]
    links, rechts = verwijzing.split(':')
    k1, r1 = re.match(r'([A-Z]+)(\d+)', links).groups()
    k2, r2 = re.match(r'([A-Z]+)(\d+)', rechts).groups()
    return [f'{chr(k)}{r}'
            for k in range(ord(k1), ord(k2) + 1)
            for r in range(int(r1), int(r2) + 1)]


def opdrachttekst(oef):
    """Alles wat de leerling te lezen krijgt bij één werkblad."""
    stukken = [oef['doel'], oef['klaar'], oef['controle']]
    stukken += [f'{waar} {wat}' for waar, wat in oef['gegeven']]
    stukken += [f'{waar} {wat} {hoe}' for waar, wat, hoe in oef['maken']]
    stukken += oef['opmaak'] + oef['hints']
    stukken += [f'{wie} {wat}' for wie, wat in oef.get('verklaringen', [])]
    return stukken


def kaarttekst(kaart):
    """Alles wat op één kaart Nieuw gereedschap staat."""
    stukken = [kaart['naam'], kaart['wat'], kaart['schrijf'], kaart['fout']]
    stukken += [f'{n} {t}' for n, t in kaart['argumenten']]
    stukken += kaart['uitleg'] + kaart['weten']
    stukken += [f'{f} {t}' for f, t in kaart.get('varianten', [])]
    for b in kaart.get('beelden', []):
        stukken.append(b['onderschrift'])
        stukken.append(b.get('formule') or '')
        stukken += [str(w) for rij in b['rijen'] for w in rij]
    return stukken


# 1. Bestaat elk genoemd werkblad?
for oef in o.OEFENINGEN:
    if oef['blad'] not in wb.sheetnames:
        fouten.append(f"werkblad {oef['blad']} bestaat niet in het startbestand")

# 2. Klopt de opdrachtomschrijving met het startbestand?
for oef in o.OEFENINGEN:
    if oef['blad'] not in wb.sheetnames:
        continue
    ws = wb[oef['blad']]

    temaken = [c for waar, *_ in oef['maken'] for c in cellen_van(waar)]
    bezet = [c for c in temaken if ws[c].value is not None]
    if bezet:
        fouten.append(f"{oef['blad']}: deze cellen moet de leerling maken, maar ze zijn al "
                      f'gevuld: {bezet[:5]}')

    gegeven = [c for waar, _ in oef['gegeven'] for c in cellen_van(waar)]
    leeg = [c for c in gegeven if ws[c].value is None]
    if leeg:
        fouten.append(f"{oef['blad']}: deze cellen zouden gegeven moeten zijn, maar zijn "
                      f'leeg: {leeg[:5]}')
    print(f"blad {oef['nr']} ({oef['blad']}): {len(gegeven)} gegeven, {len(temaken)} te maken")

# 3. De opdracht mag de formule niet weggeven. Dat is het verschil met de
#    basisbundel: hier staat alleen wat eruit moet komen.
for oef in o.OEFENINGEN:
    for stuk in opdrachttekst(oef):
        if FORMULE.search(stuk):
            fouten.append(f"blad {oef['nr']}: er staat een formule in de opdrachttekst: "
                          f'{stuk[:70]!r}')

# 4. Elke kaart die een werkblad noemt, moet bestaan.
namen = {k['naam'] for k in u.KAARTEN}
for oef in o.OEFENINGEN:
    for k in oef['kaarten']:
        if k not in namen:
            fouten.append(f"blad {oef['nr']}: kaart '{k}' bestaat niet in uitleg.py")
ongebruikt = sorted(namen - {k for oef in o.OEFENINGEN for k in oef['kaarten']})
if ongebruikt:
    print(f'\nkaarten die bij geen enkel blad horen: {ongebruikt}')

# 5. Geen enkele kaart mag gegevens uit de opdracht bevatten, en geen enkele
#    kaart mag een formule uit de sleutel letterlijk tonen.
verraad = set()
verraad |= {groep for _, groep, _, _, _ in g.B1}
verraad |= {nr for nr, _, _, _, _ in g.B1}
verraad |= {naam for naam, _ in g.B2}
verraad |= {naam for naam, _, _, _, _ in g.B3}
verraad |= {post for _, _, post, _, _ in g.B3}
verraad |= {oms for _, oms, _ in g.B1_TARIEVEN}
verraad |= set(wb.sheetnames) - {'Start'}
oplossingen = {f for blad in sl.SLEUTEL.values() for _, f, _ in blad}
for kaart in u.KAARTEN:
    tekst = ' '.join(kaarttekst(kaart))
    for geheim in verraad:
        if geheim in tekst:
            fouten.append(f"kaart '{kaart['naam']}' lekt gegevens uit de opdracht: "
                          f'{geheim!r}')
    for formule in oplossingen:
        if formule in tekst:
            fouten.append(f"kaart '{kaart['naam']}' toont letterlijk een oplossing: "
                          f'{formule}')

# 6. Staan de brongegevens waar de opdracht beweert?
controles = [
    ('1 Boekingen', f'B{g.B1_START}', g.B1[0][1]),
    ('1 Boekingen', f'C{g.B1_EIND}', g.B1[-1][2]),
    ('1 Boekingen', f'I{g.B1_TAR_START}', g.B1_TARIEVEN[0][0]),
    ('1 Boekingen', f'K{g.B1_TAR_EIND}', g.B1_TARIEVEN[-1][2]),
    ('2 Attracties', f'A{g.B2_START}', g.B2[0][0]),
    ('2 Attracties', f'F{g.B2_EIND}', g.B2[-1][1][-1]),
    ('3 Ploegen', f'A{g.B3_START}', g.B3[0][0]),
    ('3 Ploegen', f'E{g.B3_EIND}', g.B3[-1][4]),
    ('3 Ploegen', 'J3', g.B3_GRENS_TIJD),
    ('3 Ploegen', 'J4', g.B3_GRENS_UREN),
    ('4 Verdachten', 'B2', g.B4_TIJDSTIP),
    ('4 Verdachten', 'B3', g.B4_KLEUR),
    ('4 Verdachten', f'A{g.B4_START}', g.B4[0][0]),
    ('4 Verdachten', f'E{g.B4_EIND}', g.B4[-1][4]),
    ('5 Afrekening', f'B{g.B5_RIJ["bar"]}', g.B5_BAROMZET),
    ('5 Afrekening', f'B{g.B5_RIJ["uurloon"]}', g.B5_UURLOON),
]
for blad, cel, verwacht in controles:
    echt = wb[blad][cel].value
    if hasattr(echt, 'time'):
        echt = echt.time()
    if echt != verwacht:
        fouten.append(f'{blad}!{cel}: verwacht {verwacht!r}, gevonden {echt!r}')

# 7. De staffel moet oplopen en bij het laagst mogelijke aantal beginnen,
#    anders geeft benaderend zoeken een fout of een stil verkeerd antwoord.
grenzen = [vanaf for vanaf, _, _ in g.B1_TARIEVEN]
if grenzen != sorted(grenzen):
    fouten.append(f'blad 1: de staffel loopt niet op: {grenzen}')
kleinste_groep = min(personen for _, _, personen, _, _ in g.B1)
if kleinste_groep < grenzen[0]:
    fouten.append(f'blad 1: de kleinste groep telt {kleinste_groep} personen, maar de '
                  f'staffel begint pas bij {grenzen[0]} — dat geeft #N/B')
gebruikt = sorted(set(a.B1_PRIJS))
if len(gebruikt) != len(g.B1_TARIEVEN):
    fouten.append(f'blad 1: maar {len(gebruikt)} van de {len(g.B1_TARIEVEN)} schijven '
                  f'worden gebruikt — dan valt een fout in het zoeken niet op')
else:
    print(f'\nblad 1: alle {len(gebruikt)} schijven van de staffel komen voor')

# 8. INDEX met VERGELIJKEN heeft maar één antwoord als er geen gelijkspel is.
def enig(waarden, wat):
    hoogste = max(waarden)
    if waarden.count(hoogste) != 1:
        fouten.append(f'{wat}: er is een gelijkspel aan de top ({hoogste}) — dan heeft '
                      f'INDEX met VERGELIJKEN geen eenduidig antwoord')


enig(a.B2_RIJTOTAAL, 'blad 2, drukste attractie')
enig(a.B2_UURTOTAAL, 'blad 2, drukste uurblok')
enig(sorted(a.B2_RIJTOTAAL)[:-1], 'blad 2, op één na drukste attractie')
enig(a.B3_UREN, 'blad 3, langste dienst')
enig([p for _, _, p, _, _ in g.B1], 'blad 1, grootste groep')

# 9. Op blad 3 zit de les in de randgevallen: precies op de tijdsgrens, en
#    precies op het minimum aantal uren.
op_tijdsgrens = [naam for (naam, _, _, _, einde) in g.B3 if einde == g.B3_GRENS_TIJD]
op_urengrens = [naam for (naam, *_), uren in zip(g.B3, a.B3_UREN)
                if uren == g.B3_GRENS_UREN]
if not op_tijdsgrens:
    fouten.append(f'blad 3: niemand stopt precies om {g.B3_GRENS_TIJD} — dan wordt het '
                  f'verschil tussen > en >= nooit zichtbaar')
if not op_urengrens:
    fouten.append(f'blad 3: niemand werkt precies {g.B3_GRENS_UREN} uur — dan toont niets '
                  f'waarom het EN moet zijn en geen OF')
print(f'blad 3: op de tijdsgrens staat {", ".join(op_tijdsgrens)}; '
      f'op de urengrens {", ".join(op_urengrens)}')

# 10. Geen enkele dienst mag over middernacht lopen: dan wordt het verschil
#     negatief en klopt de hele kolom niet meer.
over_nacht = [naam for (naam, _, _, start, einde) in g.B3 if einde <= start]
if over_nacht:
    fouten.append(f'blad 3: deze diensten lopen over middernacht: {over_nacht}')

# 11. Elke verdachte moet in de ploegenlijst staan, anders geeft het opzoeken
#     op blad 4 een #N/B.
ploeg = {naam for naam, *_ in g.B3}
onbekend = sorted({naam for naam, *_ in g.B4} - ploeg)
if onbekend:
    fouten.append(f'blad 4: deze verdachten staan niet in de ploegenlijst: {onbekend}')

# 12. Precies één verdachte mag overblijven. Dat is de hele opdracht.
if a.B4_AANTAL_OVER != 1:
    fouten.append(f'blad 4: er blijven {a.B4_AANTAL_OVER} verdachten over in plaats van 1')
else:
    print(f'blad 4: precies één verdachte blijft over — {a.B4_DADER}')

# 13. Elke verklaring moet minstens één verdachte uitsluiten, anders staat ze
#     er voor niets.
redenen = [r for _, _, rs in a.B4_REDEN for r in rs]
for sleutelwoord, verklaring in [
    ('dienst eindigde', 'de portier (einde van de dienst)'),
    ('badgescan', 'het badgesysteem'),
    ('kostuum was', 'de bezoeker (de kleur)'),
    ('geen tas', 'de poetsvrouw (de tas)'),
    ('alibi bevestigd', 'de ploegbaas (het alibi)'),
]:
    if not any(sleutelwoord in r for r in redenen):
        fouten.append(f'blad 4: de verklaring van {verklaring} sluit niemand uit — '
                      f'dan is ze overbodig')

# 14. De kruistabel van blad 2 moet in twee richtingen hetzelfde geven.
if sum(a.B2_RIJTOTAAL) != sum(a.B2_UURTOTAAL):
    fouten.append('blad 2: rijtotalen en kolomtotalen geven niet hetzelfde')

print()
if fouten:
    print('PROBLEMEN:')
    for p in fouten:
        print('  -', p)
    raise SystemExit(1)
print('Alles sluit: opdracht, startbestand, kaarten en sleutel passen bij elkaar.')
print(f'De opdracht geeft nergens een formule weg, en geen enkele kaart lekt gegevens')
print(f'uit het startbestand. Precies één verdachte blijft over: {a.B4_DADER}.')
