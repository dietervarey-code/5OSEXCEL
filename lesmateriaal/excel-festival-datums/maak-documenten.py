# -*- coding: utf-8 -*-
"""Bouwt de opdrachtbundel en de lerarensleutel van de datumbundel.

Anders dan bij de draaitabelbundels staan de fiches hier IN de opdracht: het
is een nieuw onderwerp. Alleen de draaitabel en de draaigrafiek van blad 5
verwijzen naar het boekje dat ze al hebben.
"""
import importlib.util
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')
f = laad('f', 'fiches.py')

OPDRACHT = os.path.join(HIER, 'festival-datums-opdracht.docx')
SLEUTEL = os.path.join(HIER, 'festival-datums-lerarensleutel.docx')

BLAUW = RGBColor(0x1F, 0x38, 0x64)
ORANJE = RGBColor(0xB0, 0x4A, 0x00)
GRIJS = RGBColor(0x59, 0x59, 0x59)
GROEN = RGBColor(0x38, 0x76, 0x1D)

FORMULE = re.compile(r'(=[A-Z][A-Z.]*\((?:[^()]|\((?:[^()]|\([^()]*\))*\))*\)'
                     r'(?:&(?:"[^"]*"|[A-Z]+\d+))*'
                     r'|=[A-Z]+\d+(?:&(?:"[^"]*"|[A-Z]+\d+))+)')


def opzet(doc):
    s = doc.styles['Normal']
    s.font.name = 'Calibri'
    s.font.size = Pt(11)
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(1.5)
        sec.left_margin = sec.right_margin = Cm(1.8)


def kop(doc, tekst, niveau=1, kleur=BLAUW):
    p = doc.add_heading(tekst, level=niveau)
    for r in p.runs:
        r.font.color.rgb = kleur
    return p


def zin(p, tekst):
    """Formules in een vaste letter, zodat ze opvallen in lopende tekst."""
    for stuk in FORMULE.split(str(tekst)):
        if not stuk:
            continue
        r = p.add_run(stuk)
        if FORMULE.fullmatch(stuk):
            r.font.name = 'Consolas'
            r.font.size = Pt(10)
    return p


def vast(p, tekst):
    """Een heel stuk in de vaste letter, voor formules en voorbeelden."""
    r = p.add_run(str(tekst))
    r.font.name = 'Consolas'
    r.font.size = Pt(10)
    return r


def alinea(doc, tekst, vetgedrukt=False, schuin=False, kleur=None,
           grootte=None, na=None):
    p = doc.add_paragraph()
    if vetgedrukt or schuin or kleur or grootte:
        r = p.add_run(str(tekst))
        r.bold, r.italic = vetgedrukt, schuin
        if kleur:
            r.font.color.rgb = kleur
        if grootte:
            r.font.size = Pt(grootte)
    else:
        zin(p, tekst)
    if na is not None:
        p.paragraph_format.space_after = Pt(na)
    return p


def kader(doc, titel, regels, kleur=BLAUW):
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    c = t.rows[0].cells[0]
    r = c.paragraphs[0].add_run(titel)
    r.bold = True
    r.font.color.rgb = kleur
    for regel in regels:
        pp = c.add_paragraph()
        pp.paragraph_format.space_after = Pt(2)
        zin(pp, regel)
    doc.add_paragraph()
    return t


def tabel(doc, koppen, rijen, stijl='Light Grid Accent 1', breedtes=None,
          vaste_kolommen=()):
    t = doc.add_table(rows=1, cols=len(koppen))
    t.style = stijl
    for i, k in enumerate(koppen):
        t.rows[0].cells[i].paragraphs[0].add_run(str(k)).bold = True
    for rij in rijen:
        cellen = t.add_row().cells
        for i, w in enumerate(rij):
            if i in vaste_kolommen:
                vast(cellen[i].paragraphs[0], w)
            else:
                zin(cellen[i].paragraphs[0], w)
    if breedtes:
        for rij in t.rows:
            for i, b in enumerate(breedtes):
                rij.cells[i].width = Cm(b)
    doc.add_paragraph()
    return t


def paginaeinde(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def punten_van(oef):
    return o.punten_van(oef)


def fiche_uit(doc, nummer, kaart):
    """Eén fiche: wat het is, hoe je het schrijft, en de fout die iedereen maakt."""
    kop(doc, f"Fiche {nummer} — {kaart['naam']}", 2)
    alinea(doc, kaart['wat'], schuin=True, kleur=GRIJS, na=6)

    if kaart.get('schrijf'):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(6)
        vast(p, kaart['schrijf']).bold = True
        if kaart.get('argumenten'):
            tabel(doc, ['Stuk', 'Wat het is'],
                  [[stuk, uitleg] for stuk, uitleg in kaart['argumenten']],
                  'Light List Accent 1', [4.0, 12.0], vaste_kolommen={0})

    if kaart.get('stappen'):
        for i, stap in enumerate(kaart['stappen'], 1):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(f'{i}.  ')
            r.bold = True
            r.font.color.rgb = BLAUW
            zin(p, stap)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    for regel in kaart.get('uitleg', ()):
        alinea(doc, regel, na=4)

    if kaart.get('voorbeeld'):
        v = kaart['voorbeeld']
        breed = {2: [5.4, 10.6], 3: [3.6, 5.6, 6.8]}[len(v['koppen'])]
        tabel(doc, v['koppen'], v['rijen'], 'Light List Accent 1', breed,
              vaste_kolommen=set(range(len(v['koppen']))))

    if kaart.get('weten'):
        kader(doc, 'Goed om te weten', kaart['weten'], BLAUW)
    kader(doc, 'De fout die iedereen één keer maakt', [kaart['fout']], ORANJE)


# ==========================================================================
#  1. de opdrachtbundel
# ==========================================================================
doc = Document()
opzet(doc)
kop(doc, g.TITEL, 0)
alinea(doc, g.ONDERTITEL, schuin=True, kleur=GRIJS, na=4)
alinea(doc, 'Naam: ......................................    Klas: ..........',
       na=14)

alinea(doc,
       f'Vijf werkbladen, samen {g.DUUR} minuten, allemaal rond datums. Een '
       'datum is voor Excel gewoon een getal, en zodra je dat door hebt, kan '
       'je er alles mee: hoe lang iets duurt, hoe oud iemand is, op welke dag '
       'iets valt en hoeveel werkdagen je nog hebt.', na=8)
alinea(doc,
       f'De uitleg staat achteraan in deze bundel, {len(f.FICHES)} fiches. Bij '
       'elk blad staat welke je nodig hebt. Voor de draaitabel, het groeperen '
       'en de draaigrafiek van blad 5 pak je het boekje erbij dat je al hebt: '
       f'{f.THEORIEBUNDEL}.', na=12)

tabel(doc, ['Blad', 'Wat je oefent', 'Fiches', 'Duur', 'Pt'],
      [[f"{oef['nr']}  {oef['blad'][2:]}", oef['titel'],
        ', '.join(str(n) for n in oef['fiches']), oef['duur'],
        punten_van(oef)] for oef in o.OEFENINGEN],
      breedtes=[3.6, 5.6, 3.2, 2.6, 1.6])

kader(doc, 'Drie regels die op elk blad gelden', [
    '1.  Reken met de PEILDATUM die rechtsboven op elk blad staat, niet met '
    f'VANDAAG(). Die peildatum is {g.PEILDATUM.strftime("%d/%m/%Y")}, de '
    'eerste festivaldag. Zo staat je antwoord morgen nog altijd juist — en kan '
    'iemand anders het nakijken.',
    '2.  Zet die peildatum vast met dollartekens voor je doorvoert. Zonder '
    'dollartekens zakt hij mee naar beneden en klopt er onderaan niets meer. '
    'Dat is veruit de fout die hier het vaakst gemaakt wordt.',
    '3.  Krijg je een datum waar je een getal verwacht, of een getal van rond '
    'de 46 000 waar je een datum verwacht? Dan is het de opmaak van de cel en '
    'niet je formule. Rechtsklik, Celeigenschappen, en kies Standaard of Datum.',
], ORANJE)
kader(doc, 'Over de gegevens', [g.WAARSCHUWING], GRIJS)

for oef in o.OEFENINGEN:
    paginaeinde(doc)
    kop(doc, f"Blad {oef['nr']} — {oef['titel']}", 1)
    alinea(doc, f"werkblad '{oef['blad']}'  ·  {oef['duur']}  ·  "
                f"{punten_van(oef)} punten", schuin=True, kleur=GRIJS, na=4)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run('Fiches:  ')
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = GRIJS
    stukken = [f"{n} — {f.FICHES[n - 1]['naam']}" for n in oef['fiches']]
    for n in oef.get('kaarten_elders', ()):
        stukken.append(f"kaart {n} uit {f.THEORIEBUNDEL}")
    r = p.add_run(' · '.join(stukken))
    r.font.size = Pt(10); r.font.color.rgb = ORANJE

    for regel in oef['inleiding']:
        alinea(doc, regel, na=4)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    if oef.get('vullen'):
        kop(doc, 'Wat je maakt', 2, GRIJS)
        tabel(doc, ['Waar', 'Wat erin komt', 'Pt'],
              [[waar, wat, punten] for waar, wat, _formule, punten
               in oef['vullen']],
              breedtes=[3.4, 11.6, 1.4])

    if oef.get('los'):
        cel, wat, _formule, punten = oef['los']
        kop(doc, 'En één losse cel', 2, GRIJS)
        tabel(doc, ['Waar', 'Wat erin komt', 'Pt'],
              [[cel, wat[0].upper() + wat[1:], punten]],
              breedtes=[3.4, 11.6, 1.4])
        alinea(doc, 'Dit is de enige plaats in de bundel waar je VANDAAG() '
                    'gebruikt. Het antwoord is morgen één lager — dat hoort '
                    'zo, en daarom staat er geen vast antwoord op.', na=10)

    if oef.get('stappen'):
        kop(doc, 'Wat je doet', 2, GRIJS)
        tabel(doc, ['Stap', 'Hoe', 'Pt'],
              [[naam, hoe, punten] for naam, hoe, punten in oef['stappen']],
              breedtes=[4.4, 10.6, 1.4])

    kop(doc, 'Lees af en vul in', 2, GRIJS)
    tabel(doc, ['Cel', 'Vraag', 'Pt'],
          [[cel, wat[0].upper() + wat[1:] + '?', punten]
           for cel, wat, _antwoord, punten in oef['vragen']],
          breedtes=[2.4, 12.6, 1.4])

    kader(doc, 'Je blad is af als…', [oef['klaar'], '',
                                      'CONTROLE  ·  ' + oef['controle']], GROEN)
    kader(doc, 'Loop je vast?', oef['vast'], ORANJE)

# ---------------------------------------------------------------- de fiches
vorig_deel = None
for nummer, kaart in enumerate(f.FICHES, 1):
    paginaeinde(doc)
    if kaart['deel'] != vorig_deel:
        kop(doc, kaart['deel'], 1)
        vorig_deel = kaart['deel']
    fiche_uit(doc, nummer, kaart)

paginaeinde(doc)
kop(doc, 'Helemaal klaar?', 1)
alinea(doc, 'Loop dit lijstje af voor je indient.', schuin=True, kleur=GRIJS,
       na=8)
for regel in [
    'Op blad 1 tot 4 staat in elke te maken kolom een formule, geen getypt '
    'antwoord.',
    'Nergens staat nog #WAARDE!, #NAAM? of #GETAL!.',
    'Overal waar je een aantal dagen verwacht, staat een getal — geen datum '
    'in 1900.',
    'Overal waar je een datum verwacht, staat een datum — geen getal van rond '
    'de 46 000.',
    'De peildatum staat in al je formules met dollartekens.',
    'Op blad 5 staat een tabel, een draaitabel per maand, een draaigrafiek en '
    'een tijdlijn.',
    'Je hebt elke vraag ingevuld.',
    'Je bestand heet festival-datums-jouwnaam.xlsx.',
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(4)
    p.add_run('☐   ').font.size = Pt(13)
    zin(p, regel)

doc.save(OPDRACHT)
print('geschreven:', OPDRACHT)


# ==========================================================================
#  2. de lerarensleutel
# ==========================================================================
doc = Document()
opzet(doc)
kop(doc, g.TITEL + ' — lerarensleutel', 0)
alinea(doc, g.ONDERTITEL, schuin=True, kleur=GRIJS, na=12)

kop(doc, 'In het kort', 2, GRIJS)
alinea(doc,
       f'Vijf bladen, {g.DUUR} minuten, {o.TOTAAL} punten. Zeven nieuwe '
       'functies, allemaal rond datums, en op blad 5 geen enkele formule maar '
       'vier menuhandelingen: Ctrl+T, een draaitabel die per maand groepeert, '
       'een draaigrafiek en een tijdlijn.', na=6)
alinea(doc,
       'De fiches staan in de opdracht zelf — dit is een nieuw onderwerp, dus '
       'er is geen apart boekje. Wat ze wél moeten meebrengen is '
       f'{f.THEORIEBUNDEL}, voor de draaitabel, het groeperen (kaart 8) en de '
       'draaigrafiek van blad 5.', na=6)
alinea(doc,
       'Let op de PEILDATUM. Die staat rechtsboven op elk blad en is '
       f'{g.PEILDATUM.strftime("%d/%m/%Y")}, de eerste festivaldag. Er wordt '
       'overal op gerekend, zodat de antwoorden vastliggen en jij ze kan '
       'nakijken. VANDAAG() komt één keer aan bod, in één losse cel op blad 2, '
       'en daar staat met zoveel woorden bij dat het antwoord elke dag '
       'verandert.', na=6)
alinea(doc,
       'Blad 4 is het zwaarste: WERKDAG en NETTO.WERKDAGEN lijken op elkaar en '
       'doen het omgekeerde van elkaar. Blad 1 is de proefrit en duurt acht '
       'minuten; wie daar vastloopt op de celopmaak, loopt de hele bundel '
       'vast.', na=10)

tabel(doc, ['Blad', 'Wat', 'Functies', 'Duur', 'Punten'],
      [[f"{oef['nr']}  {oef['blad'][2:]}", oef['titel'],
        ', '.join(f.FICHES[n - 1]['naam'] for n in oef['fiches']),
        oef['duur'], punten_van(oef)] for oef in o.OEFENINGEN]
      + [['Samen', '', '', f'{g.DUUR} minuten', o.TOTAAL]],
      breedtes=[3.2, 4.4, 5.4, 2.4, 2.0])

kader(doc, 'Waar het bij iedereen misgaat', [
    '1.  De peildatum niet vastzetten. Bij het doorvoeren zakt hij mee en dan '
    'klopt de bovenste rij wel en de rest niet. Herkenbaar aan een kolom die '
    'bovenaan goed is en onderaan onzin, of aan een kolom "Op tijd?" die '
    'overal hetzelfde zegt.',
    '2.  De celopmaak. Een aantal dagen dat als 12/04/1900 verschijnt, of een '
    'einddatum die als 46 223 verschijnt. De formule klopt dan, de cel niet. '
    'Zeg het één keer klassikaal en je bespaart jezelf tien keer dezelfde '
    'vraag.',
    '3.  De ;2 vergeten bij WEEKDAG. Dat valt niet op, want er komt gewoon een '
    f'getal uit. Het weekend telt dan {a.B3_WEEKENDDAGEN + 2} of meer dagen in '
    f'plaats van {a.B3_WEEKENDDAGEN}.',
    '4.  WERKDAG en NETTO.WERKDAGEN verwisselen. De eerste geeft een DATUM, de '
    'tweede een AANTAL. Als hun kolom D getallen toont en kolom E datums, dan '
    'staan ze omgekeerd.',
    '5.  Op blad 5 vergeten te groeperen. Hun draaitabel toont dan elke datum '
    f'apart ({len(set(d for d, _k, _t in g.B5))} regels) in plaats van '
    f'{len(a.B5_MAANDEN)} maanden. Het eindtotaal klopt wel, dus ze merken het '
    'zelf niet.',
], ORANJE)

kader(doc, 'Wat je niet automatisch kan nakijken', [
    'De tabel, de draaigrafiek en de tijdlijn van blad 5 moet je zien. Een '
    'tabel herken je aan de kleur en de filterpijltjes, een draaigrafiek aan '
    'de knoppen met veldnamen erin, en een tijdlijn aan de schuifbalk met '
    'maanden.',
    'Snelle test voor de tijdlijn: klik op een maand en kijk of de grafiek '
    'meebeweegt. Doet hij dat niet, dan is het een gewone grafiek op de tabel '
    'en geen draaigrafiek op de draaitabel.',
    f'De losse cel op blad 2 kan je niet op waarde nakijken — die verandert '
    'elke dag. Kijk alleen of er een formule staat met VANDAAG() erin, en of '
    'de peildatum erin vastgezet is.',
], BLAUW)

for oef in o.OEFENINGEN:
    paginaeinde(doc)
    kop(doc, f"Blad {oef['nr']} — {oef['titel']}", 1)
    alinea(doc, f"werkblad '{oef['blad']}'  ·  {oef['duur']}  ·  "
                f"{punten_van(oef)} punten", schuin=True, kleur=GRIJS, na=8)

    if oef.get('vullen'):
        kop(doc, 'De formules', 2, GRIJS)
        tabel(doc, ['Waar', 'Formule', 'Pt'],
              [[waar, formule, punten]
               for waar, _wat, formule, punten in oef['vullen']],
              breedtes=[3.0, 12.0, 1.4], vaste_kolommen={1})
        alinea(doc, 'Alleen de bovenste cel moet getypt zijn; de rest is '
                    'doorgevoerd. Staat er in elke cel een formule met hetzelfde '
                    'rijnummer, dan is ze gekopieerd in plaats van doorgevoerd '
                    'en klopt alleen de eerste regel.',
               schuin=True, kleur=GRIJS, na=10)

    if oef.get('los'):
        cel, wat, formule, punten = oef['los']
        kop(doc, 'De losse cel', 2, GRIJS)
        tabel(doc, ['Waar', 'Formule', 'Pt'], [[cel, formule, punten]],
              breedtes=[3.0, 12.0, 1.4], vaste_kolommen={1})
        alinea(doc, 'Geen vast antwoord: dit getal verandert elke dag. Kijk of '
                    'er een formule staat en of de peildatum vastgezet is.',
               schuin=True, kleur=GRIJS, na=10)

    if oef.get('stappen'):
        kop(doc, 'De stappen', 2, GRIJS)
        tabel(doc, ['Stap', 'Hoe', 'Pt'],
              [[naam, hoe, punten] for naam, hoe, punten in oef['stappen']],
              breedtes=[4.0, 11.0, 1.4])

    nr = oef['nr']
    if nr == 1:
        tabel(doc, ['Taak', 'Start', 'Einde', 'Dagen'],
              [[taak, a._dat(s1), a._dat(e1), a.B1_DAGEN[i]]
               for i, (taak, s1, e1) in enumerate(g.B1)],
              'Light List Accent 1')
        alinea(doc, f'De langste taak is {a.B1_LANGSTE} met '
                    f'{a.B1_LANGSTE_DAGEN} dagen. Van '
                    f'{a._dat(a.B1_VROEGSTE_START)} tot '
                    f'{a._dat(a.B1_LAATSTE_EINDE)} is {a.B1_SPANNE} dagen.',
               schuin=True, kleur=GRIJS, na=10)
    elif nr == 2:
        tabel(doc, ['Vrijwilliger', 'Geboren', 'Leeftijd', 'Achter de bar?'],
              [[naam, a._dat(gb), a.B2_LEEFTIJD[i], a.B2_MAG[i]]
               for i, (naam, gb) in enumerate(g.B2)], 'Light List Accent 1')
        nipt = '; '.join(f'{naam} is {lft}' for naam, lft in a.B2_NIPT)
        alinea(doc, f'{a.B2_AANTAL_MAG} van de {len(g.B2)} mogen achter de '
                    f'bar. Let op de twee in de festivalweek: {nipt}. Eén dag '
                    'scheelt daar een heel jaar — dat is het gesprek dat dit '
                    'blad wil uitlokken.', schuin=True, kleur=GRIJS, na=10)
    elif nr == 3:
        tabel(doc, ['Datum', 'Tickets', 'Maand', 'Weekdag', 'Soort dag'],
              [[a._dat(d), t, a.B3_MAAND[i], a.B3_WEEKDAG[i], a.B3_SOORT[i]]
               for i, (d, t) in enumerate(g.B3)], 'Light List Accent 1')
        alinea(doc, f'{a.B3_WEEKENDDAGEN} van de {len(g.B3)} dagen vielen in '
                    f'het weekend. De drukste dag is '
                    f'{a._dat(a.B3_DRUKSTE_DATUM)}, een '
                    f'{a.B3_DRUKSTE_DAGNAAM} — weekdagnummer '
                    f'{a.B3_DRUKSTE_WEEKDAG}.',
               schuin=True, kleur=GRIJS, na=10)
    elif nr == 4:
        tabel(doc, ['Opbouwtaak', 'Start', 'Werkdagen', 'Klaar op',
                    'Tot het festival', 'Op tijd?'],
              [[taak, a._dat(s4), d4, a._dat(a.B4_KLAAR[i]),
                a.B4_WERKDAGEN_TOT[i], a.B4_OP_TIJD[i]]
               for i, (taak, s4, d4) in enumerate(g.B4)],
              'Light List Accent 1')
        alinea(doc, f'Precies één taak komt te laat: {a.B4_TE_LAAT_TAAK}, '
                    f'klaar op {a._dat(a.B4_TE_LAAT_KLAAR)} terwijl het '
                    f'festival op {a._dat(g.PEILDATUM)} begint. '
                    f'{a.B4_ZWAARSTE} is de zwaarste taak '
                    f'({a.B4_ZWAARSTE_DAGEN} werkdagen) en haalt het nét.',
               schuin=True, kleur=GRIJS, na=10)
    elif nr == 5:
        tabel(doc, ['Maand', 'Tickets'],
              [[g.maandnaam(m), a._getal(a.B5_PER_MAAND[m])]
               for m in a.B5_MAANDEN]
              + [['Eindtotaal', a._getal(a.B5_TOTAAL)]],
              'Light List Accent 1', [5.0, 5.0])
        alinea(doc, f'De beste maand is {a.B5_TOPMAAND} met '
                    f'{a._getal(a.B5_TOPMAAND_TICKETS)} tickets. Met de '
                    f'tijdlijn op {g.B5_TIJDLIJN_MAAND} blijft er '
                    f'{a._getal(a.B5_TIJDLIJN_TICKETS)} over en toont de '
                    'grafiek nog één balk.', schuin=True, kleur=GRIJS, na=4)
        alinea(doc, f'Het eindtotaal {a._getal(a.B5_TOTAAL)} is precies de som '
                    'van alle tickets op blad 3. Dat is geen toeval: blad 5 is '
                    'dezelfde verkoop, uitgesplitst per kanaal. Wie dat '
                    'opmerkt, heeft goed gekeken.',
               schuin=True, kleur=GRIJS, na=10)

    kop(doc, 'De antwoorden', 2, GRIJS)
    tabel(doc, ['Cel', 'Vraag', 'Antwoord', 'Pt'],
          [[cel, wat[0].upper() + wat[1:], antwoord, punten]
           for cel, wat, antwoord, punten in oef['vragen']],
          breedtes=[2.0, 8.6, 4.4, 1.4])
    kader(doc, 'Waar je naar kijkt', [oef['controle']], GROEN)

paginaeinde(doc)
kop(doc, 'Tijdsindeling en tips', 1)
tabel(doc, ['Wanneer', 'Wat'],
      [['0 - 5 min', 'Bestand opslaan onder eigen naam. Wijs de peildatum '
                     'rechtsboven aan en zeg waarom hij er staat. Herhaal dat '
                     'ze er dollartekens om moeten zetten.'],
       ['5 - 13 min', 'Blad 1. De proefrit. Hier komt de celopmaak voor het '
                      'eerst boven: laat zien hoe je een cel van Datum naar '
                      'Standaard zet, dan hoef je dat maar één keer te doen.'],
       ['13 - 23 min', 'Blad 2. DATUMVERSCHIL staat in geen enkel menu; zet ze '
                       'één keer op het bord. Vraag daarna klassikaal naar de '
                       'twee die in de festivalweek jarig zijn.'],
       ['23 - 33 min', 'Blad 3. Gaat vlot. De enige valkuil is de ;2 in '
                       'WEEKDAG.'],
       ['33 - 45 min', 'Blad 4. Het zwaarste blad. Zet WERKDAG en '
                       'NETTO.WERKDAGEN naast elkaar op het bord: de eerste '
                       'geeft een datum, de tweede een aantal.'],
       ['45 - 50 min', 'Blad 5. Alleen de snelsten halen dit. Dat mag — zie '
                       'hieronder.']],
      breedtes=[3.4, 13.0])

kader(doc, 'Als de tijd te kort is', [
    'Blad 5 spreid je het best over twee lessen. Het staat op zichzelf en '
    'heeft geen enkele formule nodig; je kan het ook als een losse les van '
    'twintig minuten geven, met de hele klas tegelijk aan het bord. De '
    'tijdlijn is trouwens het stuk waar ze het meest van onder de indruk zijn.',
    'Laat eventueel blad 3 vallen. MAAND en WEEKDAG zijn de makkelijkste '
    'functies van de bundel en komen op blad 5 nog terug in het groeperen.',
    'Wat je niet moet laten vallen is blad 1. Dat een datum een getal is, is '
    'de hele bundel in één zin.',
], BLAUW)

kader(doc, 'Wat hierbij hoort', [
    f'De zusterbundel {g.ZUSTERBUNDEL} werkt met hetzelfde festival, maar dan '
    'met de tekstfuncties: een line-up uit elkaar halen, namen opkuisen, '
    'e-mailadressen bouwen. Daar eindigt blad 5 op een draaigrafiek met een '
    'slicer, hier op een draaigrafiek met een tijdlijn.',
    'De twee bundels kunnen in elke volgorde en zijn allebei vijftig minuten.',
], BLAUW)

doc.save(SLEUTEL)
print('geschreven:', SLEUTEL)
