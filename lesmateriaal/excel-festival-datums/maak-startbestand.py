# -*- coding: utf-8 -*-
"""Bouwt het startbestand van de datumbundel.

Alle datums worden als echte datums geschreven, niet als tekst: anders werkt
er geen enkele formule op dit blad en blijft de tijdlijn van blad 5 grijs.

Rechtsboven op elk blad staat de peildatum. Daar rekenen ze op, zodat de
antwoorden niet elke dag veranderen.
"""
import importlib.util
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')

UIT = os.path.join(HIER, 'festival-datums-startbestand.xlsx')
TITEL = Font(bold=True, size=13)
GROOT = Font(bold=True, size=16)
VET = Font(bold=True)
SCHUIN = Font(italic=True, size=10)

wb = Workbook()


def breedtes(ws, waarden, start=1):
    for i, b in enumerate(waarden, start=start):
        ws.column_dimensions[get_column_letter(i)].width = b


def blad(oef, titel, ondertitel=None):
    ws = wb.create_sheet(oef['blad'])
    ws['A1'] = titel
    ws['A1'].font = TITEL
    if ondertitel:
        ws['A2'] = ondertitel
        ws['A2'].font = SCHUIN
    # de peildatum, rechtsboven, op elk blad op dezelfde plaats
    k = oef['vraagkol']
    ws.cell(row=o.PEILRIJ, column=k, value=g.PEIL_LABEL).font = VET
    c = ws.cell(row=o.PEILRIJ, column=k + 1, value=g.PEILDATUM)
    c.number_format = g.DATUMNOTATIE
    c.font = VET
    return ws


def lijst(ws, koprij, koppen, rijen):
    for i, k in enumerate(koppen, start=1):
        ws.cell(row=koprij, column=i, value=k).font = VET
    for r, waarden in enumerate(rijen, start=koprij + 1):
        for i, w in enumerate(waarden, start=1):
            c = ws.cell(row=r, column=i, value=w)
            if hasattr(w, 'year'):
                c.number_format = g.DATUMNOTATIE


def vragen(ws, rij, kolom, teksten):
    ws.cell(row=rij, column=kolom, value='Vraag').font = VET
    ws.cell(row=rij, column=kolom + 1, value='Antwoord').font = VET
    for i, t in enumerate(teksten, start=1):
        ws.cell(row=rij + i, column=kolom, value=t)


# ------------------------------------------------------------------- Start
ws = wb.active
ws.title = 'Start'
ws['A1'] = g.TITEL
ws['A1'].font = GROOT
KOP_REGELS = 'Drie regels die op elk blad gelden'
regels = [
    g.ONDERTITEL, '',
    'Naam: ......................................    Klas: ..........', '',
    f'Vijf werkbladen, samen {g.DUUR} minuten.',
    'De uitleg staat in de opdrachtbundel; bij elk blad staat welke fiches je',
    'erbij nodig hebt. Voor de draaitabel, het groeperen en de draaigrafiek',
    'van blad 5 pak je het boekje erbij dat je al hebt.', '',
]
for oef in o.OEFENINGEN:
    fiches = 'fiche ' + ', '.join(str(n) for n in oef['fiches'])
    regels.append(f"Blad {oef['nr']}  {oef['blad'][2:]:<18} "
                  f"{oef['duur'].split()[0]:>2} min   {fiches:<14} "
                  f"{oef['titel']}")
regels += [
    '', KOP_REGELS,
    '  1. Reken met de PEILDATUM die rechtsboven op elk blad staat, niet met',
    f'     VANDAAG(). Die peildatum is {g.PEILDATUM.strftime("%d/%m/%Y")}, de',
    '     eerste festivaldag. Zo staat je antwoord morgen nog altijd juist.',
    '  2. Zet die peildatum vast met dollartekens voor je doorvoert. Zonder',
    '     dollartekens zakt hij mee naar beneden en klopt er onderaan niets.',
    '  3. Krijg je een datum waar je een getal verwacht, of omgekeerd? Dat is',
    '     de opmaak van de cel, niet je formule. Rechtsklik, Celeigenschappen.',
    '',
    'EERST DOEN: sla dit bestand op als  festival-datums-jouwnaam.xlsx',
    '', g.WAARSCHUWING,
]
for i, r in enumerate(regels, start=2):
    ws.cell(row=i, column=1, value=r)
ws.cell(row=2 + regels.index(KOP_REGELS), column=1).font = VET
rij_waarschuwing = 2 + regels.index(g.WAARSCHUWING)
c = ws.cell(row=rij_waarschuwing, column=1)
c.font = SCHUIN
c.alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[rij_waarschuwing].height = 86
breedtes(ws, [92])

# ---------------------------------------------------------- 1 Voorbereiding
oef = o.OEFENINGEN[0]
ws = blad(oef, 'De voorbereiding', g.B1_VOETNOOT)
lijst(ws, g.B1_KOPRIJ, g.B1_KOP,
      [(taak, start, einde, None) for taak, start, einde in g.B1])
vragen(ws, g.B1_KOPRIJ, oef['vraagkol'], [
    'Welke taak duurt het langst?',
    'Hoeveel dagen zijn dat?',
    'Hoeveel dagen tussen de vroegste start en het laatste einde?'])
breedtes(ws, [26, 13, 13, 9])
breedtes(ws, [3, 46, 16], start=oef['vraagkol'] - 1)

# ---------------------------------------------------------- 2 Vrijwilligers
oef = o.OEFENINGEN[1]
ws = blad(oef, 'De vrijwilligers', g.B2_VOETNOOT)
lijst(ws, g.B2_KOPRIJ, g.B2_KOP,
      [(naam, geboorte, None, None) for naam, geboorte in g.B2])
k = oef['vraagkol']
ws.cell(row=o.LOSRIJ, column=k, value=oef['los'][1][0].upper()
        + oef['los'][1][1:] + '?').font = SCHUIN
vragen(ws, g.B2_KOPRIJ, k, [
    'Hoeveel vrijwilligers mogen er achter de bar?',
    'Wie is de jongste?',
    'Hoe oud is de oudste?'])
breedtes(ws, [22, 15, 28, 15])
breedtes(ws, [3, 46, 16], start=k - 1)

# ---------------------------------------------------------- 3 Ticketverkoop
oef = o.OEFENINGEN[2]
ws = blad(oef, 'De ticketverkoop per dag', g.B3_VOETNOOT)
lijst(ws, g.B3_KOPRIJ, g.B3_KOP,
      [(datum, tickets, None, None, None) for datum, tickets in g.B3])
vragen(ws, g.B3_KOPRIJ, oef['vraagkol'], [
    'Hoeveel verkoopdagen vielen er in het weekend?',
    'Hoeveel tickets zijn er op de drukste dag verkocht?',
    'Welk weekdagnummer heeft die drukste dag?'])
breedtes(ws, [13, 10, 9, 16, 13])
breedtes(ws, [3, 44, 16], start=oef['vraagkol'] - 1)

# ----------------------------------------------------------------- 4 Opbouw
oef = o.OEFENINGEN[3]
ws = blad(oef, 'De opbouw van het terrein', g.B4_VOETNOOT)
lijst(ws, g.B4_KOPRIJ, g.B4_KOP,
      [(taak, start, dagen, None, None, None) for taak, start, dagen in g.B4])
vragen(ws, g.B4_KOPRIJ, oef['vraagkol'], [
    'Welke taak komt te laat?',
    'Welke taak heeft de meeste werkdagen nodig?',
    'Wat is het hoogste getal in kolom E?'])
breedtes(ws, [24, 13, 15, 13, 24, 10])
breedtes(ws, [3, 40, 20], start=oef['vraagkol'] - 1)

# ------------------------------------------------------- 5 Verkoop per maand
oef = o.OEFENINGEN[4]
ws = blad(oef, 'Alle verkoopregels', g.B5_VOETNOOT)
lijst(ws, g.B5_KOPRIJ, g.B5_KOP, g.B5)
vragen(ws, g.B5_KOPRIJ, oef['vraagkol'], [
    'In welke maand zijn er de meeste tickets verkocht?',
    'Wat is het eindtotaal van je draaitabel?',
    f'Hoeveel tickets in {g.B5_TIJDLIJN_MAAND}, met de tijdlijn?'])
breedtes(ws, [13, 12, 10])
breedtes(ws, [3, 44, 16], start=oef['vraagkol'] - 1)

wb.save(UIT)
print('geschreven:', UIT)
for ws in wb.worksheets:
    gevuld = sum(1 for rij in ws.iter_rows() for c in rij if c.value is not None)
    print(f'  {ws.title:<20} {gevuld:>4} gevulde cellen   {ws.dimensions}')
