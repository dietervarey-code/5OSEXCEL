# -*- coding: utf-8 -*-
"""Kijkt na of opdracht, startbestand, antwoorden en sleutel bij elkaar passen.

Faalt luid zodra er iets uit elkaar loopt. Geen vaste rijnummers: alles komt
uit gegevens.py.
"""
import collections
import importlib.util
import os
import re
import sys
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')

WERKBOEK = os.path.join(HIER, 'draaitabellen-startbestand.xlsx')
FOUTEN, GETELD = [], collections.Counter()


def eis(voorwaarde, bericht):
    if not voorwaarde:
        FOUTEN.append(bericht)
    return bool(voorwaarde)


def cel_van(adres):
    m = re.match(r'^([A-Z]{1,2})(\d+)$', adres)
    return (column_index_from_string(m.group(1)), int(m.group(2))) if m else None


# ==========================================================================
#  1. de gegevens zijn intern samenhangend
# ==========================================================================
for naam, stad, regio, jaar, punten, dp, toe in g.B1:
    eis(regio in g.REGIOS, f'onbekende regio {regio!r} bij {naam}')
    eis(1850 < jaar <= 2026, f'onwaarschijnlijk oprichtingsjaar bij {naam}')
    eis(punten > 0 and dp > 0 and toe > 0, f'onmogelijk getal bij {naam}')
    GETELD['clubs nagerekend'] += 1
eis(len({r[0] for r in g.B1}) == len(g.B1), 'dubbele clubnaam')

for gerecht, soort, regio, pittig, tijd, porties in g.B2:
    eis(soort in g.B2_SOORTEN, f'onbekende soort {soort!r} bij {gerecht}')
    eis(regio in g.REGIOS, f'onbekende regio {regio!r} bij {gerecht}')
    eis(pittig in ('ja', 'nee'), f'pittig moet ja of nee zijn bij {gerecht}')
    eis(tijd > 0 and porties > 0, f'onmogelijk getal bij {gerecht}')
    GETELD['gerechten nagerekend'] += 1

for prov, regio, inw, opp, zee in g.B3:
    eis(regio in g.REGIOS, f'onbekende regio {regio!r} bij {prov}')
    eis(inw % 10000 == 0, f'{prov}: inwoners horen afgerond te zijn, niet {inw}')
    eis(opp > 0 and zee in ('ja', 'nee'), f'onmogelijke waarde bij {prov}')
    GETELD['provincies nagerekend'] += 1
eis(len({r[0] for r in g.B3}) == len(g.B3), 'dubbele provincie')

for serie, genre, jaar, seiz, afl, kijkers in g.B4:
    eis(genre in g.B4_GENRES, f'onbekend genre {genre!r} bij {serie}')
    eis(g.B4_GROEP_VAN <= jaar <= g.B4_GROEP_TOT,
        f'{serie} ({jaar}) valt buiten het groepeerbereik '
        f'{g.B4_GROEP_VAN}-{g.B4_GROEP_TOT}')
    eis(seiz > 0 and afl > 0 and kijkers > 0, f'onmogelijk getal bij {serie}')
    GETELD['series nagerekend'] += 1

for speler, positie, nummer, bedrag in g.B5 + [g.B5_EXTRA]:
    eis(positie in g.B5_POSITIES, f'onbekende positie bij {speler}')
    eis(1 <= nummer <= len(g.B1), f'clubnummer {nummer} bestaat niet ({speler})')
    eis(bedrag > 0, f'bedrag moet positief zijn bij {speler}')
    GETELD['transfers nagerekend'] += 1
eis(len({r[0] for r in g.B5}) == len(g.B5), 'dubbele spelersnaam')

# elke regio in de lijsten komt ook echt voor
for lijst, kolom, naam in ((g.B1, 2, 'clubs'), (g.B2, 2, 'gerechten'),
                           (g.B3, 1, 'provincies')):
    gebruikt = {r[kolom] for r in lijst}
    eis(gebruikt <= set(g.REGIOS), f'onbekende regio in de {naam}')

# ==========================================================================
#  2. de antwoorden kloppen, langs een tweede weg uitgerekend
# ==========================================================================
eis(a.D1_TOTAAL == len(g.B1), 'eindtotaal van draaitabel 1 klopt niet')
for regio, aantal in a.D1.items():
    eis(aantal == sum(1 for r in g.B1 if r[2] == regio), f'D1 klopt niet bij {regio}')
    GETELD['cellen nagerekend'] += 1

eis(a.D2_TOTAAL == len(g.B2), 'eindtotaal van draaitabel 2 klopt niet')
for r in a.D2_RIJEN:
    for k in a.D2_KOLLEN:
        verwacht = sum(1 for x in g.B2 if x[1] == r and x[3] == k)
        eis(a.D2_CEL[(r, k)] == verwacht, f'D2 klopt niet bij {r}/{k}')
        GETELD['cellen nagerekend'] += 1

eis(a.D3_TOTAAL == sum(r[2] for r in g.B3), 'eindtotaal van draaitabel 3 klopt niet')
for regio, inw in a.D3.items():
    eis(inw == sum(r[2] for r in g.B3 if r[1] == regio), f'D3 klopt niet bij {regio}')
    GETELD['cellen nagerekend'] += 1

for d in a.D4_LABELS:
    hoort = [r[5] for r in g.B4 if g.decennium(r[2]) == d]
    eis(a.D4_AANTAL[d] == len(hoort), f'D4 aantal klopt niet bij {d}')
    eis(abs(a.D4_GEMIDDELD[d] - sum(hoort) / len(hoort)) < 1e-9,
        f'D4 gemiddelde klopt niet bij {d}')
    GETELD['cellen nagerekend'] += 2
eis(a.D4_TOTAAL_AANTAL == len(g.B4), 'eindtotaal van draaitabel 4 klopt niet')

for regio, bedrag in a.D5.items():
    verwacht = sum(b for _, _, n, b in g.B5 if g.club(n)[2] == regio)
    eis(abs(bedrag - verwacht) < 1e-9, f'D5 klopt niet bij {regio}')
    GETELD['cellen nagerekend'] += 1
eis(abs(a.D5_TOTAAL - sum(r[3] for r in g.B5)) < 1e-9, 'D5 eindtotaal klopt niet')
eis(abs(a.D5_NA_TOTAAL - (a.D5_TOTAAL + g.B5_EXTRA[3])) < 1e-9,
    'het eindtotaal na de extra transfer klopt niet')
eis(a.D5_NA != a.D5, 'de extra transfer verandert niets — dan leert ze niets')

# elke vraag moet precies één antwoord hebben; antwoorden.py bewaakt dat al met
# een assert bij het inlezen, maar de tekstantwoorden controleren we hier nog eens
for oef in o.OEFENINGEN:
    for cel, wat, antwoord, punten in oef['vragen']:
        eis(antwoord is not None and antwoord != '',
            f"{oef['blad']} {cel}: geen antwoord")
        eis(punten > 0, f"{oef['blad']} {cel}: geen punten")
        GETELD['vragen gekeurd'] += 1

# ==========================================================================
#  3. het startbestand
# ==========================================================================
LIJSTEN = [
    ('1 Süper Lig', g.B1_KOPRIJ, g.B1_START, g.B1_EIND, g.B1_KOP, o.V1, g.B1_ANKER),
    ('2 Keuken', g.B2_KOPRIJ, g.B2_START, g.B2_EIND, g.B2_KOP, o.V2, g.B2_ANKER),
    ('3 Provincies', g.B3_KOPRIJ, g.B3_START, g.B3_EIND, g.B3_KOP, o.V3, g.B3_ANKER),
    ('4 Series', g.B4_KOPRIJ, g.B4_START, g.B4_EIND, g.B4_KOP, o.V4, g.B4_ANKER),
    ('5 Transfers', g.B5_KOPRIJ, g.B5_START, g.B5_EIND, g.B5_KOP, o.V5, g.B5_ANKER),
]

eis(os.path.exists(WERKBOEK), 'het startbestand bestaat niet — draai eerst '
                              'maak-startbestand.py')
if os.path.exists(WERKBOEK):
    wb = load_workbook(WERKBOEK)
    eis({x[0] for x in LIJSTEN} <= set(wb.sheetnames),
        f'ontbrekende bladen: {sorted({x[0] for x in LIJSTEN} - set(wb.sheetnames))}')

    for blad, koprij, start, eind, koppen, vraagkol, anker in LIJSTEN:
        if blad not in wb.sheetnames:
            continue
        ws = wb[blad]
        breedte = len(koppen)
        for i, titel in enumerate(koppen, start=1):
            eis(ws.cell(row=koprij, column=i).value == titel,
                f'{blad}: koprij verschilt in kolom {get_column_letter(i)}')
        eis(len(set(koppen)) == breedte, f'{blad}: twee kolommen heten hetzelfde')

        # te maken kolommen mogen leeg zijn; de rest moet gevuld zijn
        temaken = set()
        for oef in o.OEFENINGEN:
            if oef['blad'] != blad:
                continue
            for waar, *_rest in oef.get('vullen', []):
                temaken.add(column_index_from_string(
                    re.match(r'([A-Z]+)', waar).group(1)))
        for rij in range(start, eind + 1):
            gevuld = [c for c in range(1, breedte + 1)
                      if ws.cell(row=rij, column=c).value is not None]
            eis(gevuld, f'{blad}: rij {rij} is helemaal leeg')
        for kolom in range(1, breedte + 1):
            heeft = any(ws.cell(row=r, column=kolom).value is not None
                        for r in range(start, eind + 1))
            eis(heeft or kolom in temaken,
                f'{blad}: kolom {get_column_letter(kolom)} blijft leeg')

        # één lege kolom tussen de lijst en het vragenblok
        eis(vraagkol >= breedte + 2,
            f'{blad}: het vragenblok staat in kolom {get_column_letter(vraagkol)}, '
            f'de lijst loopt tot {get_column_letter(breedte)} — er hoort minstens '
            'één lege kolom tussen')
        for rij in range(start, eind + 1):
            eis(ws.cell(row=rij, column=breedte + 1).value is None,
                f'{blad}: kolom {get_column_letter(breedte + 1)} rij {rij} hoort '
                'leeg te blijven als scheiding')

        # de draaitabel staat onder de lijst, met ruimte ertussen
        kolom_anker, rij_anker = cel_van(anker)
        eis(kolom_anker == 1, f'{blad}: het anker {anker} staat niet in kolom A')
        ruimte = rij_anker - eind
        eis(ruimte >= 3,
            f'{blad}: de draaitabel begint op rij {rij_anker}, maar de lijst loopt '
            f'tot {eind} — laat minstens twee lege rijen ertussen')
        eis(ws.cell(row=rij_anker, column=1).value is None,
            f'{blad}: het anker {anker} is niet vrij')
        for kolom in range(1, breedte + 1):
            eis(ws.cell(row=eind + 1, column=kolom).value is None,
                f'{blad}: er staat iets op rij {eind + 1}, direct onder de lijst')
        GETELD['lijsten gekeurd'] += 1

    # de antwoordcellen zijn leeg en staan in het vragenblok
    for oef in o.OEFENINGEN:
        ws = wb[oef['blad']]
        for cel, *_rest in oef['vragen']:
            kolom, rij = cel_van(cel)
            eis(ws.cell(row=rij, column=kolom).value is None,
                f"{oef['blad']} {cel} moet leeg zijn in het startbestand")
            eis(ws.cell(row=rij, column=kolom - 1).value is not None,
                f"{oef['blad']} {cel}: er staat geen vraag naast")
            GETELD['antwoordcellen leeg'] += 1
        for waar, *_rest in oef.get('vullen', []):
            k1, r1 = cel_van(waar.split(':')[0])
            k2, r2 = cel_van(waar.split(':')[1])
            for k in range(k1, k2 + 1):
                for r in range(r1, r2 + 1):
                    eis(ws.cell(row=r, column=k).value is None,
                        f"{oef['blad']} {get_column_letter(k)}{r} moet leeg zijn")
                    GETELD['te vullen cellen leeg'] += 1

    # de rij die ze er zelf bij typen, moet nog vrij zijn
    ws = wb['5 Transfers']
    extra_rij = o.OEFENINGEN[4]['extra']['rij']
    for kolom in range(1, len(g.B5_KOP) + 1):
        eis(ws.cell(row=extra_rij, column=kolom).value is None,
            f'5 Transfers: rij {extra_rij} moet vrij zijn — daar typen ze de '
            'extra transfer')
    eis(extra_rij == g.B5_EIND + 1,
        'de extra transfer hoort direct onder de lijst te komen')

    # de clubtabel naast de transferlijst moet kloppen met blad 1
    for i, (club, stad, regio, *_rest) in enumerate(g.B1):
        r = g.B5_ZOEKRIJ + 1 + i
        for kolom, verwacht in ((g.B5_ZOEKKOL, g.clubcode(i + 1)),
                                (g.B5_ZOEKKOL + 1, club),
                                (g.B5_ZOEKKOL + 2, stad),
                                (g.B5_ZOEKKOL + 3, regio)):
            eis(ws.cell(row=r, column=kolom).value == verwacht,
                f'5 Transfers: de clubtabel wijkt af op rij {r}')
            GETELD['clubtabel nagekeken'] += 1

# ==========================================================================
#  4. de opdracht zelf
# ==========================================================================
for oef in o.OEFENINGEN:
    eis(oef['kaarten'], f"blad {oef['nr']}: geen verwijzing naar de theorie")
    for n in oef['kaarten']:
        eis(n in o.KAARTEN, f"blad {oef['nr']}: kaart {n} bestaat niet")
        eis(1 <= n <= 9, f"blad {oef['nr']}: kaart {n} gaat niet over draaitabellen")
    draai = oef['draai']
    koppen = {1: g.B1_KOP, 2: g.B2_KOP, 3: g.B3_KOP,
              4: g.B4_KOP, 5: g.B5_KOP}[oef['nr']]
    for rol in ('rijen', 'kolommen', 'waarden', 'tweede'):
        veld = draai.get(rol)
        eis(veld is None or veld in koppen,
            f"blad {oef['nr']}: {veld!r} is geen kolom van deze lijst ({rol})")
    eis(draai['bewerking'] in ('Aantal', 'Som', 'Gemiddelde', 'Max', 'Min'),
        f"blad {oef['nr']}: onbekende bewerking {draai['bewerking']!r}")
    GETELD['bladen gekeurd'] += 1

# geen theorie in de opdracht: de bundel mag niet uitleggen wat een kaart doet
VERBODEN = ['waardeveldinstelling', 'veldenlijst', 'sleep het veld naar']
tekst = ' '.join(
    r for oef in o.OEFENINGEN
    for r in oef['inleiding'] + [oef['klaar'], oef['controle']]).lower()
for woord in VERBODEN:
    eis(woord not in tekst,
        f'de opdracht legt theorie uit ("{woord}") — dat hoort op de kaarten')

eis(sum(int(oef['duur'].split()[0]) for oef in o.OEFENINGEN) == g.DUUR,
    'de duur van de bladen telt niet op tot de aangekondigde duur')
eis(o.TOTAAL == 50, f'de bundel staat op {o.TOTAAL} punten in plaats van 50')

# ==========================================================================
print('=' * 72)
for naam, hoeveel in sorted(GETELD.items()):
    print(f'  {hoeveel:>5}  {naam}')
print('=' * 72)
if FOUTEN:
    print(f'{len(FOUTEN)} PROBLEMEN:')
    for f in FOUTEN:
        print('  !', f)
    sys.exit(1)
print('alles past bij elkaar.')
