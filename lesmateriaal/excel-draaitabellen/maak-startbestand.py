# -*- coding: utf-8 -*-
"""Bouwt het startbestand van de draaitabelbundel.

Elke lijst staat kaarsrecht: één koprij, geen lege rij erin, niets eronder.
Het vragenblok staat ERNAAST met één lege kolom ertussen, en de draaitabel
komt ONDER de lijst met twee lege rijen ertussen. Zo stopt Excel bij het
bepalen van het bronbereik precies waar het hoort.

Er zit geen draaitabel in het bestand: openpyxl kan die niet schrijven, en de
leerling hoort ze zelf te maken.
"""
import importlib.util
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')

UIT = os.path.join(HIER, 'draaitabellen-startbestand.xlsx')
TITEL = Font(bold=True, size=13)
GROOT = Font(bold=True, size=16)
VET = Font(bold=True)
SCHUIN = Font(italic=True, size=10)

wb = Workbook()


def breedtes(ws, waarden, start=1):
    for i, b in enumerate(waarden, start=start):
        ws.column_dimensions[get_column_letter(i)].width = b


def blad(naam, titel, ondertitel=None):
    ws = wb.create_sheet(naam)
    ws['A1'] = titel
    ws['A1'].font = TITEL
    if ondertitel:
        ws['A2'] = ondertitel
        ws['A2'].font = SCHUIN
    return ws


def lijst(ws, koprij, koppen, rijen):
    for i, k in enumerate(koppen, start=1):
        c = ws.cell(row=koprij, column=i, value=k)
        c.font = VET
    for r, waarden in enumerate(rijen, start=koprij + 1):
        for i, w in enumerate(waarden, start=1):
            ws.cell(row=r, column=i, value=w)


def vragen(ws, rij, kolom, teksten):
    ws.cell(row=rij, column=kolom, value='Vraag').font = VET
    ws.cell(row=rij, column=kolom + 1, value='Antwoord').font = VET
    for i, t in enumerate(teksten, start=1):
        ws.cell(row=rij + i, column=kolom, value=t)


def plek(ws, oef):
    """Het plekje waar de draaitabel moet komen, met een regel erboven."""
    anker = oef['draai']['anker']
    rij = int(anker[1:])
    ws.cell(row=rij - 1, column=1,
            value=f"Zet hier je draaitabel — linkerbovenhoek in {anker}."
            ).font = SCHUIN


# ------------------------------------------------------------------- Start
ws = wb.active
ws.title = 'Start'
ws['A1'] = g.TITEL
ws['A1'].font = GROOT
regels = [
    g.ONDERTITEL, '',
    'Naam: ......................................    Klas: ..........', '',
    f'Vijf werkbladen, samen {g.DUUR} minuten. Op elk blad maak je één draaitabel.',
    'Er staat in deze bundel geen theorie. Die staat in:',
    f'   {g.THEORIEBUNDEL} — elke kaart op een eigen pagina.',
    'Bij elk blad staat welke kaarten je erbij nodig hebt.', '',
    'Blad 1  Süper Lig     6 min   kaart 3 en 4    je eerste draaitabel',
    'Blad 2  Keuken        8 min   kaart 6 en 5    een kruistabel',
    'Blad 3  Provincies    8 min   kaart 5 en 7    optellen en sorteren',
    'Blad 4  Series       12 min   kaart 8, 5, 7   groeperen en filteren',
    'Blad 5  Transfers    16 min   kaart 2 en 9    eerst aanvullen, dan samenvatten',
    '',
    'Drie regels die op elk blad gelden',
    '  1. De lijst blijft zoals ze is. Voeg er geen lege rij in en zet er niets',
    '     onder. De vragen staan ERNAAST, met een lege kolom ertussen.',
    '  2. De draaitabel komt ONDER de lijst, op de plaats die aangeduid staat.',
    '  3. Lees het kopje van je draaitabel. Staat er Aantal waar Som hoort,',
    '     dan klopt je hele tabel niet — en Excel zegt daar niets over.',
    '',
    'EERST DOEN: sla dit bestand op als  draaitabellen-jouwnaam.xlsx',
    '', g.WAARSCHUWING,
]
for i, r in enumerate(regels, start=2):
    ws.cell(row=i, column=1, value=r)
ws['A17'].font = VET
ws['A25'].font = SCHUIN
ws['A25'].alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[25].height = 86
breedtes(ws, [86])

# -------------------------------------------------------------- 1 Süper Lig
oef = o.OEFENINGEN[0]
ws = blad(oef['blad'], 'De Süper Lig')
lijst(ws, g.B1_KOPRIJ, g.B1_KOP, g.B1)
vragen(ws, g.B1_KOPRIJ, o.V1, [
    'Welke regio heeft de meeste clubs?', 'Hoeveel clubs zijn dat?'])
plek(ws, oef)
breedtes(ws, [22, 12, 12, 11, 9, 12, 13])
breedtes(ws, [3, 32, 14], start=o.V1 - 1)

# ----------------------------------------------------------------- 2 Keuken
oef = o.OEFENINGEN[1]
ws = blad(oef['blad'], 'De Turkse keuken', g.B2_VOETNOOT)
lijst(ws, g.B2_KOPRIJ, g.B2_KOP, g.B2)
vragen(ws, g.B2_KOPRIJ, o.V2, [
    'Welke soort heeft de meeste pittige gerechten?',
    'Hoeveel niet-pittige nagerechten zijn er?',
    'Welke soort bevat geen enkel pittig gerecht?'])
plek(ws, oef)
breedtes(ws, [20, 15, 20, 9, 16, 9])
breedtes(ws, [3, 42, 14], start=o.V2 - 1)

# ------------------------------------------------------------- 3 Provincies
oef = o.OEFENINGEN[2]
ws = blad(oef['blad'], 'De grootste provincies', g.B3_VOETNOOT)
lijst(ws, g.B3_KOPRIJ, g.B3_KOP, g.B3)
vragen(ws, g.B3_KOPRIJ, o.V3, [
    'Welke regio telt de meeste inwoners?', 'Hoeveel inwoners zijn dat?',
    'Welke regio telt de minste inwoners?'])
plek(ws, oef)
breedtes(ws, [18, 20, 13, 17, 9])
breedtes(ws, [3, 36, 14], start=o.V3 - 1)

# ----------------------------------------------------------------- 4 Series
oef = o.OEFENINGEN[3]
ws = blad(oef['blad'], 'Turkse series', g.B4_VOETNOOT)
lijst(ws, g.B4_KOPRIJ, g.B4_KOP, g.B4)
vragen(ws, g.B4_KOPRIJ, o.V4, [
    'In welk decennium startten de meeste series?',
    'Hoeveel series zijn dat?',
    f'Hoeveel series blijven over als je filtert op {a.D4_FILTER_GENRE}?'])
plek(ws, oef)
breedtes(ws, [30, 13, 11, 11, 13, 15])
breedtes(ws, [3, 42, 14], start=o.V4 - 1)

# -------------------------------------------------------------- 5 Transfers
oef = o.OEFENINGEN[4]
ws = blad(oef['blad'], 'Transfers van deze zomer', g.B5_VOETNOOT)
lijst(ws, g.B5_KOPRIJ, g.B5_KOP,
      [(speler, positie, g.clubcode(n), bedrag, None, None, None)
       for speler, positie, n, bedrag in g.B5])
# de zoektabel ernaast
k = g.B5_ZOEKKOL
ws.cell(row=g.B5_ZOEKRIJ - 1, column=k,
        value='Clubtabel — niet aanpassen').font = SCHUIN
for i, titel in enumerate(['Code', 'Club', 'Stad', 'Regio'], start=k):
    ws.cell(row=g.B5_ZOEKRIJ, column=i, value=titel).font = VET
for i, (club, stad, regio, *_rest) in enumerate(g.B1):
    r = g.B5_ZOEKRIJ + 1 + i
    ws.cell(row=r, column=k, value=g.clubcode(i + 1))
    ws.cell(row=r, column=k + 1, value=club)
    ws.cell(row=r, column=k + 2, value=stad)
    ws.cell(row=r, column=k + 3, value=regio)
vragen(ws, g.B5_KOPRIJ, o.V5, [
    'Naar welke regio ging het meeste geld?',
    'Hoeveel miljoen was dat?',
    'Wat is het eindtotaal na de extra transfer en na Vernieuwen?'])
plek(ws, oef)
breedtes(ws, [18, 14, 11, 15, 22, 12, 18])
breedtes(ws, [3, 10, 22, 12, 18], start=k - 1)
breedtes(ws, [3, 46, 14], start=o.V5 - 1)

wb.save(UIT)
print('geschreven:', UIT)
for ws in wb.worksheets:
    gevuld = sum(1 for rij in ws.iter_rows() for c in rij if c.value is not None)
    print(f'  {ws.title:<16} {gevuld:>4} gevulde cellen   {ws.dimensions}')
