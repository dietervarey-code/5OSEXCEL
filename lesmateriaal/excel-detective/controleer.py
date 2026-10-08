# -*- coding: utf-8 -*-
"""Controleert of de vijf zaken, het startbestand en de sleutel bij elkaar passen.

Twee controles zijn hier belangrijker dan elders. Eén: elke zaak moet precies
één geneste functie hebben, en alle vijf moeten ze verschillend zijn — anders
is het vijf keer hetzelfde trucje. Twee: die ene formule mag NERGENS letterlijk
in de opdracht staan, want dan valt de uitdaging weg.
"""
import importlib.util
import os
import re
import sys
import zipfile
from collections import Counter
from openpyxl import load_workbook

sys.dont_write_bytecode = True   # geen __pycache__ naast het lesmateriaal

HIER = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HIER, 'detective-startbestand.xlsx')


def laad(naam, bestand):
    spec = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
z = laad('z', 'zaken.py')
sl = laad('sl', 'sleutel.py')

wb = load_workbook(XLSX)
fouten = []

# De tien functies van de basisbundel, plus EN — dat is het enige nieuwe dat
# deze bundel toelaat, en alleen binnen een ALS.
TOEGELATEN = {'SOM', 'AFRONDEN', 'MAX', 'MIN', 'GEMIDDELDE', 'AANTAL',
              'ALS', 'AANTAL.ALS', 'SOM.ALS', 'VERT.ZOEKEN', 'EN'}
FUNCTIE = re.compile(r'\b([A-Z][A-Z.]{1,})\(')


def cellen_van(verwijzing):
    """'E4:E15' -> alle cellen daartussen. 'C17' -> die ene cel."""
    if ':' not in verwijzing:
        return [verwijzing]
    links, rechts = verwijzing.split(':')
    k1, r1 = re.match(r'([A-Z]+)(\d+)', links).groups()
    k2, r2 = re.match(r'([A-Z]+)(\d+)', rechts).groups()
    return [f'{chr(k)}{r}'
            for k in range(ord(k1), ord(k2) + 1)
            for r in range(int(r1), int(r2) + 1)]


def eerste_cel(verwijzing):
    return verwijzing.split(':')[0]


def binnenste(formule):
    """De functies die binnen de haakjes van een ANDERE functie staan.

    Haakjes tellen, niet zomaar zoeken. =MAX(A1:A9)-MIN(B1:B9) heeft twee
    functies maar geen nesting: ze staan naast elkaar, allebei op diepte 0.
    =ALS(EN(...);...) heeft er wel een, want EN staat op diepte 1.
    """
    genest, diepte, i = [], 0, 0
    while i < len(formule):
        m = FUNCTIE.match(formule, i)
        if m:
            if diepte > 0:
                genest.append(m.group(1))
            diepte += 1
            i = m.end()
            continue
        if formule[i] == '(':
            diepte += 1
        elif formule[i] == ')':
            diepte -= 1
        i += 1
    return genest


def opdrachttekst(zaak):
    """Alles wat de leerling bij één zaak te lezen krijgt."""
    stukken = list(zaak['dossier']) + list(zaak['opmaak'])
    stukken += [zaak['proef'], zaak['vraag']]
    stukken += [f'{waar} {wat}' for waar, wat in zaak['gegeven']]
    stukken += [f'{wat} {hoe}' for wat, hoe in zaak['vastgelopen']]
    if zaak.get('grafiek'):
        stukken += [zaak['grafiek']['titel'], zaak['grafiek']['vraag']]
    for m in zaak['maken']:
        stukken += [m['waar'], m['wat'], m.get('formule') or '', m.get('stap') or '',
                    m.get('vorm') or '']
        stukken += m.get('hulp', [])
    return stukken


# 1. Bestaat elk genoemd werkblad?
for zaak in z.ZAKEN:
    if zaak['blad'] not in wb.sheetnames:
        fouten.append(f"werkblad {zaak['blad']} bestaat niet in het startbestand")

# 2. Klopt de opdracht met het startbestand?
for zaak in z.ZAKEN:
    if zaak['blad'] not in wb.sheetnames:
        continue
    ws = wb[zaak['blad']]

    temaken = [c for m in zaak['maken'] for c in cellen_van(m['waar'])]
    bezet = [c for c in temaken if ws[c].value is not None]
    if bezet:
        fouten.append(f"{zaak['blad']}: deze cellen moet de leerling maken, maar ze zijn "
                      f'al gevuld: {bezet[:5]}')

    gegeven = [c for waar, _ in zaak['gegeven'] for c in cellen_van(waar)]
    leeg = [c for c in gegeven if ws[c].value is None]
    if leeg:
        fouten.append(f"{zaak['blad']}: deze cellen zouden gegeven moeten zijn, maar zijn "
                      f'leeg: {leeg[:5]}')
    print(f"zaak {zaak['nr']} ({zaak['blad']}): {len(gegeven)} gegeven, "
          f'{len(temaken)} te maken')

# 3. Precies één geneste functie per zaak, en die moet echt genest zijn.
nestingen = []
for zaak in z.ZAKEN:
    gemarkeerd = [m for m in zaak['maken'] if m.get('genest')]
    if len(gemarkeerd) != 1:
        fouten.append(f"zaak {zaak['nr']}: {len(gemarkeerd)} sleutelformules gemarkeerd, "
                      f'er hoort er precies één te zijn')
        continue
    cel = eerste_cel(gemarkeerd[0]['waar'])
    formule = next((f for c, f, _ in sl.SLEUTEL[zaak['nr']] if c == cel), None)
    if formule is None:
        fouten.append(f"zaak {zaak['nr']}: {cel} staat niet in de sleutel")
        continue
    binnen = binnenste(formule)
    if not binnen:
        fouten.append(f"zaak {zaak['nr']}: {formule} is niet genest — er staat geen "
                      f'functie binnen de haakjes van de buitenste')
    nestingen.append(zaak['nesting'])

    # en geen enkele ANDERE formule van die zaak mag genest zijn
    for c, f, _ in sl.SLEUTEL[zaak['nr']]:
        if f and c != cel and binnenste(f):
            fouten.append(f"zaak {zaak['nr']}: {c} is óók genest ({f}) — er hoort er maar "
                          f'één te zijn')

if len(set(nestingen)) != len(nestingen):
    dubbel = [n for n, k in Counter(nestingen).items() if k > 1]
    fouten.append(f'deze nestingen komen meer dan één keer voor: {dubbel} — vijf keer '
                  f'hetzelfde trucje leert niets')
print(f'\nnestingen: {len(set(nestingen))} verschillende — ' + ' | '.join(nestingen))

# 4. De sleutelformule mag nergens letterlijk in de opdracht staan.
for zaak in z.ZAKEN:
    tekst = ' '.join(opdrachttekst(zaak))
    for m in zaak['maken']:
        if not m.get('genest'):
            continue
        cel = eerste_cel(m['waar'])
        formule = next((f for c, f, _ in sl.SLEUTEL[zaak['nr']] if c == cel), None)
        if formule and formule in tekst:
            fouten.append(f"zaak {zaak['nr']}: de sleutelformule staat letterlijk in de "
                          f'opdracht ({formule}) — dan valt de uitdaging weg')

# 5. Elke basisformule moet wél letterlijk in de opdracht staan: zonder
#    leerkracht erbij mag niemand vastlopen op routinewerk.
for zaak in z.ZAKEN:
    for m in zaak['maken']:
        if m.get('genest') or not m.get('formule'):
            continue
        cel = eerste_cel(m['waar'])
        uit_sleutel = next((f for c, f, _ in sl.SLEUTEL[zaak['nr']] if c == cel), None)
        if uit_sleutel != m['formule']:
            fouten.append(f"zaak {zaak['nr']} {cel}: de opdracht zegt {m['formule']}, de "
                          f'sleutel {uit_sleutel}')

# 6. Sleutel en opdracht moeten over dezelfde cellen gaan.
for zaak in z.ZAKEN:
    uit_sleutel = {c for c, _, _ in sl.SLEUTEL[zaak['nr']]}
    uit_opdracht = {eerste_cel(m['waar']) for m in zaak['maken']}
    if uit_sleutel != uit_opdracht:
        fouten.append(f"zaak {zaak['nr']}: sleutel en opdracht noemen andere cellen: "
                      f'{sorted(uit_sleutel ^ uit_opdracht)}')

# 7. Alleen functies waar ze een fiche voor hebben, plus EN.
for zaak in z.ZAKEN:
    for c, f, _ in sl.SLEUTEL[zaak['nr']]:
        if not f:
            continue
        for fn in FUNCTIE.findall(f):
            if fn not in TOEGELATEN:
                fouten.append(f"zaak {zaak['nr']} {c}: {fn} hoort niet bij de "
                              f'basisfuncties')

# 8. Elke zaak heeft een dossier, een proef, een vraag, een afloop en minstens
#    drie regels 'Loop je vast?'. Zonder leerkracht is dat het vangnet.
for zaak in z.ZAKEN:
    for veld in ('dossier', 'proef', 'vraag', 'afloop'):
        if not zaak.get(veld):
            fouten.append(f"zaak {zaak['nr']}: geen {veld}")
    if len(zaak.get('vastgelopen', [])) < 3:
        fouten.append(f"zaak {zaak['nr']}: minder dan drie regels bij 'Loop je vast?'")

# 9. Staan de brongegevens waar de opdracht beweert?
controles = [
    ('1 Verdwenen avonden', 'B2', g.Z1_GRENS),
    ('1 Verdwenen avonden', f'B{g.Z1_START}', g.Z1[0][1]),
    ('1 Verdwenen avonden', f'E{g.Z1_EIND}', g.Z1[-1][4]),
    ('1 Verdwenen avonden', f'K{g.Z1_TAB_EIND}', g.Z1_HANDELAARS[-1][1]),
    ('2 Mercedes', f'A{g.Z2_START}', g.Z2[0][0]),
    ('2 Mercedes', f'G{g.Z2_EIND}', g.Z2[-1][6]),
    ('2 Mercedes', f'A{g.Z2_PER_CHAUFFEUR}', g.Z2_CHAUFFEURS[0]),
    ('3 Museumroof', f'A{g.Z3_START}', g.Z3[0][0]),
    ('3 Museumroof', f'E{g.Z3_EIND}', g.Z3[-1][3]),
    ('3 Museumroof', f'H{g.Z3_TEL_EIND}', g.Z3_TELLING[-1]),
    ('3 Museumroof', f'K{g.Z3_ZAAL_EIND}', g.Z3_ZALEN[-1][1]),
    ('4 Gelekt examen', 'B2', g.Z4_GRENS),
    ('4 Gelekt examen', f'B{g.Z4_START}', g.Z4[0][1]),
    ('4 Gelekt examen', f'E{g.Z4_EIND}', g.Z4[-1][4]),
    ('4 Gelekt examen', f'J{g.Z4_TAB_EIND}', g.Z4_GEBRUIKERS[-1][2]),
    ('5 Chocolade', 'B2', g.Z5_GRENS),
    ('5 Chocolade', f'B{g.Z5_START}', g.Z5[0][1]),
    ('5 Chocolade', f'E{g.Z5_EIND}', g.Z5[-1][4]),
    ('5 Chocolade', f'L{g.Z5_TAB_EIND}', g.Z5_MACHINES[-1][2]),
]
for blad, cel, verwacht in controles:
    echt = wb[blad][cel].value
    if echt != verwacht:
        fouten.append(f'{blad}!{cel}: verwacht {verwacht!r}, gevonden {echt!r}')

# 10. De twee grafieken moeten bestaan en naar de juiste kolommen wijzen.
xml = ''.join(zipfile.ZipFile(XLSX).read(n).decode()
              for n in zipfile.ZipFile(XLSX).namelist() if n.startswith('xl/charts/chart'))
for blad, kolom, eerste, laatste in [
    ('2 Mercedes', 'F', g.Z2_START, g.Z2_EIND),
    ('5 Chocolade', 'E', g.Z5_START, g.Z5_EIND),
]:
    verwijzing = f"'{blad}'!${kolom}${eerste}:${kolom}${laatste}"
    if verwijzing not in xml:
        fouten.append(f'de grafiek van {blad} verwijst niet naar {verwijzing}')
grafieken = xml.count('<barChart>')
if grafieken != 2:
    fouten.append(f'{grafieken} grafieken gevonden in plaats van 2')
else:
    print('grafieken: 2, allebei naar de juiste kolom')

# 11. Zaak 1 — de proef moet sluiten en de dadercode mag nergens anders staan.
if a.Z1_BEDRAG_LEUGEN != a.Z1_BEDRAG_DADER:
    fouten.append(f'zaak 1: de proef sluit niet — {a.Z1_BEDRAG_LEUGEN} tegenover '
                  f'{a.Z1_BEDRAG_DADER}. De dadercode mag alleen op leugenavonden staan.')
anders = [datum for (datum, zei, minuten, _, code) in g.Z1
          if code == g.Z1_DADERCODE and not (zei == 'overwerk' and minuten < g.Z1_GRENS)]
if anders:
    fouten.append(f'zaak 1: code {g.Z1_DADERCODE} staat ook op avonden zonder leugen '
                  f'({anders}) — dan sluit de proef niet meer')
else:
    print(f'zaak 1: {a.Z1_AANTAL_LEUGEN} leugenavonden, proef sluit op '
          f'{a.Z1_BEDRAG_LEUGEN:.2f}')

# 12. Zaak 2 — de proef moet sluiten, de hoogste stand moet uniek zijn, en de
#     langste rit mag NIET de laatste rit zijn (anders is de grafiek geen test).
if a.Z2_TOTAAL != a.Z2_PROEF:
    fouten.append(f'zaak 2: de proef sluit niet — {a.Z2_TOTAAL} tegenover {a.Z2_PROEF}')
standen = [aan for _, _, _, _, _, aan, _ in g.Z2]
if standen.count(a.Z2_HOOGSTE_STAND) != 1:
    fouten.append('zaak 2: de hoogste kilometerstand komt meer dan één keer voor')
if a.Z2_AFSTAND.count(a.Z2_LANGSTE) != 1:
    fouten.append('zaak 2: er zijn twee even lange ritten — dan is MAX niet eenduidig')
if a.Z2_LANGSTE_RIT[0] == a.Z2_LAATSTE[0]:
    fouten.append('zaak 2: de langste rit is ook de laatste rit — dan leert de grafiek '
                  'niets en is de sleutelformule overbodig')
else:
    print(f'zaak 2: langste rit {a.Z2_LANGSTE_RIT[0]} ({a.Z2_LANGSTE_RIT[6]}), laatste rit '
          f'{a.Z2_LAATSTE[0]} ({a.Z2_BESTEMMING}) — twee verschillende, zoals het hoort')

# 13. Zaak 3 — het aantal weg moet kloppen met de telling, en geen nummer mag
#     twee keer in de telling staan.
verwacht_weg = len(g.Z3) - len(g.Z3_TELLING)
if a.Z3_AANTAL_GESTOLEN != verwacht_weg:
    fouten.append(f'zaak 3: {a.Z3_AANTAL_GESTOLEN} stukken weg, maar de telling laat er '
                  f'{verwacht_weg} verwachten')
dubbel = [n for n, k in Counter(g.Z3_TELLING).items() if k > 1]
if dubbel:
    fouten.append(f'zaak 3: deze nummers staan twee keer in de telling: {dubbel}')
onbekend = sorted(set(g.Z3_TELLING) - {nr for nr, _, _, _ in g.Z3})
if onbekend:
    fouten.append(f'zaak 3: de telling bevat nummers die niet in de catalogus staan: '
                  f'{onbekend}')
if len(a.Z3_ZALEN_WEG) != 1:
    fouten.append(f'zaak 3: de gestolen stukken komen uit {len(a.Z3_ZALEN_WEG)} zalen '
                  f'({a.Z3_ZALEN_WEG}) — de clou is net dat het er één is')
else:
    print(f'zaak 3: {a.Z3_AANTAL_GESTOLEN} stukken weg, allemaal uit de '
          f'{a.Z3_ZALEN_WEG[0]}')

# 14. Zaak 4 — precies één verdachte, en de tweede moet er duidelijk onder zitten.
if a.Z4_AANTAL_VERDACHT != 1:
    fouten.append(f'zaak 4: {a.Z4_AANTAL_VERDACHT} verdachten in plaats van 1')
elif a.Z4_NIPT[4] >= g.Z4_GRENS:
    fouten.append(f'zaak 4: {a.Z4_NIPT[1]} zit met {a.Z4_NIPT[4]} bestanden niet onder de '
                  f'grens van {g.Z4_GRENS}')
else:
    print(f'zaak 4: één verdachte ({a.Z4_DADER[1]}, {a.Z4_DADER[4]}), de tweede zit op '
          f'{a.Z4_NIPT[4]} — grens {g.Z4_GRENS}')

# 15. Zaak 5 — alle slechte batches bij één technicus, en genoeg marge rond de
#     grens zodat een afrondingsverschil nooit het antwoord kan omgooien.
technici = {t for _, _, _, _, t in a.Z5_SLECHT}
if len(technici) != 1:
    fouten.append(f'zaak 5: de slechte batches komen bij {len(technici)} technici uit '
                  f'({technici}) — de clou is net dat het er één is')
if a.Z5_AANTAL_DADER != len(a.Z5_SLECHT):
    fouten.append(f'zaak 5: {a.Z5_AANTAL_DADER} batches op naam van {a.Z5_DADER}, maar '
                  f'{len(a.Z5_SLECHT)} slechte batches')
dichtbij = [(batch, pct) for (batch, *_), pct in zip(g.Z5, a.Z5_AFKEUR)
            if abs(pct - g.Z5_GRENS) < 0.02]
if dichtbij:
    fouten.append(f'zaak 5: deze batches liggen te dicht bij de grens {g.Z5_GRENS}: '
                  f'{dichtbij} — een afronding kan het antwoord dan omgooien')
else:
    marge = min(abs(pct - g.Z5_GRENS) for pct in a.Z5_AFKEUR)
    print(f'zaak 5: {len(a.Z5_SLECHT)} slechte batches, allemaal {a.Z5_DADER}; '
          f'kleinste marge tot de grens {marge:.3f}')

print()
if fouten:
    print('PROBLEMEN:')
    for p in fouten:
        print('  -', p)
    raise SystemExit(1)
print('Alles sluit: de vijf zaken, het startbestand en de sleutel passen bij elkaar.')
print('Elke zaak heeft precies één geneste functie, alle vijf verschillend, en geen')
print('enkele staat letterlijk in de opdracht. Elke basisformule staat er juist wél in.')
