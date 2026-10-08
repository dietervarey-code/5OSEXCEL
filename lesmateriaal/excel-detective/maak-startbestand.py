# -*- coding: utf-8 -*-
"""Bouwt het startbestand met de vijf zaken. Bewust kaal: de opmaak hoort bij
   de opdracht. Twee zaken krijgen wel een kant-en-klare grafiek mee — die
   moeten ze lezen, niet maken."""
import importlib.util
import os
import sys
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

sys.dont_write_bytecode = True   # geen __pycache__ naast het lesmateriaal

HIER = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location('g', os.path.join(HIER, 'gegevens.py'))
g = importlib.util.module_from_spec(spec); spec.loader.exec_module(g)

UIT = os.path.join(HIER, 'detective-startbestand.xlsx')
TITELFONT = Font(bold=True, size=13)
VET = Font(bold=True)

wb = Workbook()


def breedtes(ws, waarden):
    for i, b in enumerate(waarden, start=1):
        ws.column_dimensions[get_column_letter(i)].width = b


def koprij(ws, rij, koppen, start_kolom=1):
    for i, k in enumerate(koppen, start=start_kolom):
        ws.cell(row=rij, column=i, value=k)


def labels(ws, start_rij, teksten, kolom=1):
    for i, t in enumerate(teksten):
        ws.cell(row=start_rij + i, column=kolom, value=t)


def blad(naam, titel):
    ws = wb.create_sheet(naam)
    ws['A1'] = titel
    ws['A1'].font = TITELFONT
    return ws


def staafgrafiek(ws, titel, waardekolom, koprij_nr, eerste, laatste, anker,
                 as_min=None, y_titel=None):
    """Een kolomgrafiek over één kolom die al ingevuld is."""
    grafiek = BarChart()
    grafiek.type = 'col'
    grafiek.title = titel
    grafiek.legend = None
    grafiek.y_axis.title = y_titel
    data = Reference(ws, min_col=waardekolom, min_row=koprij_nr, max_row=laatste)
    cats = Reference(ws, min_col=1, min_row=eerste, max_row=laatste)
    grafiek.add_data(data, titles_from_data=True)
    grafiek.set_categories(cats)
    # openpyxl laat deze twee standaard leeg; zonder deze regels verbergt Excel
    # soms de assen van een grafiek die niet in Excel zelf gemaakt is.
    grafiek.x_axis.delete = False
    grafiek.y_axis.delete = False
    grafiek.x_axis.axPos = 'b'
    grafiek.y_axis.axPos = 'l'
    if as_min is not None:
        grafiek.y_axis.scaling.min = as_min
    grafiek.height = 8
    grafiek.width = 17
    ws.add_chart(grafiek, anker)
    return grafiek


# ------------------------------------------------------------------ Start
ws = wb.active
ws.title = 'Start'
ws['A1'] = g.TITEL
ws['A1'].font = Font(bold=True, size=16)
for i, r in enumerate([
    g.ONDERTITEL,
    '',
    'Naam: ......................................    Klas: ..........',
    '',
    'Vijf zaken. Elke zaak staat op zichzelf en duurt ongeveer 25 minuten.',
    'Je hoeft ze niet in volgorde te doen, maar beginnen bij zaak 1 is het',
    'gemakkelijkst.',
    '',
    'Zaak 1  De verdwenen avonden              liegt hij over zijn overuren?',
    'Zaak 2  De Mercedes van de Koning         waar staat de wagen nu?',
    'Zaak 3  De museumroof                     wat is er weg, en voor hoeveel?',
    'Zaak 4  Het gelekte examen                wie zette het online?',
    'Zaak 5  De sabotage in de chocoladefabriek wie zat aan de machines?',
    '',
    'Hoe elke zaak in elkaar zit',
    '  Zes of zeven gewone formules zetten het dossier op orde. Die staan',
    '  letterlijk in de opdracht: typ ze over en voer ze door.',
    '  Daarna is er één SLEUTELFORMULE met een functie binnen een functie.',
    '  Daarvan krijg je alleen de vorm. Die moet je zelf afmaken: dat is wat',
    '  de zaak kraakt.',
    '',
    'Hoe weet je of het klopt?',
    '  Bij elke zaak staat een PROEF: twee getallen die hetzelfde moeten geven,',
    '  of een getal dat je op voorhand kent. Komt het uit, dan heb je de zaak.',
    '  Komt het niet uit, lees dan het lijstje "Loop je vast?" bij die zaak.',
    '',
    'EERST DOEN: sla dit bestand op als  detective-jouwnaam.xlsx',
    '',
    g.WAARSCHUWING,
], start=2):
    ws.cell(row=i, column=1, value=r)
ws['A16'].font = VET
ws['A23'].font = VET
ws['A28'].font = VET
ws['A30'].font = Font(italic=True)
breedtes(ws, [86])

# ============================================= ZAAK 1 — Verdwenen avonden
ws = blad('1 Verdwenen avonden', 'Dossier Deleu — twintig avonden')
ws['A2'] = 'Geen overwerk onder (min)'
ws['B2'] = g.Z1_GRENS
koprij(ws, g.Z1_KOPRIJ, g.Z1_KOP)
for r, (datum, zei, minuten, bedrag, code) in enumerate(g.Z1, start=g.Z1_START):
    ws.cell(row=r, column=1, value=datum)
    ws.cell(row=r, column=2, value=zei)
    ws.cell(row=r, column=3, value=minuten)
    ws.cell(row=r, column=4, value=bedrag)
    ws.cell(row=r, column=5, value=code)
labels(ws, g.Z1_ANTWOORD, [
    'Aantal leugenavonden',
    'Bedrag op die avonden',
    'Code van de handelaar',
    'Bedrag bij die code',
    'Totaal uitgegeven',
    'Grootste bedrag',
])
ws['J3'] = 'Handelaars — niet aanpassen'
ws['J3'].font = VET
koprij(ws, g.Z1_KOPRIJ, g.Z1_ZOEKKOP, start_kolom=10)
for r, (code, naam) in enumerate(g.Z1_HANDELAARS, start=g.Z1_TAB_START):
    ws.cell(row=r, column=10, value=code)
    ws.cell(row=r, column=11, value=naam)
breedtes(ws, [12, 17, 13, 11, 8, 24, 13, 3, 3, 8, 24])

# ================================================== ZAAK 2 — De Mercedes
ws = blad('2 Mercedes', 'Logboek van de koninklijke Mercedes')
koprij(ws, g.Z2_KOPRIJ, g.Z2_KOP)
for r, (rit, datum, chauffeur, vertrek, kv, ka, best) in enumerate(g.Z2,
                                                                   start=g.Z2_START):
    ws.cell(row=r, column=1, value=rit)
    ws.cell(row=r, column=2, value=datum)
    ws.cell(row=r, column=3, value=chauffeur)
    ws.cell(row=r, column=4, value=vertrek)
    ws.cell(row=r, column=5, value=kv)
    ws.cell(row=r, column=6, value=ka)
    ws.cell(row=r, column=7, value=best)
koprij(ws, g.Z2_BLOK, ['Chauffeur', 'Ritten', 'Kilometers'])
for i, chauffeur in enumerate(g.Z2_CHAUFFEURS):
    ws.cell(row=g.Z2_PER_CHAUFFEUR + i, column=1, value=chauffeur)
labels(ws, g.Z2_ANTWOORD, [
    'Totaal gereden',
    'Hoogste min laagste',
    'Langste rit',
    'Gemiddelde rit',
    'Laatst gezien in',
])
breedtes(ws, [12, 10, 12, 12, 12, 13, 19, 10])
staafgrafiek(ws, 'Kilometerstand bij aankomst', waardekolom=6,
             koprij_nr=g.Z2_KOPRIJ, eerste=g.Z2_START, laatste=g.Z2_EIND,
             anker='J3', as_min=84000, y_titel='km-stand')

# ================================================== ZAAK 3 — De museumroof
ws = blad('3 Museumroof', 'Catalogus Stedelijk Museum')
koprij(ws, g.Z3_KOPRIJ, g.Z3_KOP)
for r, (nr, stuk, zaalcode, waarde) in enumerate(g.Z3, start=g.Z3_START):
    ws.cell(row=r, column=1, value=nr)
    ws.cell(row=r, column=2, value=stuk)
    ws.cell(row=r, column=3, value=zaalcode)
    ws.cell(row=r, column=5, value=waarde)
labels(ws, g.Z3_ANTWOORD, [
    'Aantal gestolen',
    'Waarde van de buit',
    'Waarde van de collectie',
    'Duurste stuk',
    'Gemiddelde waarde',
])
ws['H2'] = 'Teruggevonden — niet aanpassen'
ws['H2'].font = VET
ws.cell(row=g.Z3_KOPRIJ, column=8, value='Catalogusnr')
for r, nr in enumerate(g.Z3_TELLING, start=g.Z3_TEL_START):
    ws.cell(row=r, column=8, value=nr)
ws['J2'] = 'Zalen — niet aanpassen'
ws['J2'].font = VET
koprij(ws, g.Z3_KOPRIJ, g.Z3_ZAALKOP, start_kolom=10)
for r, (code, naam) in enumerate(g.Z3_ZALEN, start=g.Z3_ZAAL_START):
    ws.cell(row=r, column=10, value=code)
    ws.cell(row=r, column=11, value=naam)
breedtes(ws, [13, 34, 10, 20, 12, 12, 3, 13, 3, 7, 20])

# =============================================== ZAAK 4 — Gelekt examen
ws = blad('4 Gelekt examen', 'Logbestand van de map Examens')
ws['A2'] = 'Verdacht vanaf (bestanden)'
ws['B2'] = g.Z4_GRENS
koprij(ws, g.Z4_KOPRIJ, g.Z4_KOP)
for r, (log, code, dag, tijd, bestanden) in enumerate(g.Z4, start=g.Z4_START):
    ws.cell(row=r, column=1, value=log)
    ws.cell(row=r, column=2, value=code)
    ws.cell(row=r, column=3, value=dag)
    ws.cell(row=r, column=4, value=tijd)
    ws.cell(row=r, column=5, value=bestanden)
koprij(ws, g.Z4_BLOK, ['Code', 'Naam', 'Functie', 'Logins', 'Verdacht?'])
for i, (code, _, _) in enumerate(g.Z4_GEBRUIKERS):
    ws.cell(row=g.Z4_PER_GEBRUIKER + i, column=1, value=code)
labels(ws, g.Z4_ANTWOORD, [
    'Totaal bestanden',
    'Grootste in één keer',
    'Hoeveel verdachten',
])
ws['H3'] = 'Gebruikers — niet aanpassen'
ws['H3'].font = VET
koprij(ws, g.Z4_KOPRIJ, g.Z4_GEBRUIKERKOP, start_kolom=8)
for r, (code, naam, functie) in enumerate(g.Z4_GEBRUIKERS, start=g.Z4_TAB_START):
    ws.cell(row=r, column=8, value=code)
    ws.cell(row=r, column=9, value=naam)
    ws.cell(row=r, column=10, value=functie)
breedtes(ws, [11, 20, 12, 11, 12, 3, 3, 8, 20, 22])

# ================================================== ZAAK 5 — De sabotage
ws = blad('5 Chocolade', 'Productiecijfers Praliné Dumont')
ws['A2'] = 'Afkeur vanaf'
ws['B2'] = g.Z5_GRENS
koprij(ws, g.Z5_KOPRIJ, g.Z5_KOP)
for r, (batch, machine, ploeg, prod, afgekeurd) in enumerate(g.Z5, start=g.Z5_START):
    ws.cell(row=r, column=1, value=batch)
    ws.cell(row=r, column=2, value=machine)
    ws.cell(row=r, column=3, value=ploeg)
    ws.cell(row=r, column=4, value=prod)
    ws.cell(row=r, column=5, value=afgekeurd)
labels(ws, g.Z5_ANTWOORD, [
    'Totaal geproduceerd',
    'Totaal afgekeurd',
    'Hoogste afkeur',
    'Gemiddelde afkeur',
    'Naam die je vond',
    'Batches van die technicus',
])
ws['J3'] = 'Onderhoudsregister — niet aanpassen'
ws['J3'].font = VET
koprij(ws, g.Z5_KOPRIJ, g.Z5_MACHINEKOP, start_kolom=10)
for r, (code, naam, technicus) in enumerate(g.Z5_MACHINES, start=g.Z5_TAB_START):
    ws.cell(row=r, column=10, value=code)
    ws.cell(row=r, column=11, value=naam)
    ws.cell(row=r, column=12, value=technicus)
breedtes(ws, [10, 11, 9, 14, 12, 11, 22, 3, 3, 8, 22, 22])
staafgrafiek(ws, 'Afgekeurde pralines per batch', waardekolom=5,
             koprij_nr=g.Z5_KOPRIJ, eerste=g.Z5_START, laatste=g.Z5_EIND,
             anker='N3', y_titel='afgekeurd')

wb.save(UIT)
print('geschreven:', UIT)
print('bladen:', wb.sheetnames)
print(f'  zaak 1  rij {g.Z1_START}-{g.Z1_EIND}, antwoorden {g.Z1_ANTWOORD}-'
      f'{g.Z1_ANTWOORD + 5}, handelaars rij {g.Z1_TAB_START}-{g.Z1_TAB_EIND}')
print(f'  zaak 2  rij {g.Z2_START}-{g.Z2_EIND}, chauffeurs {g.Z2_PER_CHAUFFEUR}-'
      f'{g.Z2_PER_CHAUFFEUR + 3}, antwoorden {g.Z2_ANTWOORD}-{g.Z2_ANTWOORD + 4} '
      f'+ grafiek')
print(f'  zaak 3  rij {g.Z3_START}-{g.Z3_EIND}, telling rij {g.Z3_TEL_START}-'
      f'{g.Z3_TEL_EIND}, antwoorden {g.Z3_ANTWOORD}-{g.Z3_ANTWOORD + 4}')
print(f'  zaak 4  rij {g.Z4_START}-{g.Z4_EIND}, gebruikers {g.Z4_PER_GEBRUIKER}-'
      f'{g.Z4_PER_GEBRUIKER + 7}, antwoorden {g.Z4_ANTWOORD}-{g.Z4_ANTWOORD + 2}')
print(f'  zaak 5  rij {g.Z5_START}-{g.Z5_EIND}, antwoorden {g.Z5_ANTWOORD}-'
      f'{g.Z5_ANTWOORD + 5} + grafiek')
