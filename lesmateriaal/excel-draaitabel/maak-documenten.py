# -*- coding: utf-8 -*-
"""Bouwt de theoriebundel over draaitabellen en grafieken.

Eén document, veertien kaarten, elk op een eigen pagina. De beelden worden
hier getekend — er wordt niets uit een map ingelezen dat er al stond.
"""
import importlib.util
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


t = laad('t', 'tekening.py')
f = laad('f', 'fiches.py')

UIT = os.path.join(HIER, 'draaitabellen-en-grafieken.docx')
BEELDMAP = os.path.join(HIER, 'beelden')

BLAUW = RGBColor(0x1F, 0x38, 0x64)
ORANJE = RGBColor(0xB0, 0x4A, 0x00)
GRIJS = RGBColor(0x59, 0x59, 0x59)
GROENTEKST = RGBColor(0x38, 0x76, 0x1D)

KLEUREN = {'ROOD': t.ROOD, 'GROEN': t.GROEN}
MENU = re.compile(r'((?:[A-Z][a-zA-Z]+(?: [a-z]+)*)(?: → [A-Z][a-zA-Z ]+)+)')


def opzet(doc):
    s = doc.styles['Normal']
    s.font.name = 'Calibri'
    s.font.size = Pt(11)
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(1.6)
        sec.left_margin = sec.right_margin = Cm(2.0)


def kop(doc, tekst, niveau=1, kleur=BLAUW):
    p = doc.add_heading(tekst, level=niveau)
    for r in p.runs:
        r.font.color.rgb = kleur
    return p


def alinea(doc, tekst, style=None, vet=False, schuin=False, grootte=None,
           kleur=None, na=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(tekst)
    r.bold = vet
    r.italic = schuin
    if grootte:
        r.font.size = Pt(grootte)
    if kleur:
        r.font.color.rgb = kleur
    if na is not None:
        p.paragraph_format.space_after = Pt(na)
    return p


def kader(doc, titel, regels, kleur=BLAUW):
    tb = doc.add_table(rows=1, cols=1)
    tb.style = 'Table Grid'
    tb.alignment = WD_TABLE_ALIGNMENT.LEFT
    c = tb.rows[0].cells[0]
    p = c.paragraphs[0]
    r = p.add_run(titel)
    r.bold = True
    r.font.color.rgb = kleur
    for regel in regels:
        pp = c.add_paragraph()
        pp.paragraph_format.space_after = Pt(2)
        pp.add_run(regel)
    doc.add_paragraph()
    return tb


def paginaeinde(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# --------------------------------------------------------------- de beelden
def tekenbeeld(spec):
    """Tekent één beeld en geeft het pad terug. Het soort bepaalt welke
       functie uit tekening.py gebruikt wordt."""
    pad = os.path.join(BEELDMAP, spec['bestand'] + '.png')
    soort = spec['soort']
    if soort == 'rooster':
        kaders = [(b, KLEUREN[k]) for b, k in spec.get('kaders', ())]
        return t.rooster(pad, kolommen=spec['kolommen'], rijen=spec['rijen'],
                         formule=spec.get('formule'),
                         geselecteerd=spec.get('geselecteerd'),
                         gemarkeerd=spec.get('gemarkeerd', ()),
                         rechts=spec.get('rechts', ()),
                         vet_rij=spec.get('vet_rij'), kaders=kaders)
    if soort == 'paneel':
        return t.veldenpaneel(pad, spec['velden'],
                              filters=spec.get('filters', ()),
                              kolommen=spec.get('kolommen', ()),
                              rijen=spec.get('rijen', ()),
                              waarden=spec.get('waarden', ()),
                              aangevinkt=spec.get('aangevinkt', ()))
    if soort == 'dialoog':
        return t.dialoog(pad, spec['titel'], spec['regels'], spec['keuzes'],
                         gekozen=spec.get('gekozen'))
    if soort == 'staaf':
        return t.staafdiagram(pad, spec['titel'], spec['labels'], spec['waarden'],
                              astitel_x=spec.get('astitel_x'),
                              astitel_y=spec.get('astitel_y'),
                              waardelabels=spec.get('waardelabels', True),
                              markeer=spec.get('markeer'),
                              laag=spec.get('laag', 0))
    if soort == 'lijn':
        return t.lijndiagram(pad, spec['titel'], spec['labels'], spec['waarden'],
                             astitel_x=spec.get('astitel_x'),
                             astitel_y=spec.get('astitel_y'),
                             laag=spec.get('laag', 0))
    if soort == 'cirkel':
        return t.cirkeldiagram(pad, spec['titel'], spec['labels'], spec['waarden'])
    raise ValueError(soort)


BREEDTE = {'rooster': 4.5, 'paneel': 2.3, 'dialoog': 2.5,
           'staaf': 4.3, 'lijn': 4.3, 'cirkel': 3.4}


def beeld(doc, spec):
    pad = tekenbeeld(spec)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(pad, width=Inches(BREEDTE[spec['soort']]))
    po = doc.add_paragraph()
    po.paragraph_format.space_after = Pt(10)
    po.paragraph_format.left_indent = Cm(0.6)
    r = po.add_run(spec['onderschrift'])
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = GRIJS
    return pad


# ----------------------------------------------------------------- de kaart
def kaart(doc, nummer, k):
    kop(doc, f'{nummer}. {k["naam"]}', 1)
    alinea(doc, k['wat'], vet=True, na=6)

    if k.get('waar'):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run('Waar vind je het?  ')
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = GRIJS
        r = p.add_run(k['waar'])
        r.font.name = 'Consolas'
        r.font.size = Pt(10)
        r.font.color.rgb = ORANJE

    if k.get('stappen'):
        kop(doc, 'Zo doe je het', 2, GRIJS)
        for i, stap in enumerate(k['stappen'], 1):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(f'{i}.  ')
            r.bold = True
            r.font.color.rgb = BLAUW
            p.add_run(stap)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    if k.get('uitleg'):
        kop(doc, 'Wat er gebeurt', 2, GRIJS)
        for regel in k['uitleg']:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(4)
            p.add_run(regel)

    for spec in k.get('beelden', ()):
        beeld(doc, spec)

    if k.get('weten'):
        kader(doc, 'Goed om te weten', k['weten'], GROENTEKST)

    if k.get('fout'):
        kader(doc, 'De fout die iedereen één keer maakt', [k['fout']], ORANJE)


# ------------------------------------------------------------------ bouwen
doc = Document()
opzet(doc)

kop(doc, 'Draaitabellen en grafieken', 0)
alinea(doc, 'Theoriebundel — veertien kaarten', schuin=True, kleur=GRIJS, na=14)

alinea(doc,
       'Tot nu toe liet je Excel rekenen met formules: je typte zelf welke cellen '
       'het moest optellen en onder welke voorwaarde. Dat werkt, maar bij een lange '
       'lijst met veel categorieën heb je er tientallen nodig.',
       na=6)
alinea(doc,
       'Een draaitabel doet hetzelfde zonder één formule: je sleept kolomtitels in '
       'vakjes. Een grafiek maakt er daarna een beeld van. Samen zijn dat de twee '
       'gereedschappen waarmee je een lijst van honderden regels in één blad '
       'samenvat.',
       na=12)

kader(doc, 'Hoe je deze bundel gebruikt', [
    'Kaart 1 tot 9 gaan over de draaitabel, kaart 10 tot 14 over de grafiek.',
    'Elke kaart staat op zichzelf. Loop je vast bij een opdracht, pak dan die ene '
    'kaart erbij — je hoeft de rest niet gelezen te hebben.',
    'Op elke kaart staat onderaan "de fout die iedereen één keer maakt". Lees die '
    'eerst als er iets niet lukt; de kans is groot dat het die is.',
    'De voorbeelden spelen in de schoolbibliotheek, de kantine en het schoolarchief. '
    'Ze geven nooit de oplossing van een opdracht weg.',
], BLAUW)

kop(doc, 'Inhoud', 2, GRIJS)
tb = doc.add_table(rows=1, cols=3)
tb.style = 'Light Grid Accent 1'
for i, k in enumerate(['Nr', 'Kaart', 'Waarvoor']):
    tb.rows[0].cells[i].paragraphs[0].add_run(k).bold = True
for i, k in enumerate(f.KAARTEN, 1):
    c = tb.add_row().cells
    c[0].paragraphs[0].add_run(str(i))
    c[1].paragraphs[0].add_run(k['naam'])
    c[2].paragraphs[0].add_run(k['wat'].split('.')[0] + '.')

huidig_deel = None
for nummer, k in enumerate(f.KAARTEN, 1):
    paginaeinde(doc)
    if k['deel'] != huidig_deel:
        huidig_deel = k['deel']
        kop(doc, huidig_deel, 1, ORANJE)
    kaart(doc, nummer, k)

paginaeinde(doc)
kop(doc, 'Spiekblad', 1)
alinea(doc, 'Alles van deze bundel op één pagina.', schuin=True, kleur=GRIJS, na=10)

kop(doc, 'Een draaitabel in vijf stappen', 2, GRIJS)
for i, stap in enumerate([
    'Lijst nakijken: één koprij, geen lege rij of kolom, geen totaalrij eronder.',
    'Eén cel in de lijst aanklikken. Invoegen, Draaitabel.',
    'Veld naar Rijen (waarover wil je het weten?), veld naar Waarden (wat wil je '
    'weten?).',
    'Kopje lezen: staat er Som of Aantal? Zo nodig aanpassen via '
    'Waardeveldinstellingen.',
    'Sorteren van groot naar klein. Daarna pas een grafiek maken.',
], 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f'{i}.  '); r.bold = True; r.font.color.rgb = BLAUW
    p.add_run(stap)

doc.add_paragraph()
kop(doc, 'De vier vakken', 2, GRIJS)
tb = doc.add_table(rows=1, cols=2)
tb.style = 'Light Grid Accent 1'
for i, k in enumerate(['Vak', 'Wat je erin legt']):
    tb.rows[0].cells[i].paragraphs[0].add_run(k).bold = True
for naam, uitleg in [
    ('Rijen', 'het veld waarover je het wil weten — één regel per waarde'),
    ('Kolommen', 'een tweede veld, als je een kruistabel wil — één kolom per waarde'),
    ('Waarden', 'het veld waarmee gerekend wordt — Som, Aantal, Gemiddelde, Max, Min'),
    ('Filters', 'een veld waarmee je boven de tabel kan kiezen wat je ziet'),
]:
    c = tb.add_row().cells
    c[0].paragraphs[0].add_run(naam).bold = True
    c[1].paragraphs[0].add_run(uitleg)

doc.add_paragraph()
kop(doc, 'Welk grafiektype?', 2, GRIJS)
tb = doc.add_table(rows=1, cols=3)
tb.style = 'Light Grid Accent 1'
for i, k in enumerate(['Type', 'Vraag', 'Let op']):
    tb.rows[0].cells[i].paragraphs[0].add_run(k).bold = True
for rij in [
    ('Staafdiagram', 'wie is de grootste?', 'as vanaf nul, en sorteer op grootte'),
    ('Lijndiagram', 'hoe evolueert het?', 'alleen voor een verloop in de tijd'),
    ('Cirkeldiagram', 'welk aandeel van het geheel?', 'hoogstens vier of vijf delen'),
]:
    c = tb.add_row().cells
    for i, w in enumerate(rij):
        c[i].paragraphs[0].add_run(w)

doc.add_paragraph()
kader(doc, 'De drie dingen die het vaakst mislopen', [
    '1.  Je draaitabel telt in plaats van op te tellen. Kijk naar het kopje: staat er '
    'Aantal waar Som hoort? Komt je eindtotaal overeen met het aantal regels van je '
    'lijst, dan is het dat.',
    '2.  Je draaitabel is niet vernieuwd. Ze rekent niet automatisch mee. Rechtsklik, '
    'Vernieuwen — doe het uit gewoonte.',
    '3.  Je grafiek heeft geen titel of geen astitel. Dan is het een plaatje en geen '
    'antwoord.',
], ORANJE)

doc.save(UIT)
print('geschreven:', UIT)
print(f'  {len(f.KAARTEN)} kaarten, '
      f'{sum(len(k["beelden"]) for k in f.KAARTEN)} beelden')
