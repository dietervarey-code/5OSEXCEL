# -*- coding: utf-8 -*-
import importlib.util
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

def stip(t):
    """'23:55' -> een echte datum-tijd. Alles voor 12 uur hoort bij 15 maart:
    die nacht loopt over middernacht, en als tekst zou 00:04 vóór 18:04
    gesorteerd worden."""
    u, mi = (int(x) for x in t.split(':'))
    return datetime(2026, 3, 15 if u < 12 else 14, u, mi)

STIPOPMAAK = 'dd/mm hh:mm'

spec = importlib.util.spec_from_file_location('m','/tmp/ctrl/moordzaak-data.py')
m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)

KOP   = PatternFill('solid', fgColor='1F3864')
KOPTK = Font(bold=True, color='FFFFFF', size=11)
TITEL = Font(bold=True, size=14, color='1F3864')
RAND  = Border(bottom=Side('thin', color='B4C6E7'))

wb = Workbook()

def blad(naam, koppen, rijen, breedtes, titel=None, bevriezen='A4'):
    ws = wb.create_sheet(naam)
    ws['A1'] = titel or naam
    ws['A1'].font = TITEL
    start = 3
    for i, k in enumerate(koppen, start=1):
        c = ws.cell(row=start, column=i, value=k)
        c.fill, c.font = KOP, KOPTK
        c.alignment = Alignment(horizontal='center', vertical='center')
    for r, rij in enumerate(rijen, start=start + 1):
        for i, w in enumerate(rij, start=1):
            c = ws.cell(row=r, column=i, value=w)
            c.border = RAND
    for i, b in enumerate(breedtes, start=1):
        ws.column_dimensions[get_column_letter(i)].width = b
    ws.freeze_panes = bevriezen
    ws.row_dimensions[start].height = 22
    return ws

# --- Briefing ---------------------------------------------------------
ws = wb.active; ws.title = 'Zaakdossier'
ws['A1'] = 'ZAAK 2026-0315 — Hotel Astrid, Oostende'; ws['A1'].font = Font(bold=True, size=16, color='1F3864')
regels = [
    '',
    'Nacht van 14 op 15 maart 2026.',
    '',
    f'Om 06:40 vindt kamermeisje Sofie Groen de deur van kamer {m.KAMER_SLACHTOFFER} op een kier.',
    f'In de kamer ligt {m.SLACHTOFFER}, zakenman, overleden.',
    '',
    'Jij bent de speurder. In deze werkmap staat alles wat het hotel die nacht',
    'geregistreerd heeft: wie er sliep, wie er werkte, welke deuren opengingen,',
    'welke telefoons gebeld hebben, wat er gevonden is en wat de mensen verklaarden.',
    '',
    'Let op: de werkmap alleen brengt je er niet. Je leerkracht heeft dossierkaarten',
    'met gegevens die niet in het systeem staan — het verslag van de wetsdokter,',
    'wat getuigen zagen, wat er over de kleding geweten is. Vraag ze op het moment',
    'dat de opdracht het zegt.',
    '',
    'Werk netjes. Wie straks moet uitleggen hoe hij bij zijn dader komt,',
    'heeft zijn tussenstappen nodig.',
    '',
    'De tijdstippen staan als datum + uur, want de nacht loopt over middernacht heen.',
    'Sorteer dus op het tijdstip zelf, niet op de tekst.',
    '',
    'De doorgangen registreren elke keer dat een deur met een badge geopend wordt.',
    'De hoofdingang gaat van binnenuit vanzelf open: gasten die buitengaan staan',
    'dus niet in het logboek. Personeel moet wel badgen om buiten te gaan.',
]
for i, r in enumerate(regels, start=2):
    ws.cell(row=i, column=1, value=r)
ws.column_dimensions['A'].width = 95

# --- Gegevensbladen ---------------------------------------------------
blad('Gasten', ['Kamer', 'Naam', 'Aankomst', 'Vertrek', 'Land'],
     m.GASTEN, [10, 24, 14, 14, 14], 'Gastenlijst — nacht van 14 op 15 maart')

blad('Personeel', ['Nr', 'Naam', 'Functie', 'Shift van', 'Shift tot', 'Lengte (cm)'],
     m.PERSONEEL, [6, 24, 24, 12, 12, 13], 'Personeel in dienst die nacht')

blad('Badgehouders', ['Badge', 'Houder', 'Soort'],
     m.BADGES, [12, 28, 14], 'Welke badge hoort bij wie')

ws = blad('Toegangslogboek', ['Tijdstip', 'Badge', 'Deur'],
     [(stip(t), b_, d) for t, b_, d in m.LOGBOEK], [16, 12, 30],
     'Toegangslogboek — elke deuropening met badge')
for r in range(4, 4 + len(m.LOGBOEK)):
    ws.cell(row=r, column=1).number_format = STIPOPMAAK

ws = blad('Telefoonlogboek', ['Tijdstip', 'Van', 'Naar', 'Duur (min)'],
     [(stip(t), v, n, d) for t, v, n, d in m.TELEFOON], [16, 28, 30, 12],
     'Telefoonlogboek van de centrale')
for r in range(4, 4 + len(m.TELEFOON)):
    ws.cell(row=r, column=1).number_format = STIPOPMAAK

blad('Voorwerpen', ['Nr', 'Voorwerp', 'Gevonden in', 'Bijzonderheid'],
     m.VOORWERPEN, [7, 30, 32, 40], 'Voorwerpen aangetroffen door de spoorzoekers')

ws = blad('Verklaringen', ['Tijdstip', 'Wie', 'Verklaring'],
     [(stip(t), w, v) for t, w, v in m.VERKLARINGEN], [16, 26, 78],
     'Verklaringen afgenomen na de vondst')
for r in range(4, 4 + len(m.VERKLARINGEN)):
    ws.cell(row=r, column=1).number_format = STIPOPMAAK

# --- Antwoordblad -----------------------------------------------------
ws = wb.create_sheet('Antwoordblad')
ws['A1'] = 'ANTWOORDBLAD — vul in terwijl je werkt'; ws['A1'].font = Font(bold=True, size=14, color='1F3864')
vragen = [
    ('1', 'Hoeveel deuropeningen staan er in totaal in het toegangslogboek?', ''),
    ('2', 'Welke badges openden tussen 23:40 en 00:20 een deur op de 3e verdieping?', ''),
    ('3', 'Bij wie horen die badges?', ''),
    ('4', 'Welke badge hoort bij niemand persoonlijk?', ''),
    ('5', 'Hoe vaak is die badge die nacht in totaal gebruikt?', ''),
    ('6', 'Wie kwam tussen 21:58 en 22:15 in de receptie achterkamer?', ''),
    ('7', 'Wie van het personeel is 1m85 of groter?', ''),
    ('8', 'Wiens shift was al gedaan terwijl die persoon nergens bij een uitgang staat?', ''),
    ('9', 'Welk telefoongesprek is verdacht, en waarom?', ''),
    ('10', 'Welke voorwerpen wijzen naar de dader?', ''),
    ('', '', ''),
    ('DADER', 'Wie heeft het gedaan?', ''),
    ('BEWIJS', 'Geef je drie sterkste bewijsstukken.', ''),
]
ws['A3'] = 'Nr'; ws['B3'] = 'Vraag'; ws['C3'] = 'Jouw antwoord'
for c in ('A3','B3','C3'):
    ws[c].fill, ws[c].font = KOP, KOPTK
    ws[c].alignment = Alignment(horizontal='center')
for r, (nr, vraag, antw) in enumerate(vragen, start=4):
    ws.cell(row=r, column=1, value=nr).font = Font(bold=True)
    ws.cell(row=r, column=2, value=vraag)
    ws.cell(row=r, column=3, value=antw).fill = PatternFill('solid', fgColor='FFF2CC')
    for k in range(1, 4):
        ws.cell(row=r, column=k).border = RAND
        ws.cell(row=r, column=k).alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[r].height = 30
ws.column_dimensions['A'].width = 9
ws.column_dimensions['B'].width = 62
ws.column_dimensions['C'].width = 46
ws.freeze_panes = 'A4'

uit = '/home/user/5OSEXCEL/lesmateriaal/moordzaak/moordzaak-startbestand.xlsx'
wb.save(uit)
print('geschreven:', uit)
print('bladen:', wb.sheetnames)
