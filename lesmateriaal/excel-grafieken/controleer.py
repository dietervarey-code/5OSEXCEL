# -*- coding: utf-8 -*-
"""Kijkt na of opdracht, startbestand, antwoorden en sleutel bij elkaar passen."""
import collections
import importlib.util
import os
import re
import sys
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')

WERKBOEK = os.path.join(HIER, 'grafieken-startbestand.xlsx')
FOUTEN, GETELD = [], collections.Counter()


def eis(voorwaarde, bericht):
    if not voorwaarde:
        FOUTEN.append(bericht)
    return bool(voorwaarde)


def cel_van(adres):
    m = re.match(r'^([A-Z]{1,2})(\d+)$', adres)
    return (column_index_from_string(m.group(1)), int(m.group(2))) if m else None


# ==========================================================================
#  1. de gegevens
# ==========================================================================
eis(len({r[0] for r in g.B1}) == len(g.B1), 'dubbele clubnaam')
for club, stad, toe in g.B1:
    eis(toe > 0, f'{club} heeft geen toeschouwers')
    GETELD['clubs nagerekend'] += 1
eis([r[2] for r in g.B1] == sorted((r[2] for r in g.B1), reverse=True),
    'de clublijst staat niet aflopend gesorteerd — dan klopt blad 1 niet')

eis([j for j, _ in g.B2] == sorted(j for j, _ in g.B2),
    'de bevolkingsreeks staat niet op jaartal gesorteerd')
eis(all(b > a_ for (_, a_), (_, b) in zip(g.B2, g.B2[1:])),
    'de bevolking daalt ergens — dan klopt het verhaal van blad 2 niet')
eis(g.B2[-1][0] - g.B2[-2][0] != g.B2[1][0] - g.B2[0][0],
    'alle metingen liggen even ver uit elkaar — dan vervalt de les van kaart 14')
GETELD['metingen nagerekend'] += len(g.B2)

eis(len(g.B3) <= 5,
    f'{len(g.B3)} delen in een cirkeldiagram — kaart 11 zegt hoogstens vier of vijf')
eis(abs(sum(a.B3_AANDEEL) - 1) < 1e-9, 'de aandelen tellen niet op tot 100 %')
for naam, waarde in g.B3:
    eis(waarde > 0, f'{naam} heeft geen waarde')
    GETELD['productgroepen nagerekend'] += 1

eis(len(g.B4_KOP) == len(g.B4_LANDEN) + 1, 'de koprij van blad 4 klopt niet')
for rij in g.B4:
    eis(len(rij) == len(g.B4_LANDEN) + 1, f'rij {rij[0]} heeft niet alle landen')
    eis(all(w > 0 for w in rij[1:]), f'nul of negatief bezoekersaantal in {rij[0]}')
    GETELD['toerismejaren nagerekend'] += 1

for speler, nummer, dp in g.B5:
    eis(1 <= nummer <= len(g.B1), f'clubnummer {nummer} bestaat niet ({speler})')
    eis(dp > 0, f'{speler} heeft geen doelpunten')
    GETELD['spelers nagerekend'] += 1
eis(len({r[0] for r in g.B5}) == len(g.B5), 'dubbele spelersnaam')

# ==========================================================================
#  2. de antwoorden, langs een tweede weg
# ==========================================================================
eis(a.B1_VERSCHIL == max(r[2] for r in g.B1) - min(r[2] for r in g.B1),
    'B1_VERSCHIL klopt niet')
eis(abs(sum(v for v in a.B2_GROEI if v is not None)
        - (a.B2_INWONERS[-1] - a.B2_INWONERS[0])) < 1e-9,
    'de som van de groeikolom klopt niet met het verschil eerste-laatste')
eis(a.B3_TOTAAL == sum(w for _, w in g.B3), 'B3_TOTAAL klopt niet')
eis(a.B4_REEKSEN == len(g.B4_LANDEN), 'B4_REEKSEN klopt niet')
eis(a.B5_TOPSCORER_DOELPUNTEN == max(r[2] for r in g.B5), 'topscorer klopt niet')
eis(a.B5_TOPSCORER_CLUB == g.club(
    next(n for s, n, _ in g.B5 if s == a.B5_TOPSCORER))[0], 'club van de topscorer')

for oef in o.OEFENINGEN:
    for cel, wat, antwoord, punten in oef['vragen']:
        eis(antwoord is not None and antwoord != '',
            f"{oef['blad']} {cel}: geen antwoord")
        eis(punten > 0, f"{oef['blad']} {cel}: geen punten")
        GETELD['vragen gekeurd'] += 1

# ==========================================================================
#  3. het startbestand
# ==========================================================================
LIJSTEN = [
    ('1 Toeschouwers', g.B1_KOPRIJ, g.B1_START, g.B1_EIND, g.B1_KOP, o.V1),
    ('2 Bevolking', g.B2_KOPRIJ, g.B2_START, g.B2_EIND, g.B2_KOP, o.V2),
    ('3 Export', g.B3_KOPRIJ, g.B3_START, g.B3_EIND, g.B3_KOP, o.V3),
    ('4 Toerisme', g.B4_KOPRIJ, g.B4_START, g.B4_EIND, g.B4_KOP, o.V4),
    ('5 Topscorers', g.B5_KOPRIJ, g.B5_START, g.B5_EIND, g.B5_KOP, o.V5),
]

eis(os.path.exists(WERKBOEK), 'het startbestand bestaat niet')
if os.path.exists(WERKBOEK):
    wb = load_workbook(WERKBOEK)
    eis({x[0] for x in LIJSTEN} <= set(wb.sheetnames), 'ontbrekende bladen')

    for blad, koprij, start, eind, koppen, vraagkol in LIJSTEN:
        if blad not in wb.sheetnames:
            continue
        ws = wb[blad]
        breedte = len(koppen)
        for i, titel in enumerate(koppen, start=1):
            eis(ws.cell(row=koprij, column=i).value == titel,
                f'{blad}: koprij verschilt in kolom {get_column_letter(i)}')
        temaken = set()
        for oef in o.OEFENINGEN:
            if oef['blad'] != blad:
                continue
            for waar, *_rest in oef.get('vullen', []):
                temaken.add(column_index_from_string(
                    re.match(r'([A-Z]+)', waar).group(1)))
        for kolom in range(1, breedte + 1):
            heeft = any(ws.cell(row=r, column=kolom).value is not None
                        for r in range(start, eind + 1))
            eis(heeft or kolom in temaken,
                f'{blad}: kolom {get_column_letter(kolom)} blijft leeg')
        eis(vraagkol >= breedte + 2,
            f'{blad}: het vragenblok staat te dicht bij de lijst')
        for rij in range(start, eind + 1):
            eis(ws.cell(row=rij, column=breedte + 1).value is None,
                f'{blad}: kolom {get_column_letter(breedte + 1)} rij {rij} hoort '
                'leeg te blijven als scheiding')
        GETELD['lijsten gekeurd'] += 1

    for oef in o.OEFENINGEN:
        ws = wb[oef['blad']]
        for cel, *_rest in oef['vragen']:
            kolom, rij = cel_van(cel)
            eis(ws.cell(row=rij, column=kolom).value is None,
                f"{oef['blad']} {cel} moet leeg zijn in het startbestand")
            eis(ws.cell(row=rij, column=kolom - 1).value is not None,
                f"{oef['blad']} {cel}: er staat geen vraag naast")
            GETELD['antwoordcellen leeg'] += 1
        for waar, *_rest in oef.get('vullen', []):
            links, rechts = waar.split(':')
            k1, r1 = cel_van(links); k2, r2 = cel_van(rechts)
            for k in range(k1, k2 + 1):
                for r in range(r1, r2 + 1):
                    eis(ws.cell(row=r, column=k).value is None,
                        f"{oef['blad']} {get_column_letter(k)}{r} moet leeg zijn")
                    GETELD['te vullen cellen leeg'] += 1

    # het startbestand bevat geen enkele grafiek: die moeten ze zelf maken
    for ws in wb.worksheets:
        eis(not getattr(ws, '_charts', None),
            f'{ws.title}: er zit al een grafiek in het startbestand')

    # de clubtabel naast de topscorerslijst klopt met blad 1
    ws = wb['5 Topscorers']
    for i, (club, stad, _toe) in enumerate(g.B1):
        r = g.B5_ZOEKRIJ + 1 + i
        for kolom, verwacht in ((g.B5_ZOEKKOL, g.clubcode(i + 1)),
                                (g.B5_ZOEKKOL + 1, club),
                                (g.B5_ZOEKKOL + 2, stad)):
            eis(ws.cell(row=r, column=kolom).value == verwacht,
                f'5 Topscorers: de clubtabel wijkt af op rij {r}')
            GETELD['clubtabel nagekeken'] += 1

# ==========================================================================
#  4. de opdracht zelf
# ==========================================================================
TYPES = {'staafdiagram', 'lijndiagram', 'cirkeldiagram',
         'lijndiagram met drie reeksen'}
for oef in o.OEFENINGEN:
    eis(oef['kaarten'], f"blad {oef['nr']}: geen verwijzing naar de theorie")
    for n in oef['kaarten']:
        eis(n in o.KAARTEN, f"blad {oef['nr']}: kaart {n} bestaat niet")
        eis(10 <= n <= 14,
            f"blad {oef['nr']}: kaart {n} gaat niet over grafieken")
    gr = oef['grafiek']
    eis(gr['soort'] in TYPES, f"blad {oef['nr']}: onbekend type {gr['soort']!r}")
    eis(gr['titel'] and 'Grafiektitel' not in gr['titel'],
        f"blad {oef['nr']}: geen zinnige titel")
    eis(gr['selectie'], f"blad {oef['nr']}: er staat niet wat ze selecteren")
    GETELD['bladen gekeurd'] += 1

# elk grafiektype komt minstens één keer aan bod
gebruikt = {oef['grafiek']['soort'].split(' met ')[0] for oef in o.OEFENINGEN}
eis({'staafdiagram', 'lijndiagram', 'cirkeldiagram'} <= gebruikt,
    f'niet alle drie de basistypes komen aan bod: {sorted(gebruikt)}')

VERBODEN = ['gegevenslabels zijn', 'een legenda heeft alleen zin',
            'klik de grafiek aan. rechtsboven']
tekst = ' '.join(r for oef in o.OEFENINGEN
                 for r in oef['inleiding'] + [oef['klaar'], oef['controle']]).lower()
for woord in VERBODEN:
    eis(woord not in tekst,
        f'de opdracht legt theorie uit ("{woord}") — dat hoort op de kaarten')

eis(sum(int(oef['duur'].split()[0]) for oef in o.OEFENINGEN) == g.DUUR,
    'de duur van de bladen telt niet op tot de aangekondigde duur')
eis(o.TOTAAL == 50, f'de bundel staat op {o.TOTAAL} punten in plaats van 50')

# ==========================================================================
print('=' * 72)
for naam, hoeveel in sorted(GETELD.items()):
    print(f'  {hoeveel:>5}  {naam}')
print('=' * 72)
if FOUTEN:
    print(f'{len(FOUTEN)} PROBLEMEN:')
    for f in FOUTEN:
        print('  !', f)
    sys.exit(1)
print('alles past bij elkaar.')
