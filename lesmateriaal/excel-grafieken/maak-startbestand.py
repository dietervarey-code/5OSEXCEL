# -*- coding: utf-8 -*-
"""Bouwt het startbestand van de grafiekbundel.

Er zitten geen grafieken in: de leerling hoort ze zelf te maken. Elke lijst
staat klaar met de koprij erboven en het vragenblok ernaast.
"""
import importlib.util
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')

UIT = os.path.join(HIER, 'grafieken-startbestand.xlsx')
TITEL = Font(bold=True, size=13)
GROOT = Font(bold=True, size=16)
VET = Font(bold=True)
SCHUIN = Font(italic=True, size=10)

wb = Workbook()


def breedtes(ws, waarden, start=1):
    for i, b in enumerate(waarden, start=start):
        ws.column_dimensions[get_column_letter(i)].width = b


def blad(naam, titel, ondertitel=None):
    ws = wb.create_sheet(naam)
    ws['A1'] = titel
    ws['A1'].font = TITEL
    if ondertitel:
        ws['A2'] = ondertitel
        ws['A2'].font = SCHUIN
    return ws


def lijst(ws, koprij, koppen, rijen):
    for i, k in enumerate(koppen, start=1):
        ws.cell(row=koprij, column=i, value=k).font = VET
    for r, waarden in enumerate(rijen, start=koprij + 1):
        for i, w in enumerate(waarden, start=1):
            ws.cell(row=r, column=i, value=w)


def vragen(ws, rij, kolom, teksten):
    ws.cell(row=rij, column=kolom, value='Vraag').font = VET
    ws.cell(row=rij, column=kolom + 1, value='Antwoord').font = VET
    for i, t in enumerate(teksten, start=1):
        ws.cell(row=rij + i, column=kolom, value=t)


def plek(ws, rij, soort):
    ws.cell(row=rij, column=1,
            value=f'Zet je {soort} hieronder.').font = SCHUIN


# ------------------------------------------------------------------- Start
ws = wb.active
ws.title = 'Start'
ws['A1'] = g.TITEL
ws['A1'].font = GROOT
regels = [
    g.ONDERTITEL, '',
    'Naam: ......................................    Klas: ..........', '',
    f'Vijf werkbladen, samen {g.DUUR} minuten. Op elk blad maak je één grafiek.',
    'Er staat in deze bundel geen theorie. Die staat in:',
    f'   {g.THEORIEBUNDEL} — kaart 10 tot 14, elk op een eigen pagina.',
    'Bij elk blad staat welke kaarten je erbij nodig hebt.', '',
    'Blad 1  Toeschouwers   8 min   kaart 10 en 12      staafdiagram',
    'Blad 2  Bevolking     10 min   kaart 11, 12, 14    lijndiagram',
    'Blad 3  Export         8 min   kaart 11 en 14      cirkeldiagram',
    'Blad 4  Toerisme      10 min   kaart 12 en 14      drie reeksen samen',
    'Blad 5  Topscorers    14 min   kaart 10 en 13      eerst aanvullen, dan tekenen',
    '',
    'Drie regels die op elk blad gelden',
    '  1. Het type grafiek is geen smaakkwestie. Vergelijken is een staafdiagram,',
    '     een verloop in de tijd is een lijn, een aandeel van een geheel is een',
    '     cirkel. Bij elk blad staat welke vraag je beantwoordt.',
    '  2. Elke grafiek krijgt een titel die zegt wat je ziet. "Grafiektitel" laten',
    '     staan is geen titel.',
    '  3. Laat je verticale as bij nul beginnen, tenzij er uitdrukkelijk iets',
    '     anders staat.',
    '',
    'EERST DOEN: sla dit bestand op als  grafieken-jouwnaam.xlsx',
    '', g.WAARSCHUWING,
]
for i, r in enumerate(regels, start=2):
    ws.cell(row=i, column=1, value=r)
ws['A17'].font = VET
ws['A27'].font = SCHUIN
ws['A27'].alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[27].height = 86
breedtes(ws, [88])

# ------------------------------------------------------------ 1 Toeschouwers
oef = o.OEFENINGEN[0]
ws = blad(oef['blad'], 'Toeschouwers per club')
lijst(ws, g.B1_KOPRIJ, g.B1_KOP, g.B1)
vragen(ws, g.B1_KOPRIJ, o.V1, [
    'Welke club heeft de langste balk?',
    'Hoeveel toeschouwers zitten er tussen de langste en de kortste balk?',
    'Hoeveel reeksen heeft je grafiek?'])
plek(ws, g.B1_EIND + 2, oef['grafiek']['soort'])
breedtes(ws, [22, 12, 14])
breedtes(ws, [3, 52, 16], start=o.V1 - 1)

# --------------------------------------------------------------- 2 Bevolking
oef = o.OEFENINGEN[1]
ws = blad(oef['blad'], 'De bevolking van Türkiye', g.B2_VOETNOOT)
lijst(ws, g.B2_KOPRIJ, g.B2_KOP, [(j, i, None) for j, i in g.B2])
vragen(ws, g.B2_KOPRIJ, o.V2, [
    'Tussen welke twee jaren groeide de bevolking het sterkst?',
    'Hoeveel miljoen was dat?',
    'Hoeveel jaar zit er tussen de laatste twee metingen?'])
plek(ws, g.B2_EIND + 2, oef['grafiek']['soort'])
breedtes(ws, [10, 20, 28])
breedtes(ws, [3, 52, 16], start=o.V2 - 1)

# ------------------------------------------------------------------ 3 Export
oef = o.OEFENINGEN[2]
ws = blad(oef['blad'], 'De export per productgroep', g.B3_VOETNOOT)
lijst(ws, g.B3_KOPRIJ, g.B3_KOP, [(naam, w, None) for naam, w in g.B3])
vragen(ws, g.B3_KOPRIJ, o.V3, [
    'Welk aandeel heeft de grootste groep, in procent?',
    'Welk grafiektype neem je bij vijftien productgroepen?'])
plek(ws, g.B3_EIND + 2, oef['grafiek']['soort'])
breedtes(ws, [24, 20, 12])
breedtes(ws, [3, 48, 16], start=o.V3 - 1)

# ---------------------------------------------------------------- 4 Toerisme
oef = o.OEFENINGEN[3]
ws = blad(oef['blad'], 'Bezoekers per herkomstland', g.B4_VOETNOOT)
lijst(ws, g.B4_KOPRIJ, g.B4_KOP, g.B4)
vragen(ws, g.B4_KOPRIJ, o.V4, [
    f'Welk land viel in {a.B4_VALJAAR} het sterkst terug?',
    f'Welk land staat in {g.B4[-1][0]} het hoogst?',
    'Hoeveel reeksen heeft je grafiek?'])
plek(ws, g.B4_EIND + 2, oef['grafiek']['soort'])
breedtes(ws, [10, 13, 13, 13])
breedtes(ws, [3, 46, 16], start=o.V4 - 1)

# -------------------------------------------------------------- 5 Topscorers
oef = o.OEFENINGEN[4]
ws = blad(oef['blad'], 'De topscorers van het seizoen', g.B5_VOETNOOT)
lijst(ws, g.B5_KOPRIJ, g.B5_KOP,
      [(speler, g.clubcode(n), dp, None, None) for speler, n, dp in g.B5])
k = g.B5_ZOEKKOL
ws.cell(row=g.B5_ZOEKRIJ - 1, column=k,
        value='Clubtabel — niet aanpassen').font = SCHUIN
for i, titel in enumerate(['Code', 'Club', 'Stad'], start=k):
    ws.cell(row=g.B5_ZOEKRIJ, column=i, value=titel).font = VET
for i, (club, stad, _toe) in enumerate(g.B1):
    r = g.B5_ZOEKRIJ + 1 + i
    ws.cell(row=r, column=k, value=g.clubcode(i + 1))
    ws.cell(row=r, column=k + 1, value=club)
    ws.cell(row=r, column=k + 2, value=stad)
vragen(ws, g.B5_KOPRIJ, o.V5, [
    'Wie maakte de meeste doelpunten?',
    'Bij welke club zit die speler?',
    'Hoeveel reeksen heeft je grafiek?'])
plek(ws, g.B5_ZOEKRIJ + len(g.B1) + 2, oef['grafiek']['soort'])
breedtes(ws, [18, 11, 13, 22, 12])
breedtes(ws, [3, 10, 22, 12], start=k - 1)
breedtes(ws, [3, 40, 16], start=o.V5 - 1)

wb.save(UIT)
print('geschreven:', UIT)
for ws in wb.worksheets:
    gevuld = sum(1 for rij in ws.iter_rows() for c in rij if c.value is not None)
    print(f'  {ws.title:<16} {gevuld:>4} gevulde cellen   {ws.dimensions}')
