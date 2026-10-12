# -*- coding: utf-8 -*-
"""Bouwt het startbestand van de Ottomaanse bundel.

Bewust kaal: de opmaak hoort bij de opdracht. Alleen de kolombreedtes staan
goed, zodat niemand tijd verliest met slepen.

Twee dingen zijn hier anders dan in de vorige bundels, en dat is met opzet:

1. De lijsten op blad 2 en blad 3 zijn BRON van een draaitabel. Daarom staat
   er geen totaalrij onder en geen lege rij in, en staan de vragenblokjes
   ERNAAST met één lege kolom ertussen. Excel zou een blokje eronder anders
   als deel van de lijst beschouwen.
2. Er zit géén draaitabel in het bestand. openpyxl kan die niet schrijven, en
   dat komt goed uit: de leerling moet ze zelf maken.
"""
import importlib.util
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')

UIT = os.path.join(HIER, 'ottomaans-startbestand.xlsx')
TITEL = Font(bold=True, size=13)
GROOT = Font(bold=True, size=16)
VET = Font(bold=True)
SCHUIN = Font(italic=True, size=10)

wb = Workbook()


def breedtes(ws, waarden, start=1):
    for i, b in enumerate(waarden, start=start):
        ws.column_dimensions[get_column_letter(i)].width = b


def koprij(ws, rij, koppen, kolom=1):
    for i, k in enumerate(koppen, start=kolom):
        ws.cell(row=rij, column=i, value=k)


def labels(ws, rij, teksten, kolom=1):
    for i, t in enumerate(teksten):
        ws.cell(row=rij + i, column=kolom, value=t)


def blad(naam, titel, ondertitel=None):
    ws = wb.create_sheet(naam)
    ws['A1'] = titel
    ws['A1'].font = TITEL
    if ondertitel:
        ws['A2'] = ondertitel
        ws['A2'].font = SCHUIN
    return ws


# ------------------------------------------------------------------- Start
ws = wb.active
ws.title = 'Start'
ws['A1'] = g.TITEL
ws['A1'].font = GROOT
labels(ws, 2, [
    g.ONDERTITEL,
    '',
    'Naam: ......................................    Klas: ..........',
    '',
    f'Vijf werkbladen, samen ongeveer {g.DUUR} minuten. Werk ze in volgorde af:',
    'blad 2 heeft blad 1 nodig, en blad 4 heeft blad 2 en 3 nodig.',
    '',
    'Blad 1  Sultans          8 min   herhaling: rekenen, ALS, tellen',
    'Blad 2  Gebeurtenissen  12 min   herhaling: VERT.ZOEKEN uit de sultantabel',
    'Blad 3  Oorlogen        10 min   herhaling: SOM, SOM.ALS, ALS',
    'Blad 4  Draaitabel      12 min   NIEUW: twee lijsten samenvatten',
    'Blad 5  Grafiek          8 min   NIEUW: een staafdiagram en een lijndiagram',
    '',
    'Wat je ernaast legt',
    f'  {g.FUNCTIEBUNDEL} — de functiefiches die je al kent, voor blad 1 tot 3.',
    f'  {g.THEORIEBUNDEL} — de nieuwe theoriebundel, voor blad 4 en 5.',
    '',
    'Twee regels die op dit bestand anders zijn dan vroeger',
    '  1. De lijsten op blad 2 en 3 worden straks de BRON van een draaitabel.',
    '     Voeg er geen lege rij in en zet er geen totaalrij onder. De vragen',
    '     staan daarom ERNAAST, met een lege kolom ertussen. Hou dat zo.',
    '  2. Blad 4 en 5 hebben geen formules nodig. Daar wordt gesleept en getekend.',
    '',
    'Controleer jezelf',
    '  Onderaan elk blad in de opdrachtbundel staat een CONTROLEGETAL.',
    '  Komt jouw cel daarop uit, dan mag je verder. Typ het nooit zelf in.',
    '',
    'EERST DOEN: sla dit bestand op als  ottomaans-jouwnaam.xlsx',
    '',
    g.WAARSCHUWING,
])
for adres in ('A16', 'A20', 'A26'):
    ws[adres].font = VET
ws['A31'].font = SCHUIN
ws['A31'].alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[31].height = 72
breedtes(ws, [86])

# ---------------------------------------------------------------- 1 Sultans
ws = blad('1 Sultans', 'De sultans en hun regeerperiode',
          'Dit is ook de zoektabel voor blad 2. Verander er niets aan.')
ws[f'A{g.B1_KOPRIJ - 1}'] = g.B1_DREMPELTEKST
ws[f'B{g.B1_KOPRIJ - 1}'] = g.B1_DREMPEL
koprij(ws, g.B1_KOPRIJ, g.B1_KOP)
for r, (code, naam, begin, einde) in enumerate(g.B1, start=g.B1_START):
    ws.cell(row=r, column=1, value=code)
    ws.cell(row=r, column=2, value=naam)
    ws.cell(row=r, column=3, value=begin)
    ws.cell(row=r, column=4, value=einde)
labels(ws, g.B1_SAMENVATTING, [
    'Totaal aantal regeringsjaren',
    'Gemiddelde regeerperiode',
    'Langste regeerperiode',
    'Kortste regeerperiode',
    'Aantal sultans in de lijst',
    'Aantal met een lange regering',
])
ws.cell(row=g.B1_SAMENVATTING + 7, column=1, value=g.B1_VOETNOOT).font = SCHUIN
breedtes(ws, [10, 17, 13, 14, 16, 20])

# --------------------------------------------------------- 2 Gebeurtenissen
ws = blad('2 Gebeurtenissen', 'Vijfendertig gebeurtenissen uit de geschiedenis '
                              'van het rijk')
koprij(ws, g.B2_KOPRIJ, g.B2_KOP)
for r, (jaar, naam, soort, regio, code) in enumerate(g.B2, start=g.B2_START):
    ws.cell(row=r, column=1, value=jaar)
    ws.cell(row=r, column=2, value=naam)
    ws.cell(row=r, column=3, value=soort)
    ws.cell(row=r, column=4, value=regio)
    ws.cell(row=r, column=5, value=g.eeuw(jaar))
    ws.cell(row=r, column=6, value=code)
k = g.B2_VRAAGKOL
ws.cell(row=g.B2_VRAAGRIJ, column=k, value='Vraag').font = VET
ws.cell(row=g.B2_VRAAGRIJ, column=k + 1, value='Antwoord').font = VET
labels(ws, g.B2_VRAAGRIJ + 1, [
    'Aantal gebeurtenissen in de lijst',
    f'Aantal van de soort {g.B2_TELSOORT}',
    f'Aantal in de regio {g.B2_TELREGIO}',
    f'Aantal in de {g.B2_TELEEUW}',
], kolom=k)
breedtes(ws, [8, 38, 15, 18, 11, 13, 17, 16, 26])
breedtes(ws, [3, 30, 12], start=k - 1)

# --------------------------------------------------------------- 3 Oorlogen
ws = blad('3 Oorlogen', 'De grote oorlogen van het rijk',
          g.B3_VOETNOOT)
ws[f'A{g.B3_KOPRIJ - 1}'] = g.B3_DREMPELTEKST
ws[f'B{g.B3_KOPRIJ - 1}'] = g.B3_DREMPEL
koprij(ws, g.B3_KOPRIJ, g.B3_KOP)
for r, (naam, tegen, begin, einde, uitkomst) in enumerate(g.B3, start=g.B3_START):
    ws.cell(row=r, column=1, value=naam)
    ws.cell(row=r, column=2, value=tegen)
    ws.cell(row=r, column=3, value=begin)
    ws.cell(row=r, column=4, value=einde)
    ws.cell(row=r, column=6, value=uitkomst)
k = g.B3_VRAAGKOL
ws.cell(row=g.B3_VRAAGRIJ, column=k, value='Vraag').font = VET
ws.cell(row=g.B3_VRAAGRIJ, column=k + 1, value='Antwoord').font = VET
labels(ws, g.B3_VRAAGRIJ + 1, [
    'Totaal aantal oorlogsjaren',
    'Gemiddelde duur van een oorlog',
    'Duur van de langste oorlog',
    'Duur van de kortste oorlog',
    f'Aantal oorlogen met uitkomst {g.B3_TELUITKOMST}',
    f'Oorlogsjaren tegen {g.B3_SOMLAND}',
    f'Oorlogsjaren met uitkomst {g.B3_SOMUITKOMST}',
], kolom=k)
breedtes(ws, [34, 15, 9, 9, 14, 13, 16])
breedtes(ws, [3, 34, 12], start=k - 1)

# ------------------------------------------------------------- 4 Draaitabel
ws = blad('4 Draaitabel', 'Twee draaitabellen')
v, k = g.B4_VRAAGRIJ, g.B4_VRAAGKOL
ws.cell(row=v - 1, column=k,
        value='Lees de antwoorden af van je eigen draaitabellen en typ ze hier.'
        ).font = SCHUIN
ws.cell(row=v, column=k, value='Vraag').font = VET
ws.cell(row=v, column=k + 1, value='Antwoord').font = VET
labels(ws, v + 1, [
    'In welke eeuw gebeurde er het meest?',
    'Hoeveel gebeurtenissen waren dat?',
    'Welke soort gebeurtenis komt het vaakst voor?',
    'Tegen welke tegenstander was het rijk het langst in oorlog?',
    'Hoeveel oorlogsjaren waren dat samen?',
], kolom=k)
for nummer, draai in ((1, g.B4_DRAAI1), (2, g.B4_DRAAI2)):
    rij = int(draai['anker'][1:]) - 2
    ws.cell(row=rij, column=1, value=draai['titel']).font = VET
    ws.cell(row=rij + 1, column=1,
            value=f"(bron: blad '{draai['bron']}' — zet de linkerbovenhoek van je "
                  f"draaitabel in {draai['anker']})").font = SCHUIN
breedtes(ws, [16, 11, 11, 11, 11, 11, 11, 11, 11])
breedtes(ws, [3, 58, 18], start=g.B4_VRAAGKOL - 1)

# ---------------------------------------------------------------- 5 Grafiek
ws = blad('5 Grafiek', 'Het geschatte aantal inwoners', g.B5_BRONTEKST)
koprij(ws, g.B5_KOPRIJ, g.B5_KOP)
for r, (jaar, inwoners) in enumerate(g.B5, start=g.B5_START):
    ws.cell(row=r, column=1, value=jaar)
    ws.cell(row=r, column=2, value=inwoners)
k = g.B5_VRAAGKOL
ws.cell(row=g.B5_VRAAGRIJ, column=k, value='Vraag').font = VET
ws.cell(row=g.B5_VRAAGRIJ, column=k + 1, value='Antwoord').font = VET
labels(ws, g.B5_VRAAGRIJ + 1, [
    'Tussen welke twee metingen daalt het sterkst?',
    'Hoeveel miljoen daalde de schatting toen?',
    'Welk grafiektype past bij een verloop in de tijd?',
    'Welke tegenstander heeft de langste balk in je staafdiagram?',
], kolom=k)
ws.cell(row=g.B5_EIND + 3, column=1,
        value='Zet je twee grafieken hieronder, naast elkaar.').font = SCHUIN
breedtes(ws, [10, 30, 26])
breedtes(ws, [3, 52, 20], start=k - 1)

wb.save(UIT)
print('geschreven:', UIT)
for ws in wb.worksheets:
    gevuld = sum(1 for rij in ws.iter_rows() for c in rij if c.value is not None)
    print(f'  {ws.title:<18} {gevuld} gevulde cellen')
