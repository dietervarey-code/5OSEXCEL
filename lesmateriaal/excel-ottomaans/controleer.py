# -*- coding: utf-8 -*-
"""Kijkt na of opdracht, startbestand, antwoorden en sleutel bij elkaar passen.

Faalt luid zodra er iets uit elkaar loopt. Dat is het hele punt: de gegevens
staan op één plaats, en als iemand daar een regel bij zet moet de rest
meeschuiven of stuk gaan — niet stilletjes verkeerd worden.

Er staan geen vaste rijnummers in dit script.
"""
import importlib.util
import os
import re
import sys
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')

WERKBOEK = os.path.join(HIER, 'ottomaans-startbestand.xlsx')
FOUTEN = []
GETELD = {}


def eis(voorwaarde, bericht):
    if not voorwaarde:
        FOUTEN.append(bericht)
    return bool(voorwaarde)


def tel(naam, hoeveel=1):
    GETELD[naam] = GETELD.get(naam, 0) + hoeveel


CEL = re.compile(r'^([A-Z]{1,2})(\d+)$')
BEREIK = re.compile(r'^([A-Z]{1,2})(\d+):([A-Z]{1,2})(\d+)$')


def cellen_van(waar):
    """'E5:E17' -> [(5,5),(5,6),...]. Geeft None als het geen celadres is."""
    m = CEL.match(waar)
    if m:
        return [(column_index_from_string(m.group(1)), int(m.group(2)))]
    m = BEREIK.match(waar)
    if m:
        k1, r1 = column_index_from_string(m.group(1)), int(m.group(2))
        k2, r2 = column_index_from_string(m.group(3)), int(m.group(4))
        return [(k, r) for k in range(k1, k2 + 1) for r in range(r1, r2 + 1)]
    return None


# ==========================================================================
#  1. de gegevens zijn intern samenhangend
# ==========================================================================
codes = {c for c, *_ in g.B1}
eis(len(codes) == len(g.B1), 'dubbele sultancode in B1')
for jaar, naam, soort, regio, code in g.B2:
    if not eis(code in codes, f'onbekende sultancode {code} bij "{naam}"'):
        continue
    _, sultan, begin, einde = g.sultan(code)
    eis(begin <= jaar <= einde,
        f'{jaar} "{naam}" valt buiten de regeerperiode van {sultan} '
        f'({begin}-{einde})')
    eis(soort in g.B2_SOORTEN, f'onbekende soort "{soort}" bij "{naam}"')
    eis(regio in g.B2_REGIOS, f'onbekende regio "{regio}" bij "{naam}"')
    tel('gebeurtenissen nagerekend')

for naam, tegen, begin, einde, uitkomst in g.B3:
    eis(einde >= begin, f'oorlog "{naam}" eindigt voor ze begint')
    eis(uitkomst in g.B3_UITKOMSTEN,
        f'onbekende uitkomst "{uitkomst}" bij "{naam}"')
    tel('oorlogen nagerekend')

for code, naam, begin, einde in g.B1:
    eis(einde > begin, f'{naam} heeft geen regeerperiode')
    tel('sultans nagerekend')

eis(len({j for j, _ in g.B5}) == len(g.B5), 'dubbel jaartal in de bevolkingsreeks')
eis([j for j, _ in g.B5] == sorted(j for j, _ in g.B5),
    'de bevolkingsreeks staat niet op jaartal gesorteerd')

# de telwoorden in de opdracht moeten echt in de lijst voorkomen
eis(g.B2_TELSOORT in g.B2_SOORTEN, 'er wordt geteld op een soort die niet bestaat')
eis(g.B2_TELREGIO in g.B2_REGIOS, 'er wordt geteld op een regio die niet bestaat')
eis(g.B2_TELEEUW in {g.eeuw(j) for j, *_ in g.B2},
    'er wordt geteld op een eeuw die niet in de lijst zit')
eis(g.B3_SOMLAND in {t for _, t, *_ in g.B3},
    'er wordt opgeteld voor een tegenstander die niet bestaat')
eis(g.B3_TELUITKOMST in g.B3_UITKOMSTEN, 'onbekende uitkomst in de telling')
eis(g.B3_SOMUITKOMST in g.B3_UITKOMSTEN, 'onbekende uitkomst in de som')

# ==========================================================================
#  2. de antwoorden kloppen, langs een tweede weg uitgerekend
# ==========================================================================
eis(a.B1_SOM == sum(e - b for _, _, b, e in g.B1), 'B1_SOM klopt niet')
eis(a.B1_LANGE == len([1 for _, _, b, e in g.B1 if e - b >= g.B1_DREMPEL]),
    'B1_LANGE klopt niet')
eis(a.B2_AANTAL == len(g.B2), 'B2_AANTAL klopt niet')
eis(a.B3_SOM == sum(e - b for _, _, b, e, _ in g.B3), 'B3_SOM klopt niet')
eis(a.B3_SOMLAND == sum(e - b for _, t, b, e, _ in g.B3 if t == g.B3_SOMLAND),
    'B3_SOMLAND klopt niet')
eis(a.B3_SOMUITKOMST == sum(e - b for _, _, b, e, u in g.B3
                            if u == g.B3_SOMUITKOMST),
    'B3_SOMUITKOMST klopt niet')

# de draaitabellen, cel per cel opnieuw geteld
for r in a.D1_RIJLABELS:
    for k in a.D1_KOLLABELS:
        verwacht = len([1 for j, _, s, _, _ in g.B2
                        if s == r and g.eeuw(j) == k])
        eis(a.D1_CEL[(r, k)] == verwacht,
            f'draaitabel 1 klopt niet bij {r} / {k}')
        tel('cellen van draaitabel 1')
eis(a.D1_EINDTOTAAL == len(g.B2), 'eindtotaal van draaitabel 1 klopt niet')
eis(a.D1_EINDTOTAAL == sum(a.D1_KOLTOTAAL.values()) == sum(a.D1_RIJTOTAAL.values()),
    'de randtotalen van draaitabel 1 lopen niet gelijk')

for t, jaren in a.D2.items():
    verwacht = sum(e - b for _, tt, b, e, _ in g.B3 if tt == t)
    eis(jaren == verwacht, f'draaitabel 2 klopt niet bij {t}')
    tel('regels van draaitabel 2')
eis(a.D2_EINDTOTAAL == a.B3_SOM,
    'het eindtotaal van draaitabel 2 is niet hetzelfde als de SOM op blad 3 — '
    'dan deugt het controlegetal in de opdracht niet')

eis(sum(v for v in a.B5_VERSCHIL if v is not None)
    == a.B5_INWONERS[-1] - a.B5_INWONERS[0],
    'de som van de verschilkolom is niet het verschil tussen eerste en laatste')

# ==========================================================================
#  3. het startbestand: te maken cellen leeg, gegeven cellen gevuld
# ==========================================================================
eis(os.path.exists(WERKBOEK), 'het startbestand bestaat niet — draai eerst '
                              'maak-startbestand.py')
if os.path.exists(WERKBOEK):
    wb = load_workbook(WERKBOEK)
    bladen = {oef['blad'] for oef in o.OEFENINGEN}
    eis(bladen <= set(wb.sheetnames),
        f'ontbrekende werkbladen: {sorted(bladen - set(wb.sheetnames))}')

    for oef in o.OEFENINGEN:
        if oef['blad'] not in wb.sheetnames:
            continue
        ws = wb[oef['blad']]
        for regel in oef['regels']:
            plekken = cellen_van(regel['waar'])
            if plekken is None:
                eis(regel['hoe'] == o.TEKENEN,
                    f"{oef['blad']} {regel['waar']}: geen celadres en toch geen "
                    'grafiek')
                continue
            for kolom, rij in plekken:
                cel = ws.cell(row=rij, column=kolom)
                eis(cel.value is None,
                    f"{oef['blad']}!{get_column_letter(kolom)}{rij} moet leeg zijn "
                    f"in het startbestand (de leerling maakt ze), maar bevat "
                    f"{cel.value!r}")
                tel('te maken cellen leeg')

    # de gegeven gegevens staan er echt
    ws = wb['1 Sultans']
    for i, (code, naam, begin, einde) in enumerate(g.B1):
        rij = g.B1_START + i
        eis(ws.cell(row=rij, column=1).value == code, f'1 Sultans A{rij}')
        eis(ws.cell(row=rij, column=2).value == naam, f'1 Sultans B{rij}')
        eis(ws.cell(row=rij, column=3).value == begin, f'1 Sultans C{rij}')
        eis(ws.cell(row=rij, column=4).value == einde, f'1 Sultans D{rij}')
        tel('gegeven cellen nagekeken', 4)
    eis(ws.cell(row=g.B1_KOPRIJ - 1, column=2).value == g.B1_DREMPEL,
        'de drempel staat niet in B2 van blad 1')

    ws = wb['2 Gebeurtenissen']
    for i, (jaar, naam, soort, regio, code) in enumerate(g.B2):
        rij = g.B2_START + i
        for kolom, verwacht in ((1, jaar), (2, naam), (3, soort), (4, regio),
                                (5, g.eeuw(jaar)), (6, code)):
            eis(ws.cell(row=rij, column=kolom).value == verwacht,
                f'2 Gebeurtenissen {get_column_letter(kolom)}{rij}')
            tel('gegeven cellen nagekeken')

    ws = wb['3 Oorlogen']
    for i, (naam, tegen, begin, einde, uitkomst) in enumerate(g.B3):
        rij = g.B3_START + i
        for kolom, verwacht in ((1, naam), (2, tegen), (3, begin), (4, einde),
                                (6, uitkomst)):
            eis(ws.cell(row=rij, column=kolom).value == verwacht,
                f'3 Oorlogen {get_column_letter(kolom)}{rij}')
            tel('gegeven cellen nagekeken')

    ws = wb['5 Grafiek']
    for i, (jaar, inwoners) in enumerate(g.B5):
        rij = g.B5_START + i
        eis(ws.cell(row=rij, column=1).value == jaar, f'5 Grafiek A{rij}')
        eis(ws.cell(row=rij, column=2).value == inwoners, f'5 Grafiek B{rij}')
        tel('gegeven cellen nagekeken', 2)

    # ------------------------------------------------------------------
    #  4. de twee bronlijsten zijn geschikt voor een draaitabel
    # ------------------------------------------------------------------
    for blad, koprij, start, eind, koppen, vraagkol in (
            ('2 Gebeurtenissen', g.B2_KOPRIJ, g.B2_START, g.B2_EIND, g.B2_KOP,
             g.B2_VRAAGKOL),
            ('3 Oorlogen', g.B3_KOPRIJ, g.B3_START, g.B3_EIND, g.B3_KOP,
             g.B3_VRAAGKOL)):
        ws = wb[blad]
        breedte = len(koppen)
        for i, titel in enumerate(koppen, start=1):
            eis(ws.cell(row=koprij, column=i).value == titel,
                f'{blad}: koprij mist of verschilt in kolom {get_column_letter(i)}')
        eis(len(set(koppen)) == breedte, f'{blad}: twee kolommen heten hetzelfde')

        # geen lege rij in de lijst: elke rij heeft minstens één gegeven cel
        for rij in range(start, eind + 1):
            gevuld = [c for c in range(1, breedte + 1)
                      if ws.cell(row=rij, column=c).value is not None]
            eis(gevuld, f'{blad}: rij {rij} is helemaal leeg — een draaitabel '
                        'stopt daar')
        # geen lege kolom in de lijst: elke kolom is gevuld of wordt gemaakt
        temaken = set()
        for oef in o.OEFENINGEN:
            if oef['blad'] != blad:
                continue
            for regel in oef['regels']:
                for kolom, _ in (cellen_van(regel['waar']) or []):
                    temaken.add(kolom)
        for kolom in range(1, breedte + 1):
            gevuld = any(ws.cell(row=r, column=kolom).value is not None
                         for r in range(start, eind + 1))
            eis(gevuld or kolom in temaken,
                f'{blad}: kolom {get_column_letter(kolom)} blijft leeg — dat '
                'breekt de draaitabel')
        # niets direct onder de lijst, en een lege kolom voor het vragenblokje
        for kolom in range(1, breedte + 1):
            eis(ws.cell(row=eind + 1, column=kolom).value is None,
                f'{blad}: er staat iets op rij {eind + 1}, direct onder de lijst. '
                'Een draaitabel zuigt dat mee naar binnen.')
        eis(vraagkol == breedte + 2,
            f'{blad}: het vragenblokje staat in kolom {get_column_letter(vraagkol)}, '
            f'maar de lijst loopt tot {get_column_letter(breedte)} — er hoort '
            'precies één lege kolom tussen')
        for rij in range(start, eind + 1):
            eis(ws.cell(row=rij, column=breedte + 1).value is None,
                f'{blad}: kolom {get_column_letter(breedte + 1)} rij {rij} hoort '
                'leeg te blijven als scheiding')
        tel('bronlijsten gekeurd')

    # het ankerpunt van elke draaitabel moet vrij zijn
    ws = wb['4 Draaitabel']
    for draai in (g.B4_DRAAI1, g.B4_DRAAI2):
        kolom, rij = cellen_van(draai['anker'])[0]
        eis(ws.cell(row=rij, column=kolom).value is None,
            f"4 Draaitabel: het anker {draai['anker']} is niet vrij")
        eis(draai['bron'] in wb.sheetnames,
            f"4 Draaitabel: onbekende bron {draai['bron']}")
        eis(draai['rijen'] in (g.B2_KOP if draai['bron'] == '2 Gebeurtenissen'
                               else g.B3_KOP),
            f"4 Draaitabel: {draai['rijen']} is geen kolom van {draai['bron']}")
        eis(draai['waarden'] in (g.B2_KOP if draai['bron'] == '2 Gebeurtenissen'
                                 else g.B3_KOP),
            f"4 Draaitabel: {draai['waarden']} is geen kolom van {draai['bron']}")
        if draai['kolommen']:
            eis(draai['kolommen'] in (g.B2_KOP if draai['bron'] == '2 Gebeurtenissen'
                                      else g.B3_KOP),
                f"4 Draaitabel: {draai['kolommen']} is geen kolom van "
                f"{draai['bron']}")
        tel('draaitabellen gekeurd')

    # de twee ankers mogen elkaar niet in de weg zitten
    r1 = cellen_van(g.B4_DRAAI1['anker'])[0][1]
    r2 = cellen_van(g.B4_DRAAI2['anker'])[0][1]
    nodig = len(a.D1_RIJLABELS) + 3
    eis(r2 - r1 >= nodig,
        f'de twee draaitabellen staan te dicht bij elkaar: draaitabel 1 heeft '
        f'{nodig} rijen nodig, er zitten er {r2 - r1} tussen de ankers')

# ==========================================================================
#  5. de opdracht zelf
# ==========================================================================
BASISFUNCTIES = {'SOM', 'GEMIDDELDE', 'MAX', 'MIN', 'AANTAL', 'AANTAL.ALS',
                 'SOM.ALS', 'ALS', 'VERT.ZOEKEN', 'AFRONDEN'}
NIEUW = {o.SLEPEN, o.LEZEN, o.TEKENEN, o.GEEN_FUNCTIE}
FUNCTIE = re.compile(r'\b([A-Z]+(?:\.[A-Z]+)*)\s*\(')

for oef in o.OEFENINGEN:
    for regel in oef['regels']:
        hoe = regel['hoe']
        eis(hoe in BASISFUNCTIES or hoe in NIEUW,
            f"{oef['blad']} {regel['waar']}: '{hoe}' is geen bekende functie en "
            'geen van de nieuwe werkwijzen')
        if hoe in BASISFUNCTIES:
            eis(regel.get('formule'),
                f"{oef['blad']} {regel['waar']}: functie {hoe} zonder formule in "
                'de sleutel')
            eis(regel['formule'] and hoe in regel['formule'],
                f"{oef['blad']} {regel['waar']}: de sleutelformule gebruikt {hoe} "
                'niet')
        if hoe in (o.LEZEN,):
            eis(regel.get('antwoord') is not None,
                f"{oef['blad']} {regel['waar']}: leesvraag zonder antwoord")
        tel('opdrachtregels gekeurd')

    # geen enkele formule mag een functie in een functie bevatten
    for regel in oef['regels']:
        f = regel.get('formule')
        if not f:
            continue
        diepte, binnenin = 0, []
        for stuk in re.finditer(r'[A-Z.]+\(|\(|\)', f):
            tekst = stuk.group()
            if tekst.endswith('('):
                if diepte > 0 and len(tekst) > 1:
                    binnenin.append(tekst[:-1])
                diepte += 1
            else:
                diepte -= 1
        eis(not binnenin,
            f"{oef['blad']} {regel['waar']}: geneste functie {binnenin} in "
            f"{f} — deze bundel hoort er geen te bevatten")

# een antwoord dat ze van hun draaitabel moeten AFLEZEN, mag nergens anders in
# de opdracht al uitgerekend worden — anders is blad 4 een invuloefening
# Het gaat om de CATEGORIE die ze moeten vinden (de drukste eeuw, de soort die
# het vaakst voorkomt). Staat die al als criterium in een formule van blad 1-3,
# dan heeft de draaitabel het antwoord al weggegeven. Een getal dat toevallig
# ook elders uitkomt, is geen verklapping — daarom alleen de tekstantwoorden.
_criteria = {}
for oef in o.OEFENINGEN:
    for r in oef['regels']:
        for literal in re.findall(r'"([^"]*)"', r.get('formule') or ''):
            _criteria.setdefault(literal, []).append(f"{oef['blad']} {r['waar']}")
for oef in o.OEFENINGEN:
    for r in oef['regels']:
        if r['hoe'] != o.LEZEN or not isinstance(r.get('antwoord'), str):
            continue
        eis(r['antwoord'] not in _criteria,
            f"{oef['blad']} {r['waar']}: \"{r['antwoord']}\" staat al als "
            f"criterium in {_criteria.get(r['antwoord'])} — dan hoeft de "
            'draaitabel er niet meer aan te pas te komen')
        tel('leesantwoorden afgeschermd')

eis(sum(int(oef['duur'].split()[0]) for oef in o.OEFENINGEN) == g.DUUR,
    'de duur van de bladen telt niet op tot de aangekondigde duur')

# de controlegetallen in de opdracht moeten echt uit de antwoorden komen
controles = ' '.join(oef['controle'] for oef in o.OEFENINGEN)
for getal in (a.B1_SOM, a.B1_LANGE, a.B2_TEL_EEUW, a.B2_TEL_REGIO, a.B3_SOM,
              a.D1_EINDTOTAAL, a.D2_EINDTOTAAL):
    eis(str(getal) in controles,
        f'het controlegetal {getal} staat in geen enkele controlezin')

# ==========================================================================
#  verslag
# ==========================================================================
print('=' * 72)
for naam, hoeveel in sorted(GETELD.items()):
    print(f'  {hoeveel:>5}  {naam}')
print('=' * 72)
if FOUTEN:
    print(f'{len(FOUTEN)} PROBLEMEN:')
    for f in FOUTEN:
        print('  !', f)
    sys.exit(1)
print('alles past bij elkaar.')
