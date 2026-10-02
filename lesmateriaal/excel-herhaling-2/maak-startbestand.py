# -*- coding: utf-8 -*-
"""Bouwt het startbestand van de tweede herhalingsbundel. Bewust kaal: de
   opmaak hoort bij de opdracht. Alleen de kolommen staan breed genoeg."""
import importlib.util
import sys

sys.dont_write_bytecode = True   # geen __pycache__ naast het lesmateriaal
import os
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HIER = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location('g', os.path.join(HIER, 'gegevens.py'))
g = importlib.util.module_from_spec(spec); spec.loader.exec_module(g)

UIT = os.path.join(HIER, 'excel-herhaling-2-startbestand.xlsx')
TITELFONT = Font(bold=True, size=13)

wb = Workbook()


def breedtes(ws, waarden):
    for i, b in enumerate(waarden, start=1):
        ws.column_dimensions[get_column_letter(i)].width = b


def koprij(ws, rij, koppen, start_kolom=1):
    for i, k in enumerate(koppen, start=start_kolom):
        ws.cell(row=rij, column=i, value=k)


def blad(naam, titel):
    ws = wb.create_sheet(naam)
    ws['A1'] = titel
    ws['A1'].font = TITELFONT
    return ws


# ------------------------------------------------------------------ Start
ws = wb.active
ws.title = 'Start'
ws['A1'] = g.TITEL
ws['A1'].font = Font(bold=True, size=16)
for i, r in enumerate([
    '',
    g.BEDRIJF,
    '',
    'Zeven korte oefeningen, samen ongeveer 40 minuten. Reken op vijf à zes',
    'minuten per blad. Werk ze in volgorde af.',
    '',
    'Dit is de tweede herhalingsbundel. Dezelfde tien functies, dezelfde soort',
    'opmaak, een ander bedrijf en andere cijfers. Herken je de opbouw van de',
    'vorige keer? Dat is de bedoeling.',
    '',
    'Loop je vast, pak er dan de functiefiches bij die je al hebt — daar staat',
    'per functie hoe ze werkt, met een voorbeeld en een schermbeeld.',
    '',
    'Blad 1 Weekomzet       — optellen in twee richtingen',
    'Blad 2 Kwekers         — een reeks samenvatten',
    'Blad 3 Prijsverhoging  — een percentage toepassen en afronden',
    'Blad 4 Bezoekers       — hoogste, laagste en gemiddelde',
    'Blad 5 Snoeicursus     — laten beslissen en daarna tellen',
    'Blad 6 Leveranciers    — tellen en optellen per categorie',
    'Blad 7 Zaden           — gegevens opzoeken en doorrekenen',
    '',
    'Sla het bestand eerst op onder je eigen naam.',
    'Reken nooit iets uit met je rekenmachine om het daarna over te typen.',
], start=2):
    ws.cell(row=i, column=1, value=r)
breedtes(ws, [88])

# ------------------------------------------------------------ 1 Weekomzet
ws = blad('1 Weekomzet', 'Omzet per dag')
koprij(ws, 3, g.B1_KOP)
for r, (dag, planten, meubels, gereedschap) in enumerate(g.B1, start=g.B1_START):
    ws.cell(row=r, column=1, value=dag)
    ws.cell(row=r, column=2, value=planten)
    ws.cell(row=r, column=3, value=meubels)
    ws.cell(row=r, column=4, value=gereedschap)
ws.cell(row=g.B1_TOTAAL, column=1, value='Totaal')
breedtes(ws, [14, 13, 14, 14, 14])

# -------------------------------------------------------------- 2 Kwekers
ws = blad('2 Kwekers', 'Geleverde planten per kweker')
koprij(ws, 3, g.B2_KOP)
for r, (naam, april, mei, juni) in enumerate(g.B2, start=g.B2_START):
    ws.cell(row=r, column=1, value=naam)
    ws.cell(row=r, column=2, value=april)
    ws.cell(row=r, column=3, value=mei)
    ws.cell(row=r, column=4, value=juni)
for i, label in enumerate([
    'Totaal aantal planten', 'Gemiddelde per kweker',
    'Hoogste totaal', 'Laagste totaal', 'Aantal kwekers',
    'Aantal grote kwekers',
]):
    ws.cell(row=g.B2_SAMENVATTING + i, column=1, value=label)
breedtes(ws, [26, 11, 11, 11, 12, 15])

# ------------------------------------------------------- 3 Prijsverhoging
ws = blad('3 Prijsverhoging', 'Prijslijst aanpassen')
ws['A2'] = 'Verhoging'
ws['B2'] = g.B3_PERCENTAGE
koprij(ws, 4, g.B3_KOP)
for r, (artikel, prijs) in enumerate(g.B3, start=g.B3_START):
    ws.cell(row=r, column=1, value=artikel)
    ws.cell(row=r, column=2, value=prijs)
ws.cell(row=g.B3_TOTAAL, column=1, value='Totaal')
breedtes(ws, [28, 13, 13, 15])

# ------------------------------------------------------------ 4 Bezoekers
ws = blad('4 Bezoekers', 'Bezoekers per maand')
koprij(ws, 3, g.B4_KOP)
for r, (maand, bezoekers) in enumerate(g.B4, start=g.B4_START):
    ws.cell(row=r, column=1, value=maand)
    ws.cell(row=r, column=2, value=bezoekers)
for i, label in enumerate([
    'Totaal aantal bezoekers', 'Gemiddeld per maand',
    'Drukste maand', 'Rustigste maand', 'Aantal maanden',
    f'Maanden boven {g.B4_DREMPEL} bezoekers',
]):
    ws.cell(row=g.B4_SAMENVATTING + i, column=1, value=label)
breedtes(ws, [26, 14])

# ---------------------------------------------------------- 5 Snoeicursus
ws = blad('5 Snoeicursus', 'Resultaten snoeicursus')
ws['A2'] = 'Geslaagd vanaf'
ws['B2'] = g.B5_GRENS
koprij(ws, 4, g.B5_KOP)
for r, (naam, punten) in enumerate(g.B5, start=g.B5_START):
    ws.cell(row=r, column=1, value=naam)
    ws.cell(row=r, column=2, value=punten)
for i, label in enumerate([
    'Aantal geslaagd', 'Aantal niet geslaagd', 'Gemiddelde punten',
]):
    ws.cell(row=g.B5_ANTWOORD + i, column=1, value=label)
breedtes(ws, [24, 11, 16])

# --------------------------------------------------------- 6 Leveranciers
ws = blad('6 Leveranciers', 'Bestellingen per leverancier')
koprij(ws, 3, g.B6_KOP)
for r, (bon, leverancier, bedrag) in enumerate(g.B6, start=g.B6_START):
    ws.cell(row=r, column=1, value=bon)
    ws.cell(row=r, column=2, value=leverancier)
    ws.cell(row=r, column=3, value=bedrag)
# De drie leveranciersnamen staan er al: die dienen als criterium.
ws.cell(row=g.B6_SAMENVATTING - 1, column=1, value='Leverancier')
ws.cell(row=g.B6_SAMENVATTING - 1, column=2, value='Aantal bonnen')
ws.cell(row=g.B6_SAMENVATTING - 1, column=3, value='Totaal bedrag')
for i, leverancier in enumerate(g.B6_LEVERANCIERS):
    ws.cell(row=g.B6_SAMENVATTING + i, column=1, value=leverancier)
ws.cell(row=g.B6_CONTROLE, column=1, value='Controle: alle bonnen samen')
breedtes(ws, [24, 16, 15])

# ---------------------------------------------------------------- 7 Zaden
ws = blad('7 Zaden', 'Zaadbestelling')
koprij(ws, 3, g.B7_KOP)
for r, (code, aantal) in enumerate(g.B7_REGELS, start=g.B7_START):
    ws.cell(row=r, column=1, value=code)
    ws.cell(row=r, column=4, value=aantal)
ws.cell(row=g.B7_TOTAAL, column=4, value='Totaal')
ws.cell(row=g.B7_DEEL, column=2, value=f'Waarvan bloemenmengsel ({g.B7_DEELTOTAAL})')

ws['G2'] = 'Zoektabel — niet aanpassen'
ws['G2'].font = Font(bold=True)
koprij(ws, 3, g.B7_ZOEKKOP, start_kolom=7)
for r, (code, oms, prijs) in enumerate(g.B7_TABEL, start=g.B7_TAB_START):
    ws.cell(row=r, column=7, value=code)
    ws.cell(row=r, column=8, value=oms)
    ws.cell(row=r, column=9, value=prijs)
breedtes(ws, [10, 24, 11, 10, 13, 4, 10, 24, 11])

wb.save(UIT)
print('geschreven:', UIT)
print('bladen:', wb.sheetnames)
for nr, (naam, start, eind) in enumerate([
    ('1 Weekomzet', g.B1_START, g.B1_EIND),
    ('2 Kwekers', g.B2_START, g.B2_EIND),
    ('3 Prijsverhoging', g.B3_START, g.B3_EIND),
    ('4 Bezoekers', g.B4_START, g.B4_EIND),
    ('5 Snoeicursus', g.B5_START, g.B5_EIND),
    ('6 Leveranciers', g.B6_START, g.B6_EIND),
    ('7 Zaden', g.B7_START, g.B7_EIND),
], start=1):
    print(f'  blad {nr} {naam:18s} gegevens rij {start}-{eind}')
