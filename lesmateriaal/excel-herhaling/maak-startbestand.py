# -*- coding: utf-8 -*-
"""Bouwt het startbestand van de herhalingsbundel. Bewust kaal: de opmaak
   hoort bij de opdracht. Alleen de kolommen staan breed genoeg."""
import importlib.util
import sys

sys.dont_write_bytecode = True   # geen __pycache__ naast het lesmateriaal
import os
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HIER = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location('g', os.path.join(HIER, 'gegevens.py'))
g = importlib.util.module_from_spec(spec); spec.loader.exec_module(g)

UIT = os.path.join(HIER, 'excel-herhaling-startbestand.xlsx')
TITELFONT = Font(bold=True, size=13)

wb = Workbook()


def breedtes(ws, waarden):
    for i, b in enumerate(waarden, start=1):
        ws.column_dimensions[get_column_letter(i)].width = b


def koprij(ws, rij, koppen, start_kolom=1):
    for i, k in enumerate(koppen, start=start_kolom):
        ws.cell(row=rij, column=i, value=k)


def blad(naam, titel):
    ws = wb.create_sheet(naam)
    ws['A1'] = titel
    ws['A1'].font = TITELFONT
    return ws


# ------------------------------------------------------------------ Start
ws = wb.active
ws.title = 'Start'
ws['A1'] = g.TITEL
ws['A1'].font = Font(bold=True, size=16)
for i, r in enumerate([
    '',
    g.BEDRIJF,
    '',
    'Zeven korte oefeningen, samen ongeveer 40 minuten. Reken op vijf à zes',
    'minuten per blad. Werk ze in volgorde af.',
    '',
    'Dit is een herhalingsbundel: je kent alle functies al uit de vorige bundel.',
    'Loop je vast, pak er dan de functiefiches van toen bij — dezelfde tien',
    'functies en dezelfde vijf hulpfiches komen hier terug.',
    '',
    'Blad 1 Dagontvangsten  — optellen in twee richtingen',
    'Blad 2 Werkuren        — een reeks samenvatten',
    'Blad 3 Kortingen       — een percentage toepassen en afronden',
    'Blad 4 Verbruik        — hoogste, laagste en gemiddelde',
    'Blad 5 Inschrijvingen  — laten beslissen en daarna tellen',
    'Blad 6 Afdelingen      — tellen en optellen per categorie',
    'Blad 7 Materiaal       — gegevens opzoeken en doorrekenen',
    '',
    'Sla het bestand eerst op onder je eigen naam.',
    'Reken nooit iets uit met je rekenmachine om het daarna over te typen.',
], start=2):
    ws.cell(row=i, column=1, value=r)
breedtes(ws, [88])

# ------------------------------------------------------- 1 Dagontvangsten
ws = blad('1 Dagontvangsten', 'Ontvangsten per dag')
koprij(ws, 3, g.B1_KOP)
for r, (dag, contant, kaart, online) in enumerate(g.B1, start=g.B1_START):
    ws.cell(row=r, column=1, value=dag)
    ws.cell(row=r, column=2, value=contant)
    ws.cell(row=r, column=3, value=kaart)
    ws.cell(row=r, column=4, value=online)
ws.cell(row=g.B1_TOTAAL, column=1, value='Totaal')
breedtes(ws, [14, 13, 14, 12, 14])

# ------------------------------------------------------- 2 Werkuren
ws = blad('2 Werkuren', 'Gewerkte uren per medewerker')
koprij(ws, 3, g.B2_KOP)
for r, (naam, w1, w2, w3) in enumerate(g.B2, start=g.B2_START):
    ws.cell(row=r, column=1, value=naam)
    ws.cell(row=r, column=2, value=w1)
    ws.cell(row=r, column=3, value=w2)
    ws.cell(row=r, column=4, value=w3)
for i, label in enumerate([
    'Totaal aantal uren', 'Gemiddelde per medewerker',
    'Hoogste totaal', 'Laagste totaal', 'Aantal medewerkers',
    'Aantal voltijds',
]):
    ws.cell(row=g.B2_SAMENVATTING + i, column=1, value=label)
breedtes(ws, [26, 11, 11, 11, 12, 13])

# ------------------------------------------------------- 3 Kortingen
ws = blad('3 Kortingen', 'Kortingsactie gereedschap')
ws['A2'] = 'Korting'
ws['B2'] = g.B3_PERCENTAGE
koprij(ws, 4, g.B3_KOP)
for r, (artikel, prijs) in enumerate(g.B3, start=g.B3_START):
    ws.cell(row=r, column=1, value=artikel)
    ws.cell(row=r, column=2, value=prijs)
ws.cell(row=g.B3_TOTAAL, column=1, value='Totaal')
breedtes(ws, [26, 13, 13, 15])

# ------------------------------------------------------- 4 Verbruik
ws = blad('4 Verbruik', 'Elektriciteitsverbruik per maand')
koprij(ws, 3, g.B4_KOP)
for r, (maand, verbruik) in enumerate(g.B4, start=g.B4_START):
    ws.cell(row=r, column=1, value=maand)
    ws.cell(row=r, column=2, value=verbruik)
for i, label in enumerate([
    'Totaal verbruik', 'Gemiddeld per maand',
    'Hoogste maand', 'Laagste maand', 'Aantal maanden',
    f'Maanden boven {g.B4_DREMPEL} kWh',
]):
    ws.cell(row=g.B4_SAMENVATTING + i, column=1, value=label)
breedtes(ws, [22, 16])

# ------------------------------------------------------- 5 Inschrijvingen
ws = blad('5 Inschrijvingen', 'Inschrijvingen workshop')
ws['A2'] = 'Grens volwassene'
ws['B2'] = g.B5_GRENS
koprij(ws, 4, g.B5_KOP)
for r, (naam, leeftijd) in enumerate(g.B5, start=g.B5_START):
    ws.cell(row=r, column=1, value=naam)
    ws.cell(row=r, column=2, value=leeftijd)
for i, label in enumerate([
    'Aantal volwassenen', 'Aantal jeugd', 'Gemiddelde leeftijd',
]):
    ws.cell(row=g.B5_ANTWOORD + i, column=1, value=label)
breedtes(ws, [24, 11, 16])

# ------------------------------------------------------- 6 Afdelingen
ws = blad('6 Afdelingen', 'Verkoop per afdeling')
koprij(ws, 3, g.B6_KOP)
for r, (bon, afdeling, bedrag) in enumerate(g.B6, start=g.B6_START):
    ws.cell(row=r, column=1, value=bon)
    ws.cell(row=r, column=2, value=afdeling)
    ws.cell(row=r, column=3, value=bedrag)
# De drie afdelingsnamen staan er al: die dienen als criterium in de formules.
ws.cell(row=g.B6_SAMENVATTING - 1, column=1, value='Afdeling')
ws.cell(row=g.B6_SAMENVATTING - 1, column=2, value='Aantal bonnen')
ws.cell(row=g.B6_SAMENVATTING - 1, column=3, value='Totaal bedrag')
for i, afdeling in enumerate(g.B6_AFDELINGEN):
    ws.cell(row=g.B6_SAMENVATTING + i, column=1, value=afdeling)
ws.cell(row=g.B6_CONTROLE, column=1, value='Controle: alle bonnen samen')
breedtes(ws, [22, 16, 15])

# ------------------------------------------------------- 7 Materiaal
ws = blad('7 Materiaal', 'Materiaalbestelling')
koprij(ws, 3, g.B7_KOP)
for r, (code, aantal) in enumerate(g.B7_REGELS, start=g.B7_START):
    ws.cell(row=r, column=1, value=code)
    ws.cell(row=r, column=4, value=aantal)
ws.cell(row=g.B7_TOTAAL, column=4, value='Totaal')
ws.cell(row=g.B7_TOTAAL + 2, column=2, value='Waarvan cement (M-01)')

ws['G2'] = 'Zoektabel — niet aanpassen'
ws['G2'].font = Font(bold=True)
koprij(ws, 3, g.B7_ZOEKKOP, start_kolom=7)
for r, (code, oms, prijs) in enumerate(g.B7_TABEL, start=g.B7_TAB_START):
    ws.cell(row=r, column=7, value=code)
    ws.cell(row=r, column=8, value=oms)
    ws.cell(row=r, column=9, value=prijs)
breedtes(ws, [10, 24, 11, 10, 13, 4, 10, 24, 11])

wb.save(UIT)
print('geschreven:', UIT)
print('bladen:', wb.sheetnames)
for nr, (naam, start, eind) in enumerate([
    ('1 Dagontvangsten', g.B1_START, g.B1_EIND),
    ('2 Werkuren', g.B2_START, g.B2_EIND),
    ('3 Kortingen', g.B3_START, g.B3_EIND),
    ('4 Verbruik', g.B4_START, g.B4_EIND),
    ('5 Inschrijvingen', g.B5_START, g.B5_EIND),
    ('6 Afdelingen', g.B6_START, g.B6_EIND),
    ('7 Materiaal', g.B7_START, g.B7_EIND),
], start=1):
    print(f'  blad {nr} {naam:18s} gegevens rij {start}-{eind}')
