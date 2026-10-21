# -*- coding: utf-8 -*-
"""Bouwt de opdrachtbundel en de lerarensleutel van de Minecraft-bundel.

Het raamwerk is hetzelfde als in de vorige bundels. Nieuw is wat er NIET
in staat: geen theorie. Elk blad verwijst naar de kaarten.
"""
import importlib.util
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')

OPDRACHT = os.path.join(HIER, 'minecraft-draaitabellen-opdracht.docx')
SLEUTEL = os.path.join(HIER, 'minecraft-draaitabellen-lerarensleutel.docx')

BLAUW = RGBColor(0x1F, 0x38, 0x64)
ORANJE = RGBColor(0xB0, 0x4A, 0x00)
GRIJS = RGBColor(0x59, 0x59, 0x59)
GROEN = RGBColor(0x38, 0x76, 0x1D)

FORMULE = re.compile(r"(=[A-Z.]*\((?:[^()]|\([^()]*\))*\)|=[A-Z$]+\d+(?:[*+\-/]"
                     r"[A-Z$]*\d+)*)")


def opzet(doc):
    s = doc.styles['Normal']
    s.font.name = 'Calibri'
    s.font.size = Pt(11)
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(1.5)
        sec.left_margin = sec.right_margin = Cm(1.8)


def kop(doc, tekst, niveau=1, kleur=BLAUW):
    p = doc.add_heading(tekst, level=niveau)
    for r in p.runs:
        r.font.color.rgb = kleur
    return p


def zin(p, tekst):
    """Zet formules in een vaste letter, zodat ze opvallen in lopende tekst."""
    for stuk in FORMULE.split(str(tekst)):
        if not stuk:
            continue
        r = p.add_run(stuk)
        if FORMULE.fullmatch(stuk):
            r.font.name = 'Consolas'
            r.font.size = Pt(10)
    return p


def alinea(doc, tekst, style=None, vet=False, schuin=False, kleur=None,
           grootte=None, na=None):
    p = doc.add_paragraph(style=style)
    if vet or schuin or kleur or grootte:
        r = p.add_run(str(tekst))
        r.bold, r.italic = vet, schuin
        if kleur:
            r.font.color.rgb = kleur
        if grootte:
            r.font.size = Pt(grootte)
    else:
        zin(p, tekst)
    if na is not None:
        p.paragraph_format.space_after = Pt(na)
    return p


def kader(doc, titel, regels, kleur=BLAUW):
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    c = t.rows[0].cells[0]
    r = c.paragraphs[0].add_run(titel)
    r.bold = True
    r.font.color.rgb = kleur
    for regel in regels:
        pp = c.add_paragraph()
        pp.paragraph_format.space_after = Pt(2)
        zin(pp, regel)
    doc.add_paragraph()
    return t


def tabel(doc, koppen, rijen, stijl='Light Grid Accent 1', breedtes=None):
    t = doc.add_table(rows=1, cols=len(koppen))
    t.style = stijl
    for i, k in enumerate(koppen):
        t.rows[0].cells[i].paragraphs[0].add_run(str(k)).bold = True
    for rij in rijen:
        cellen = t.add_row().cells
        for i, w in enumerate(rij):
            zin(cellen[i].paragraphs[0], w)
    if breedtes:
        for rij in t.rows:
            for i, b in enumerate(breedtes):
                rij.cells[i].width = Cm(b)
    doc.add_paragraph()
    return t


def paginaeinde(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def genummerd(doc, regels, kleur=BLAUW):
    for i, regel in enumerate(regels, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f'{i}.  ')
        r.bold = True
        r.font.color.rgb = kleur
        zin(p, regel)



def kaartregel(doc, oef):
    """De verwijzing naar de theorie. Hier staat de theorie zelf niet."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run('Theorie:  ')
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = GRIJS
    stukken = [f"kaart {n} — {o.KAARTEN[n]}" for n in oef['kaarten']]
    r = p.add_run(' · '.join(stukken))
    r.font.size = Pt(10)
    r.font.color.rgb = ORANJE
    return p


def draaispec(draai):
    """De draaitabel als rijtjes voor een tabel in het document."""
    rijen = [['Rijen', draai['rijen']]]
    if draai['kolommen']:
        rijen.append(['Kolommen', draai['kolommen']])
    waarde = f"{draai['waarden']}  als  {draai['bewerking']}"
    if draai.get('tweede'):
        waarde = f"Aantal van {draai['tweede']}   én   " + waarde
    rijen.append(['Waarden', waarde])
    if draai.get('groeperen'):
        rijen.append(['Groeperen', draai['groeperen']])
    if draai.get('sorteren'):
        rijen.append(['Sorteren', f"{draai['sorteren']} op de waarden"])
    rijen.append(['Waar', f"linkerbovenhoek in {draai['anker']}"])
    return rijen


# ==========================================================================
#  1. de opdrachtbundel
# ==========================================================================
doc = Document()
opzet(doc)
kop(doc, g.TITEL, 0)
alinea(doc, g.ONDERTITEL, schuin=True, kleur=GRIJS, na=4)
alinea(doc, 'Naam: ......................................    Klas: ..........',
       na=14)

alinea(doc,
       f'Vijf werkbladen, samen {g.DUUR} minuten. Op elk blad staat een lijst klaar '
       'en maak je er één draaitabel van. Het materiaal verschilt telkens; de '
       'techniek is vijf keer dezelfde.', na=8)
alinea(doc,
       'In deze bundel staat geen theorie. Bij elk blad staat welke kaarten uit '
       f'{g.THEORIEBUNDEL} je erbij nodig hebt — elke kaart staat daar op een eigen '
       'pagina. Loop je vast, dan pak je die kaart erbij.', na=12)

tabel(doc, ['Blad', 'Wat je oefent', 'Theorie', 'Duur', 'Punten'],
      [[f"{oef['nr']}  {oef['blad'][2:]}", oef['titel'],
        ', '.join(f"kaart {n}" for n in oef['kaarten']), oef['duur'],
        oef['punten_draai'] + sum(p for *_, p in oef['vragen'])
        + sum(p for *_, p in oef.get('vullen', []))]
       for oef in o.OEFENINGEN],
      breedtes=[3.6, 5.4, 3.4, 2.6, 1.8])

kader(doc, 'Drie regels die op elk blad gelden', [
    '1.  De lijst blijft zoals ze is. Voeg er geen lege rij in, zet er niets onder '
    'en verwijder er niets uit. De vragen staan ERNAAST, met een lege kolom ertussen.',
    '2.  De draaitabel komt ONDER de lijst, op de plaats die op het blad aangeduid '
    'staat. Zo blijft de lijst een lijst.',
    '3.  Lees altijd het kopje van je draaitabel. Staat er Aantal waar Som of '
    'Gemiddelde hoort, dan klopt je hele tabel niet — en Excel waarschuwt daar niet '
    'voor.',
], ORANJE)

kader(doc, 'Over de cijfers', [g.WAARSCHUWING], GRIJS)

for oef in o.OEFENINGEN:
    paginaeinde(doc)
    kop(doc, f"Blad {oef['nr']} — {oef['titel']}", 1)
    alinea(doc, f"werkblad '{oef['blad']}'  ·  {oef['duur']}", schuin=True,
           kleur=GRIJS, na=4)
    kaartregel(doc, oef)
    for regel in oef['inleiding']:
        alinea(doc, regel, na=4)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    if oef.get('vullen'):
        kop(doc, 'Eerst: de lijst aanvullen', 2, GRIJS)
        tabel(doc, ['Waar', 'Wat', 'Kolom uit de beroepentabel', 'Pt'],
              [[waar, wat, kolom, punten]
               for waar, wat, kolom, _formule, punten in oef['vullen']],
              breedtes=[3.4, 6.4, 5.6, 1.4])
        alinea(doc,
               'Zet de beroepentabel vast met dollartekens voordat je doorvoert, '
               'anders zakt ze mee naar beneden en krijg je onderaan #N/B.', na=10)

    kop(doc, 'Dan: de draaitabel', 2, GRIJS)
    tabel(doc, ['Vak', 'Wat erin komt'], draaispec(oef['draai']),
          breedtes=[3.4, 13.0])

    if oef.get('extra'):
        e = oef['extra']
        naam, nummer, aanbiedingen = e['waarden']
        kop(doc, 'Daarna: één dorpeling erbij', 2, GRIJS)
        alinea(doc,
               f'Typ onderaan de lijst, op rij {e["rij"]}, deze dorpeling erbij:',
               na=4)
        tabel(doc, ['Dorpeling', 'Beroepcode', 'Aanbiedingen'],
              [[naam, g.beroepcode(nummer), aanbiedingen]],
              breedtes=[4.4, 4.0, 4.0])
        alinea(doc,
               'Vul daarna kolom D, E en F voor die rij ook aan, vernieuw je '
               'draaitabel (rechtsklik erin, Vernieuwen) en kijk of de dorpeling '
               'meegeteld is. Zo niet: je bronbereik loopt nog tot de oude laatste '
               'rij. Draaitabel analyseren, Gegevensbron wijzigen.', na=10)

    kop(doc, 'Lees af en vul in', 2, GRIJS)
    tabel(doc, ['Cel', 'Vraag', 'Pt'],
          [[cel, wat[0].upper() + wat[1:] + '?', punten]
           for cel, wat, _antwoord, punten in oef['vragen']],
          breedtes=[2.4, 12.6, 1.4])

    kader(doc, 'Je blad is af als…', [oef['klaar'], '',
                                      'CONTROLE  ·  ' + oef['controle']], GROEN)
    kader(doc, 'Loop je vast?', oef['vast'], ORANJE)

paginaeinde(doc)
kop(doc, 'Helemaal klaar?', 1)
alinea(doc, 'Loop dit lijstje af voor je indient.', schuin=True, kleur=GRIJS, na=8)
for regel in [
    'Op elk van de vijf bladen staat een draaitabel.',
    'Elk controlegetal klopt.',
    'Nergens staat #N/B of (leeg) in een draaitabel.',
    'Je hebt elke vraag ingevuld — ook die van blad 5 na het vernieuwen.',
    'Je lijsten zijn nog altijd lijsten: geen lege rij, niets eronder.',
    'Je bestand heet minecraft-draaitabellen-jouwnaam.xlsx.',
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(4)
    p.add_run('☐   ').font.size = Pt(13)
    zin(p, regel)

doc.save(OPDRACHT)
print('geschreven:', OPDRACHT)


# ==========================================================================
#  2. de lerarensleutel
# ==========================================================================
def draaitabel_af(doc, koppen, rijen, breedtes=None):
    tabel(doc, koppen, rijen, 'Light List Accent 1', breedtes)


doc = Document()
opzet(doc)
kop(doc, g.TITEL + ' — lerarensleutel', 0)
alinea(doc, g.ONDERTITEL, schuin=True, kleur=GRIJS, na=12)

kop(doc, 'In het kort', 2, GRIJS)
alinea(doc,
       f'Vijf bladen, {g.DUUR} minuten, {o.TOTAAL} punten. Geen enkele formule, '
       'behalve op blad 5 waar ze de lijst eerst moeten aanvullen met VERT.ZOEKEN. '
       'Er staat geen theorie in de opdracht: elk blad verwijst naar de kaarten '
       f'van {g.THEORIEBUNDEL}.', na=6)
alinea(doc,
       'De bladen lopen op in moeilijkheid: Rijen en Waarden, dan Kolommen erbij, '
       'dan Som in plaats van Aantal, dan groeperen en filteren, en ten slotte een '
       'lijst die eerst aangevuld moet worden. Wie blad 1 niet rond krijgt, hoeft '
       'aan blad 4 niet te beginnen.', na=6)
alinea(doc,
       'Dit is dezelfde bundel als de Türkiye-versie, met ander materiaal. Wie die '
       'al maakte, herkent elk blad en kan sneller. Wie ze niet maakte, merkt er '
       'niets van.', na=10)

tabel(doc, ['Blad', 'Wat', 'Kaarten', 'Duur', 'Punten'],
      [[f"{oef['nr']}  {oef['blad'][2:]}", oef['titel'],
        ', '.join(str(n) for n in oef['kaarten']), oef['duur'],
        oef['punten_draai'] + sum(p for *_, p in oef['vragen'])
        + sum(p for *_, p in oef.get('vullen', []))]
       for oef in o.OEFENINGEN]
      + [['Samen', '', '', f'{g.DUUR} minuten', o.TOTAAL]],
      breedtes=[3.6, 5.4, 2.6, 2.6, 2.2])

kader(doc, 'Een draaitabel kan niet in het startbestand zitten', [
    'De bibliotheek waarmee deze bestanden gemaakt worden, kan er geen schrijven — '
    'en dat komt goed uit, want de leerling hoort ze zelf te maken. Het betekent '
    'wel dat dit blad voor blad met de hand nagekeken moet worden. Hieronder staan '
    'alle vijf de draaitabellen volledig uitgeschreven, zodat je het blad van een '
    'leerling er gewoon naast kan leggen.',
    'Of ze de tabel op exact de juiste cel gezet hebben, is bijzaak. Of de velden '
    'in de juiste vakken zitten en de bewerking klopt, niet.',
], BLAUW)

kader(doc, 'Waar het bij iedereen misgaat', [
    '1.  Het waardeveld staat op Aantal waar Som of Gemiddelde hoort. Dat is '
    'blad 3 en blad 4 in één zin. Herkenbaar aan het eindtotaal: krijg je het '
    'aantal regels van de lijst, dan staat het op Aantal. Blad 3 is daar met '
    'opzet gemeen in: met Aantal geeft élke rij hetzelfde getal, en dat ziet er '
    'geloofwaardig genoeg uit om niet op te vallen.',
    '2.  Groeperen lukt niet omdat ze rechtsklikken op de verkeerde cel. Het moet '
    'op een rijlabel met een jaartal erin.',
    '3.  Op blad 5 vergeten ze de dollartekens rond de beroepentabel, of ze '
    'vergeten na het toevoegen van de extra rij het bronbereik aan te passen. '
    'Vernieuwen alleen volstaat daar niet — dat is precies de les van kaart 9.',
], ORANJE)

for oef in o.OEFENINGEN:
    paginaeinde(doc)
    punten = (oef['punten_draai'] + sum(p for *_, p in oef['vragen'])
              + sum(p for *_, p in oef.get('vullen', [])))
    kop(doc, f"Blad {oef['nr']} — {oef['titel']}", 1)
    alinea(doc, f"werkblad '{oef['blad']}'  ·  {oef['duur']}  ·  {punten} punten  ·  "
                f"kaart {', '.join(str(n) for n in oef['kaarten'])}",
           schuin=True, kleur=GRIJS, na=8)

    if oef.get('vullen'):
        kop(doc, 'De formules', 2, GRIJS)
        tabel(doc, ['Waar', 'Formule', 'Pt'],
              [[waar, formule, punten]
               for waar, _wat, _kol, formule, punten in oef['vullen']],
              breedtes=[3.0, 12.0, 1.4])

    kop(doc, 'De draaitabel', 2, GRIJS)
    tabel(doc, ['Vak', 'Wat erin komt'], draaispec(oef['draai']),
          breedtes=[3.0, 13.0])

    nr = oef['nr']
    if nr == 1:
        draaitabel_af(doc, ['Dimensie', 'Aantal mobs'],
                      [[d, a.D1[d]] for d in a.D1_LABELS]
                      + [['Eindtotaal', a.D1_TOTAAL]], [6.0, 5.0])
    elif nr == 2:
        draaitabel_af(doc, ['Aantal'] + a.D2_KOLLEN + ['Eindtotaal'],
                      [[r] + [a.D2_CEL[(r, k)] or '' for k in a.D2_KOLLEN]
                       + [a.D2_RIJTOT[r]] for r in a.D2_RIJEN]
                      + [['Eindtotaal'] + [a.D2_KOLTOT[k] for k in a.D2_KOLLEN]
                         + [a.D2_TOTAAL]])
        alinea(doc, f'{a.D2_LEGE_VAKJES} vakjes blijven leeg. Dat is geen fout: '
                    'die combinaties komen niet voor. De derde vraag gaat er zelfs '
                    f'over — {a.D2_EEN_GEREEDSCHAP} is de enige rij met maar één '
                    'gevuld vakje.', schuin=True, kleur=GRIJS, na=10)
    elif nr == 3:
        draaitabel_af(doc, ['Materiaal', 'Som van duurzaamheid'],
                      [[m, a._getal(a.D3[m])] for m in a.D3_AFLOPEND]
                      + [['Eindtotaal', a._getal(a.D3_TOTAAL)]], [6.0, 5.0])
        alinea(doc, f'Staat er {len(g.B3)} als eindtotaal, dan telt hun draaitabel '
                    'de stuks gereedschap in plaats van hun duurzaamheid — en dan '
                    f'staat er op élke rij {len(g.B3_SOORTEN)}.',
               schuin=True, kleur=GRIJS, na=4)
        alinea(doc, f'De pointe van dit blad: {a.D3_LAAGSTE} komt het laagst uit, '
                    f'lager nog dan {a.D3_AFLOPEND[-2]}. Dat weet elke speler, maar '
                    'nu staat het er zwart op wit.',
               schuin=True, kleur=GRIJS, na=10)
    elif nr == 4:
        draaitabel_af(doc, ['Periode', 'Aantal updates',
                            'Gemiddeld nieuwe blokken'],
                      [[d, a.D4_AANTAL[d], a._getal(a.D4_GEMIDDELD[d])]
                       for d in a.D4_LABELS]
                      + [['Eindtotaal', a.D4_TOTAAL_AANTAL,
                          a._getal(a.D4_TOTAAL_GEMIDDELD)]], [4.6, 4.2, 5.6])
        alinea(doc, f'Let op de tweede kolom tegenover de derde: {a.D4_DRUKSTE} had '
                    f'de meeste updates ({a.D4_DRUKSTE_AANTAL}), maar het laagste '
                    f'gemiddelde. {a.D4_HOOGSTE_GEM} heeft er het minst '
                    f'({a.D4_AANTAL[a.D4_HOOGSTE_GEM]}) en het hoogste gemiddelde. '
                    'Minder updates, grotere updates — die vraag is het waard om '
                    'klassikaal te stellen.', schuin=True, kleur=GRIJS, na=10)
    elif nr == 5:
        draaitabel_af(doc, ['Categorie', 'Som van de aanbiedingen',
                            'Na de extra dorpeling'],
                      [[c, a._getal(a.D5[c]), a._getal(a.D5_NA[c])]
                       for c in a.D5_AFLOPEND]
                      + [['Eindtotaal', a._getal(a.D5_TOTAAL),
                          a._getal(a.D5_NA_TOTAAL)]], [5.0, 5.4, 5.0])
        alinea(doc, f'De extra dorpeling is een {a.D5_EXTRA_BEROEP.lower()}, dus '
                    f'{oef["extra"]["categorie"]}. Alleen die regel en het '
                    'eindtotaal veranderen.', schuin=True, kleur=GRIJS, na=10)

    kop(doc, 'De antwoorden', 2, GRIJS)
    tabel(doc, ['Cel', 'Vraag', 'Antwoord', 'Pt'],
          [[cel, wat[0].upper() + wat[1:], a._getal(antwoord)
            if isinstance(antwoord, (int, float)) else antwoord, punten]
           for cel, wat, antwoord, punten in oef['vragen']],
          breedtes=[2.0, 8.6, 4.4, 1.4])

    kader(doc, 'Waar je naar kijkt', [oef['controle']], GROEN)

paginaeinde(doc)
kop(doc, 'Tijdsindeling en tips', 1)
tabel(doc, ['Wanneer', 'Wat'],
      [['0 - 4 min', 'Bestand opslaan onder eigen naam. Theoriebundel uitdelen. '
                     'Zeg welke gegevens kloppen en welke verzonnen zijn, en dat '
                     'ze alles mogen nakijken op de wiki.'],
       ['4 - 10 min', 'Blad 1. Dit is de proefrit; iedereen hoort hier door te '
                      'raken. Wie na zes minuten nog geen draaitabel heeft, help je '
                      'nu — anders haalt die de rest niet.'],
       ['10 - 26 min', 'Blad 2 en 3. Kruistabel en Som. Loop rond en kijk naar de '
                       'kopjes van hun draaitabellen.'],
       ['26 - 38 min', 'Blad 4. Groeperen is het nieuwste stuk; verwacht hier de '
                       'meeste vragen.'],
       ['38 - 50 min', 'Blad 5. Wie hier aankomt heeft de les gehad; wie blad 5 '
                       'afwerkt, begrijpt draaitabellen.']],
      breedtes=[3.4, 13.0])

kader(doc, 'Als je weinig tijd hebt', [
    'Laat blad 2 vallen. De kruistabel is leuk maar niet nodig om blad 3 tot 5 te '
    'kunnen; dat scheelt acht minuten.',
    'Blad 5 kan ook een apart lesje van twintig minuten worden. Het staat op '
    'zichzelf: de beroepentabel die het nodig heeft, staat op dat blad zelf.',
], BLAUW)

kader(doc, 'Wat je hierna kan doen', [
    'Elke lijst kan nog andere vragen beantwoorden. Laat wie snel klaar is een '
    'tweede draaitabel van dezelfde lijst maken met andere velden — op blad 1 '
    'bijvoorbeeld Categorie in Kolommen naast Dimensie in Rijen, of op blad 2 '
    'Gemiddelde van Hardheid in plaats van Aantal van Blok.',
    'De grafiekbundel gebruikt hetzelfde materiaal. Wie deze bundel af heeft, kan '
    'daar meteen aan beginnen.',
], BLAUW)

kader(doc, 'Over de gegevens, als er een discussie komt', [
    'Die discussie is winst, geen tijdverlies. De namen, de levenspunten, de '
    'duurzaamheid en de jaartallen komen uit de Java-editie en mogen nagekeken '
    'worden. De indeling van de blokken in vier categorieën is onze keuze: '
    'netherrack als natuurlijk blok in plaats van bouwblok is verdedigbaar, en '
    'als een leerling dat opmerkt heeft die precies begrepen waar het op blad 2 '
    'om draait. Een draaitabel is maar zo goed als de kolom waarop ze gebouwd is.',
    'Het aantal aanbiedingen van de dorpelingen en het aantal nieuwe blokken per '
    'update zijn verzonnen. Die staan er om mee te rekenen, niet om te citeren.',
], GRIJS)

doc.save(SLEUTEL)
print('geschreven:', SLEUTEL)
