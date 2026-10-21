# -*- coding: utf-8 -*-
"""Kijkt na of opdracht, startbestand, antwoorden en sleutel bij elkaar passen.

Faalt luid zodra er iets uit elkaar loopt. Geen vaste rijnummers: alles komt
uit gegevens.py.
"""
import collections
import importlib.util
import os
import re
import sys
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')

WERKBOEK = os.path.join(HIER, 'minecraft-draaitabellen-startbestand.xlsx')
FOUTEN, GETELD = [], collections.Counter()


def eis(voorwaarde, bericht):
    if not voorwaarde:
        FOUTEN.append(bericht)
    return bool(voorwaarde)


def cel_van(adres):
    m = re.match(r'^([A-Z]{1,2})(\d+)$', adres)
    return (column_index_from_string(m.group(1)), int(m.group(2))) if m else None


# ==========================================================================
#  1. de gegevens zijn intern samenhangend
# ==========================================================================
for mob, categorie, dimensie, levenspunten in g.B1:
    eis(categorie in g.B1_CATEGORIEEN, f'onbekende categorie {categorie!r} bij {mob}')
    eis(dimensie in g.DIMENSIES, f'onbekende dimensie {dimensie!r} bij {mob}')
    eis(levenspunten > 0, f'onmogelijke levenspunten bij {mob}')
    eis(levenspunten % 2 == 0 or levenspunten == 3,
        f'{mob}: {levenspunten} levenspunten is geen heel aantal hartjes')
    GETELD['mobs nagerekend'] += 1
eis(len({r[0] for r in g.B1}) == len(g.B1), 'dubbele mobnaam')

for blok, categorie, gereedschap, hardheid in g.B2:
    eis(categorie in g.B2_CATEGORIEEN, f'onbekende categorie {categorie!r} bij {blok}')
    eis(gereedschap in g.B2_GEREEDSCHAP,
        f'onbekend gereedschap {gereedschap!r} bij {blok}')
    eis(hardheid > 0, f'onmogelijke hardheid bij {blok}')
    GETELD['blokken nagerekend'] += 1
eis(len({r[0] for r in g.B2}) == len(g.B2), 'dubbel blok')

for naam, materiaal, soort, duurzaamheid in g.B3:
    eis(materiaal in g.B3_MATERIALEN, f'onbekend materiaal {materiaal!r} bij {naam}')
    eis(soort in g.B3_SOORTEN, f'onbekende soort {soort!r} bij {naam}')
    eis(duurzaamheid == g.B3_DUURZAAMHEID[materiaal],
        f'{naam}: duurzaamheid {duurzaamheid} hoort {g.B3_DUURZAAMHEID[materiaal]} '
        'te zijn')
    GETELD['gereedschap nagerekend'] += 1
eis(len({r[0] for r in g.B3}) == len(g.B3), 'dubbel stuk gereedschap')
# Het hele punt van blad 3 is Som tegenover Aantal. Dat werkt alleen als er van
# elk materiaal evenveel stuks zijn: dan geeft Aantal overal hetzelfde getal.
_per_materiaal = collections.Counter(r[1] for r in g.B3)
eis(len(set(_per_materiaal.values())) == 1,
    'niet van elk materiaal evenveel stuks — dan verdwijnt de les van blad 3, '
    f'want Aantal geeft dan toch verschillen: {dict(_per_materiaal)}')

for versie, naam, jaar, thema, blokken in g.B4:
    eis(thema in g.B4_THEMAS, f'onbekend hoofdthema {thema!r} bij {versie}')
    eis(g.B4_GROEP_VAN <= jaar <= g.B4_GROEP_TOT,
        f'{versie} ({jaar}) valt buiten het groepeerbereik '
        f'{g.B4_GROEP_VAN}-{g.B4_GROEP_TOT}')
    eis(blokken > 0, f'onmogelijk aantal nieuwe blokken bij {versie}')
    GETELD['updates nagerekend'] += 1
eis(len({r[0] for r in g.B4}) == len(g.B4), 'dubbel versienummer')
eis([r[2] for r in g.B4] == sorted(r[2] for r in g.B4),
    'de updates staan niet op jaartal — dat maakt het groeperen onnodig moeilijk')
eis((g.B4_GROEP_TOT - g.B4_GROEP_VAN + 1) % g.B4_GROEP_PER == 0,
    f'het bereik {g.B4_GROEP_VAN}-{g.B4_GROEP_TOT} gaat niet op in groepen van '
    f'{g.B4_GROEP_PER}')
eis(g.B4_FILTER in g.B4_THEMAS, f'er valt niet te filteren op {g.B4_FILTER!r}')

for beroep, werkblok, categorie in g.B5_BEROEPEN:
    eis(beroep and werkblok and categorie, f'onvolledig beroep: {beroep!r}')
    GETELD['beroepen nagerekend'] += 1
eis(len({r[0] for r in g.B5_BEROEPEN}) == len(g.B5_BEROEPEN), 'dubbel beroep')
eis(len({r[1] for r in g.B5_BEROEPEN}) == len(g.B5_BEROEPEN),
    'twee beroepen delen hetzelfde werkblok')

for naam, nummer, aanbiedingen in list(g.B5) + [g.B5_EXTRA]:
    eis(1 <= nummer <= len(g.B5_BEROEPEN),
        f'beroepnummer {nummer} bestaat niet ({naam})')
    eis(aanbiedingen > 0, f'aanbiedingen moeten positief zijn bij {naam}')
    GETELD['dorpelingen nagerekend'] += 1
eis(len({r[0] for r in g.B5}) == len(g.B5), 'dubbele dorpelingennaam')
eis(g.B5_EXTRA[0] not in {r[0] for r in g.B5},
    'de extra dorpeling staat al in de lijst')
# elke categorie van de beroepentabel moet ook echt in de dorpelingenlijst zitten,
# anders staat er een rij in de zoektabel die nergens toe leidt
_gebruikt = {g.beroep(n)[2] for _, n, _ in g.B5}
eis(_gebruikt == {r[2] for r in g.B5_BEROEPEN},
    f'niet elke categorie komt voor in de lijst: {sorted(_gebruikt)}')

# ==========================================================================
#  2. de antwoorden kloppen, langs een tweede weg uitgerekend
# ==========================================================================
eis(a.D1_TOTAAL == len(g.B1), 'eindtotaal van draaitabel 1 klopt niet')
for dimensie, aantal in a.D1.items():
    eis(aantal == sum(1 for r in g.B1 if r[2] == dimensie),
        f'D1 klopt niet bij {dimensie}')
    GETELD['cellen nagerekend'] += 1

eis(a.D2_TOTAAL == len(g.B2), 'eindtotaal van draaitabel 2 klopt niet')
for r in a.D2_RIJEN:
    for k in a.D2_KOLLEN:
        verwacht = sum(1 for x in g.B2 if x[1] == r and x[2] == k)
        eis(a.D2_CEL[(r, k)] == verwacht, f'D2 klopt niet bij {r}/{k}')
        GETELD['cellen nagerekend'] += 1
eis(a.D2_LEGE_VAKJES > 0,
    'de kruistabel heeft geen enkel leeg vakje — dan vervalt de derde vraag')

eis(a.D3_TOTAAL == sum(r[3] for r in g.B3), 'eindtotaal van draaitabel 3 klopt niet')
for materiaal, duurzaamheid in a.D3.items():
    eis(duurzaamheid == sum(r[3] for r in g.B3 if r[1] == materiaal),
        f'D3 klopt niet bij {materiaal}')
    GETELD['cellen nagerekend'] += 1

for p in a.D4_LABELS:
    hoort = [r[4] for r in g.B4 if g.periode(r[2]) == p]
    eis(a.D4_AANTAL[p] == len(hoort), f'D4 aantal klopt niet bij {p}')
    eis(abs(a.D4_GEMIDDELD[p] - sum(hoort) / len(hoort)) < 1e-9,
        f'D4 gemiddelde klopt niet bij {p}')
    GETELD['cellen nagerekend'] += 2
eis(a.D4_TOTAAL_AANTAL == len(g.B4), 'eindtotaal van draaitabel 4 klopt niet')
eis(0 < a.D4_FILTER_AANTAL < len(g.B4),
    f'het filter {a.D4_FILTER_THEMA} laat {a.D4_FILTER_AANTAL} updates over — '
    'dan leert het filteren niets')
# de pointe van blad 4: de drukste periode is niet die met de grootste updates
eis(a.D4_DRUKSTE != a.D4_HOOGSTE_GEM,
    'de drukste periode heeft ook het hoogste gemiddelde — dan valt er niets '
    'te zien in de tweede waardekolom')

for categorie, aanbiedingen in a.D5.items():
    verwacht = sum(x for _, n, x in g.B5 if g.beroep(n)[2] == categorie)
    eis(aanbiedingen == verwacht, f'D5 klopt niet bij {categorie}')
    GETELD['cellen nagerekend'] += 1
eis(a.D5_TOTAAL == sum(r[2] for r in g.B5), 'D5 eindtotaal klopt niet')
eis(a.D5_NA_TOTAAL == a.D5_TOTAAL + g.B5_EXTRA[2],
    'het eindtotaal na de extra dorpeling klopt niet')
eis(a.D5_NA != a.D5, 'de extra dorpeling verandert niets — dan leert ze niets')

# elke vraag moet precies één antwoord hebben; antwoorden.py bewaakt dat al met
# een assert bij het inlezen, maar de tekstantwoorden controleren we hier nog eens
for oef in o.OEFENINGEN:
    for cel, wat, antwoord, punten in oef['vragen']:
        eis(antwoord is not None and antwoord != '',
            f"{oef['blad']} {cel}: geen antwoord")
        eis(punten > 0, f"{oef['blad']} {cel}: geen punten")
        GETELD['vragen gekeurd'] += 1

# geen enkele vraag mag met een antwoord uit een ander blad samenvallen: dan
# kan een leerling het ergens anders aflezen zonder de draaitabel te maken
_tekst = collections.defaultdict(list)
for oef in o.OEFENINGEN:
    for cel, wat, antwoord, _p in oef['vragen']:
        if isinstance(antwoord, str):
            _tekst[antwoord].append(f"{oef['blad']} {cel}")
for antwoord, plekken in _tekst.items():
    eis(len(plekken) == 1,
        f'het antwoord {antwoord!r} staat op {len(plekken)} plaatsen: '
        f'{", ".join(plekken)}')

# ==========================================================================
#  3. het startbestand
# ==========================================================================
LIJSTEN = [
    ('1 Mobs', g.B1_KOPRIJ, g.B1_START, g.B1_EIND, g.B1_KOP, o.V1, g.B1_ANKER),
    ('2 Blokken', g.B2_KOPRIJ, g.B2_START, g.B2_EIND, g.B2_KOP, o.V2, g.B2_ANKER),
    ('3 Gereedschap', g.B3_KOPRIJ, g.B3_START, g.B3_EIND, g.B3_KOP, o.V3,
     g.B3_ANKER),
    ('4 Updates', g.B4_KOPRIJ, g.B4_START, g.B4_EIND, g.B4_KOP, o.V4, g.B4_ANKER),
    ('5 Dorpelingen', g.B5_KOPRIJ, g.B5_START, g.B5_EIND, g.B5_KOP, o.V5,
     g.B5_ANKER),
]
eis({x[0] for x in LIJSTEN} == {oef['blad'] for oef in o.OEFENINGEN},
    'de bladen van LIJSTEN en van OEFENINGEN lopen uit elkaar')

eis(os.path.exists(WERKBOEK), 'het startbestand bestaat niet — draai eerst '
                              'maak-startbestand.py')
if os.path.exists(WERKBOEK):
    wb = load_workbook(WERKBOEK)
    eis({x[0] for x in LIJSTEN} <= set(wb.sheetnames),
        f'ontbrekende bladen: {sorted({x[0] for x in LIJSTEN} - set(wb.sheetnames))}')

    for blad, koprij, start, eind, koppen, vraagkol, anker in LIJSTEN:
        if blad not in wb.sheetnames:
            continue
        ws = wb[blad]
        breedte = len(koppen)
        for i, titel in enumerate(koppen, start=1):
            eis(ws.cell(row=koprij, column=i).value == titel,
                f'{blad}: koprij verschilt in kolom {get_column_letter(i)}')
        eis(len(set(koppen)) == breedte, f'{blad}: twee kolommen heten hetzelfde')

        # te maken kolommen mogen leeg zijn; de rest moet gevuld zijn
        temaken = set()
        for oef in o.OEFENINGEN:
            if oef['blad'] != blad:
                continue
            for waar, *_rest in oef.get('vullen', []):
                temaken.add(column_index_from_string(
                    re.match(r'([A-Z]+)', waar).group(1)))
        for rij in range(start, eind + 1):
            gevuld = [c for c in range(1, breedte + 1)
                      if ws.cell(row=rij, column=c).value is not None]
            eis(gevuld, f'{blad}: rij {rij} is helemaal leeg')
        for kolom in range(1, breedte + 1):
            heeft = any(ws.cell(row=r, column=kolom).value is not None
                        for r in range(start, eind + 1))
            eis(heeft or kolom in temaken,
                f'{blad}: kolom {get_column_letter(kolom)} blijft leeg')

        # één lege kolom tussen de lijst en het vragenblok
        eis(vraagkol >= breedte + 2,
            f'{blad}: het vragenblok staat in kolom {get_column_letter(vraagkol)}, '
            f'de lijst loopt tot {get_column_letter(breedte)} — er hoort minstens '
            'één lege kolom tussen')
        for rij in range(start, eind + 1):
            eis(ws.cell(row=rij, column=breedte + 1).value is None,
                f'{blad}: kolom {get_column_letter(breedte + 1)} rij {rij} hoort '
                'leeg te blijven als scheiding')

        # de draaitabel staat onder de lijst, met ruimte ertussen
        kolom_anker, rij_anker = cel_van(anker)
        eis(kolom_anker == 1, f'{blad}: het anker {anker} staat niet in kolom A')
        ruimte = rij_anker - eind
        eis(ruimte >= 3,
            f'{blad}: de draaitabel begint op rij {rij_anker}, maar de lijst loopt '
            f'tot {eind} — laat minstens twee lege rijen ertussen')
        eis(ws.cell(row=rij_anker, column=1).value is None,
            f'{blad}: het anker {anker} is niet vrij')
        for kolom in range(1, breedte + 1):
            eis(ws.cell(row=eind + 1, column=kolom).value is None,
                f'{blad}: er staat iets op rij {eind + 1}, direct onder de lijst')
        GETELD['lijsten gekeurd'] += 1

    # er mag nog geen draaitabel in het startbestand zitten
    for ws in wb.worksheets:
        eis(not getattr(ws, '_pivots', None),
            f'{ws.title}: er zit al een draaitabel in het startbestand')

    # de antwoordcellen zijn leeg en staan in het vragenblok
    for oef in o.OEFENINGEN:
        ws = wb[oef['blad']]
        for cel, *_rest in oef['vragen']:
            kolom, rij = cel_van(cel)
            eis(ws.cell(row=rij, column=kolom).value is None,
                f"{oef['blad']} {cel} moet leeg zijn in het startbestand")
            eis(ws.cell(row=rij, column=kolom - 1).value is not None,
                f"{oef['blad']} {cel}: er staat geen vraag naast")
            GETELD['antwoordcellen leeg'] += 1
        for waar, *_rest in oef.get('vullen', []):
            k1, r1 = cel_van(waar.split(':')[0])
            k2, r2 = cel_van(waar.split(':')[1])
            for k in range(k1, k2 + 1):
                for r in range(r1, r2 + 1):
                    eis(ws.cell(row=r, column=k).value is None,
                        f"{oef['blad']} {get_column_letter(k)}{r} moet leeg zijn")
                    GETELD['te vullen cellen leeg'] += 1

    # de rij die ze er zelf bij typen, moet nog vrij zijn
    ws = wb['5 Dorpelingen']
    extra_rij = o.OEFENINGEN[4]['extra']['rij']
    for kolom in range(1, len(g.B5_KOP) + 1):
        eis(ws.cell(row=extra_rij, column=kolom).value is None,
            f'5 Dorpelingen: rij {extra_rij} moet vrij zijn — daar typen ze de '
            'extra dorpeling')
    eis(extra_rij == g.B5_EIND + 1,
        'de extra dorpeling hoort direct onder de lijst te komen')

    # de beroepentabel naast de dorpelingenlijst moet kloppen met gegevens.py
    for i, (beroep, werkblok, categorie) in enumerate(g.B5_BEROEPEN):
        r = g.B5_ZOEKRIJ + 1 + i
        for kolom, verwacht in ((g.B5_ZOEKKOL, g.beroepcode(i + 1)),
                                (g.B5_ZOEKKOL + 1, beroep),
                                (g.B5_ZOEKKOL + 2, werkblok),
                                (g.B5_ZOEKKOL + 3, categorie)):
            eis(ws.cell(row=r, column=kolom).value == verwacht,
                f'5 Dorpelingen: de beroepentabel wijkt af op rij {r}')
            GETELD['beroepentabel nagekeken'] += 1
    # de zoektabel mag de lijst noch het vragenblok raken
    eis(g.B5_ZOEKKOL > len(g.B5_KOP) + 1,
        'de beroepentabel staat tegen de dorpelingenlijst aan')
    eis(o.V5 > g.B5_ZOEKKOL + 4,
        'het vragenblok staat tegen de beroepentabel aan')

# ==========================================================================
#  4. de opdracht zelf
# ==========================================================================
for oef in o.OEFENINGEN:
    eis(oef['kaarten'], f"blad {oef['nr']}: geen verwijzing naar de theorie")
    for n in oef['kaarten']:
        eis(n in o.KAARTEN, f"blad {oef['nr']}: kaart {n} bestaat niet")
        eis(1 <= n <= 9, f"blad {oef['nr']}: kaart {n} gaat niet over draaitabellen")
    draai = oef['draai']
    koppen = {1: g.B1_KOP, 2: g.B2_KOP, 3: g.B3_KOP,
              4: g.B4_KOP, 5: g.B5_KOP}[oef['nr']]
    for rol in ('rijen', 'kolommen', 'waarden', 'tweede'):
        veld = draai.get(rol)
        eis(veld is None or veld in koppen,
            f"blad {oef['nr']}: {veld!r} is geen kolom van deze lijst ({rol})")
    eis(draai['bewerking'] in ('Aantal', 'Som', 'Gemiddelde', 'Max', 'Min'),
        f"blad {oef['nr']}: onbekende bewerking {draai['bewerking']!r}")
    # een VERT.ZOEKEN mag alleen een kolom aanvullen die ook echt gebruikt wordt
    for waar, _wat, kolom, formule, _p in oef.get('vullen', []):
        eis(formule.startswith('=VERT.ZOEKEN('),
            f"blad {oef['nr']} {waar}: geen VERT.ZOEKEN")
        eis(f';{kolom};' in formule,
            f"blad {oef['nr']} {waar}: de formule haalt kolom {kolom} niet op")
        eis(formule.count('$') == 4,
            f"blad {oef['nr']} {waar}: de zoektabel is niet volledig vastgezet")
    if oef.get('vullen'):
        aangevuld = {koppen[column_index_from_string(
            re.match(r'([A-Z]+)', waar).group(1)) - 1]
            for waar, *_r in oef['vullen']}
        eis(draai['rijen'] in aangevuld,
            f"blad {oef['nr']}: de draaitabel groepeert niet op een kolom die ze "
            'eerst moeten aanvullen — dan heeft het aanvullen geen zin')
    GETELD['bladen gekeurd'] += 1

# geen theorie in de opdracht: de bundel mag niet uitleggen wat een kaart doet
VERBODEN = ['waardeveldinstelling', 'veldenlijst', 'sleep het veld naar']
tekst = ' '.join(
    r for oef in o.OEFENINGEN
    for r in oef['inleiding'] + [oef['klaar'], oef['controle']]).lower()
for woord in VERBODEN:
    eis(woord not in tekst,
        f'de opdracht legt theorie uit ("{woord}") — dat hoort op de kaarten')

eis(sum(int(oef['duur'].split()[0]) for oef in o.OEFENINGEN) == g.DUUR,
    'de duur van de bladen telt niet op tot de aangekondigde duur')
eis(o.TOTAAL == 50, f'de bundel staat op {o.TOTAAL} punten in plaats van 50')

# ==========================================================================
#  5. de documenten zijn gebouwd
# ==========================================================================
for naam in ('minecraft-draaitabellen-opdracht.docx',
             'minecraft-draaitabellen-lerarensleutel.docx'):
    eis(os.path.exists(os.path.join(HIER, naam)),
        f'{naam} bestaat niet — draai eerst maak-documenten.py')

# ==========================================================================
print('=' * 72)
for naam, hoeveel in sorted(GETELD.items()):
    print(f'  {hoeveel:>5}  {naam}')
print('=' * 72)
if FOUTEN:
    print(f'{len(FOUTEN)} PROBLEMEN:')
    for f in FOUTEN:
        print('  !', f)
    sys.exit(1)
print('alles past bij elkaar.')
