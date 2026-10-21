# -*- coding: utf-8 -*-
"""Bouwt het startbestand van de Minecraft-draaitabelbundel.

Elke lijst staat kaarsrecht: één koprij, geen lege rij erin, niets eronder.
Het vragenblok staat ERNAAST met één lege kolom ertussen, en de draaitabel
komt ONDER de lijst met twee lege rijen ertussen. Zo stopt Excel bij het
bepalen van het bronbereik precies waar het hoort.

Er zit geen draaitabel in het bestand: openpyxl kan die niet schrijven, en de
leerling hoort ze zelf te maken.
"""
import importlib.util
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')

UIT = os.path.join(HIER, 'minecraft-draaitabellen-startbestand.xlsx')
TITEL = Font(bold=True, size=13)
GROOT = Font(bold=True, size=16)
VET = Font(bold=True)
SCHUIN = Font(italic=True, size=10)

wb = Workbook()


def breedtes(ws, waarden, start=1):
    for i, b in enumerate(waarden, start=start):
        ws.column_dimensions[get_column_letter(i)].width = b


def blad(naam, titel, ondertitel=None):
    ws = wb.create_sheet(naam)
    ws['A1'] = titel
    ws['A1'].font = TITEL
    if ondertitel:
        ws['A2'] = ondertitel
        ws['A2'].font = SCHUIN
    return ws


def lijst(ws, koprij, koppen, rijen):
    for i, k in enumerate(koppen, start=1):
        ws.cell(row=koprij, column=i, value=k).font = VET
    for r, waarden in enumerate(rijen, start=koprij + 1):
        for i, w in enumerate(waarden, start=1):
            ws.cell(row=r, column=i, value=w)


def vragen(ws, rij, kolom, teksten):
    ws.cell(row=rij, column=kolom, value='Vraag').font = VET
    ws.cell(row=rij, column=kolom + 1, value='Antwoord').font = VET
    for i, t in enumerate(teksten, start=1):
        ws.cell(row=rij + i, column=kolom, value=t)


def plek(ws, oef):
    """Het plekje waar de draaitabel moet komen, met een regel erboven."""
    anker = oef['draai']['anker']
    rij = int(anker[1:])
    ws.cell(row=rij - 1, column=1,
            value=f"Zet hier je draaitabel — linkerbovenhoek in {anker}."
            ).font = SCHUIN


# ------------------------------------------------------------------- Start
ws = wb.active
ws.title = 'Start'
ws['A1'] = g.TITEL
ws['A1'].font = GROOT
KOP_REGELS = 'Drie regels die op elk blad gelden'
regels = [
    g.ONDERTITEL, '',
    'Naam: ......................................    Klas: ..........', '',
    f'Vijf werkbladen, samen {g.DUUR} minuten. Op elk blad maak je één draaitabel.',
    'Er staat in deze bundel geen theorie. Die staat in:',
    f'   {g.THEORIEBUNDEL} — elke kaart op een eigen pagina.',
    'Bij elk blad staat welke kaarten je erbij nodig hebt.', '',
]
for oef in o.OEFENINGEN:
    kaarten = 'kaart ' + ' en '.join(str(n) for n in oef['kaarten']) \
        if len(oef['kaarten']) == 2 else \
        'kaart ' + ', '.join(str(n) for n in oef['kaarten'])
    regels.append(f"Blad {oef['nr']}  {oef['blad'][2:]:<13} "
                  f"{oef['duur'].split()[0]:>2} min   {kaarten:<15} {oef['titel']}")
regels += [
    '', KOP_REGELS,
    '  1. De lijst blijft zoals ze is. Voeg er geen lege rij in en zet er niets',
    '     onder. De vragen staan ERNAAST, met een lege kolom ertussen.',
    '  2. De draaitabel komt ONDER de lijst, op de plaats die aangeduid staat.',
    '  3. Lees het kopje van je draaitabel. Staat er Aantal waar Som hoort,',
    '     dan klopt je hele tabel niet — en Excel zegt daar niets over.',
    '',
    'EERST DOEN: sla dit bestand op als  minecraft-draaitabellen-jouwnaam.xlsx',
    '', g.WAARSCHUWING,
]
for i, r in enumerate(regels, start=2):
    ws.cell(row=i, column=1, value=r)
ws.cell(row=2 + regels.index(KOP_REGELS), column=1).font = VET
rij_waarschuwing = 2 + regels.index(g.WAARSCHUWING)
c = ws.cell(row=rij_waarschuwing, column=1)
c.font = SCHUIN
c.alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[rij_waarschuwing].height = 100
breedtes(ws, [92])

# ------------------------------------------------------------------- 1 Mobs
oef = o.OEFENINGEN[0]
ws = blad(oef['blad'], 'De mobs', g.B1_VOETNOOT)
lijst(ws, g.B1_KOPRIJ, g.B1_KOP, g.B1)
vragen(ws, g.B1_KOPRIJ, o.V1, [
    'In welke dimensie horen de meeste mobs thuis?', 'Hoeveel mobs zijn dat?'])
plek(ws, oef)
breedtes(ws, [16, 12, 12, 14])
breedtes(ws, [3, 40, 14], start=o.V1 - 1)

# ---------------------------------------------------------------- 2 Blokken
oef = o.OEFENINGEN[1]
ws = blad(oef['blad'], 'De blokken', g.B2_VOETNOOT)
lijst(ws, g.B2_KOPRIJ, g.B2_KOP, g.B2)
vragen(ws, g.B2_KOPRIJ, o.V2, [
    'Welk gereedschap heb je voor de meeste blokken nodig?',
    'Hoeveel natuurlijke blokken breek je met een schep?',
    'In welke categorie vraagt élk blok hetzelfde gereedschap?'])
plek(ws, oef)
breedtes(ws, [18, 13, 14, 11])
breedtes(ws, [3, 46, 14], start=o.V2 - 1)

# ------------------------------------------------------------ 3 Gereedschap
oef = o.OEFENINGEN[2]
ws = blad(oef['blad'], 'Het gereedschap', g.B3_VOETNOOT)
lijst(ws, g.B3_KOPRIJ, g.B3_KOP, g.B3)
vragen(ws, g.B3_KOPRIJ, o.V3, [
    'Welk materiaal haalt samen de hoogste duurzaamheid?',
    'Hoeveel is dat?',
    'Welk materiaal komt het laagst uit?'])
plek(ws, oef)
breedtes(ws, [22, 13, 13, 14])
breedtes(ws, [3, 46, 14], start=o.V3 - 1)

# ---------------------------------------------------------------- 4 Updates
oef = o.OEFENINGEN[3]
ws = blad(oef['blad'], 'De grote updates', g.B4_VOETNOOT)
lijst(ws, g.B4_KOPRIJ, g.B4_KOP, g.B4)
vragen(ws, g.B4_KOPRIJ, o.V4, [
    'In welke periode verschenen de meeste updates?',
    'Hoeveel updates zijn dat?',
    f'Hoeveel updates blijven over als je filtert op {a.D4_FILTER_THEMA}?'])
plek(ws, oef)
breedtes(ws, [9, 34, 8, 14, 16])
breedtes(ws, [3, 46, 14], start=o.V4 - 1)

# ------------------------------------------------------------ 5 Dorpelingen
oef = o.OEFENINGEN[4]
ws = blad(oef['blad'], 'De dorpelingen', g.B5_VOETNOOT)
lijst(ws, g.B5_KOPRIJ, g.B5_KOP,
      [(naam, g.beroepcode(n), aanbiedingen, None, None, None)
       for naam, n, aanbiedingen in g.B5])
# de beroepentabel ernaast
k = g.B5_ZOEKKOL
ws.cell(row=g.B5_ZOEKRIJ - 1, column=k,
        value='Beroepentabel — niet aanpassen').font = SCHUIN
for i, titel in enumerate(['Code', 'Beroep', 'Werkblok', 'Categorie'], start=k):
    ws.cell(row=g.B5_ZOEKRIJ, column=i, value=titel).font = VET
for i, (beroep, werkblok, categorie) in enumerate(g.B5_BEROEPEN):
    r = g.B5_ZOEKRIJ + 1 + i
    ws.cell(row=r, column=k, value=g.beroepcode(i + 1))
    ws.cell(row=r, column=k + 1, value=beroep)
    ws.cell(row=r, column=k + 2, value=werkblok)
    ws.cell(row=r, column=k + 3, value=categorie)
vragen(ws, g.B5_KOPRIJ, o.V5, [
    'Welke categorie heeft samen de meeste aanbiedingen?',
    'Hoeveel aanbiedingen zijn dat?',
    'Wat is het eindtotaal na de extra dorpeling en na Vernieuwen?'])
plek(ws, oef)
breedtes(ws, [14, 12, 14, 18, 18, 13])
breedtes(ws, [3, 10, 18, 18, 13], start=k - 1)
breedtes(ws, [3, 46, 14], start=o.V5 - 1)

wb.save(UIT)
print('geschreven:', UIT)
for ws in wb.worksheets:
    gevuld = sum(1 for rij in ws.iter_rows() for c in rij if c.value is not None)
    print(f'  {ws.title:<16} {gevuld:>4} gevulde cellen   {ws.dimensions}')
