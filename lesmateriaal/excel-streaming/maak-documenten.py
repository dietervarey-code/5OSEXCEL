# -*- coding: utf-8 -*-
"""Bouwt de opdrachtbundel en de lerarensleutel van de streamingbundel.

Het raamwerk (kaders, tabellen, formuleopmaak) is hetzelfde als in
excel-ottomaans: de twee bundels horen er voor de leerling ook
hetzelfde uit te zien.
"""
import importlib.util
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')

OPDRACHT = os.path.join(HIER, 'streaming-opdracht.docx')
SLEUTEL = os.path.join(HIER, 'streaming-lerarensleutel.docx')

BLAUW = RGBColor(0x1F, 0x38, 0x64)
ORANJE = RGBColor(0xB0, 0x4A, 0x00)
GRIJS = RGBColor(0x59, 0x59, 0x59)
GROEN = RGBColor(0x38, 0x76, 0x1D)

FORMULE = re.compile(r"(=[A-Z.]*\((?:[^()]|\([^()]*\))*\)|=[A-Z$]+\d+(?:[*+\-/]"
                     r"[A-Z$]*\d+)*)")


def opzet(doc):
    s = doc.styles['Normal']
    s.font.name = 'Calibri'
    s.font.size = Pt(11)
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(1.5)
        sec.left_margin = sec.right_margin = Cm(1.8)


def kop(doc, tekst, niveau=1, kleur=BLAUW):
    p = doc.add_heading(tekst, level=niveau)
    for r in p.runs:
        r.font.color.rgb = kleur
    return p


def zin(p, tekst):
    """Zet formules in een vaste letter, zodat ze opvallen in lopende tekst."""
    for stuk in FORMULE.split(str(tekst)):
        if not stuk:
            continue
        r = p.add_run(stuk)
        if FORMULE.fullmatch(stuk):
            r.font.name = 'Consolas'
            r.font.size = Pt(10)
    return p


def alinea(doc, tekst, style=None, vet=False, schuin=False, kleur=None,
           grootte=None, na=None):
    p = doc.add_paragraph(style=style)
    if vet or schuin or kleur or grootte:
        r = p.add_run(str(tekst))
        r.bold, r.italic = vet, schuin
        if kleur:
            r.font.color.rgb = kleur
        if grootte:
            r.font.size = Pt(grootte)
    else:
        zin(p, tekst)
    if na is not None:
        p.paragraph_format.space_after = Pt(na)
    return p


def kader(doc, titel, regels, kleur=BLAUW):
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    c = t.rows[0].cells[0]
    r = c.paragraphs[0].add_run(titel)
    r.bold = True
    r.font.color.rgb = kleur
    for regel in regels:
        pp = c.add_paragraph()
        pp.paragraph_format.space_after = Pt(2)
        zin(pp, regel)
    doc.add_paragraph()
    return t


def tabel(doc, koppen, rijen, stijl='Light Grid Accent 1', breedtes=None):
    t = doc.add_table(rows=1, cols=len(koppen))
    t.style = stijl
    for i, k in enumerate(koppen):
        t.rows[0].cells[i].paragraphs[0].add_run(str(k)).bold = True
    for rij in rijen:
        cellen = t.add_row().cells
        for i, w in enumerate(rij):
            zin(cellen[i].paragraphs[0], w)
    if breedtes:
        for rij in t.rows:
            for i, b in enumerate(breedtes):
                rij.cells[i].width = Cm(b)
    doc.add_paragraph()
    return t


def paginaeinde(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def genummerd(doc, regels, kleur=BLAUW):
    for i, regel in enumerate(regels, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f'{i}.  ')
        r.bold = True
        r.font.color.rgb = kleur
        zin(p, regel)


# ==========================================================================
#  1. de opdrachtbundel
# ==========================================================================
doc = Document()
opzet(doc)

kop(doc, g.TITEL, 0)
alinea(doc, g.ONDERTITEL, schuin=True, kleur=GRIJS, na=4)
alinea(doc, 'Naam: ......................................    Klas: ..........',
       na=14)

alinea(doc,
       f'Vijf werkbladen, samen ongeveer {g.DUUR} minuten. De eerste drie zijn '
       'herhaling: functies die je al kent, op een streamingaanbod en op een '
       'enquête over schermtijd. Blad 4 en 5 zijn het nieuwe werk. Daar typ je geen '
       'formules meer, maar sleep je een lijst in elkaar tot een overzicht en maak '
       'je er een grafiek van.', na=8)
alinea(doc,
       'Werk in volgorde. Blad 2 haalt gegevens uit blad 1, en blad 4 en 5 werken '
       'allebei verder op blad 2. Sla een blad over en je loopt vast.', na=12)

tabel(doc, ['Blad', 'Wat', 'Duur', 'Waarmee'],
      [[f"{oef['nr']}  {oef['blad'][2:]}", oef['titel'], oef['duur'],
        ', '.join(dict.fromkeys(
            r['hoe'] for r in oef['regels']
            if r['hoe'] not in (o.LEZEN,)))]
       for oef in o.OEFENINGEN],
      breedtes=[3.4, 4.6, 2.2, 7.0])

kader(doc, 'Wat je erbij nodig hebt', [
    f'{g.FUNCTIEBUNDEL} — de functiefiches die je al kent. Voor blad 1 tot 3.',
    f'{g.THEORIEBUNDEL} — de nieuwe theoriebundel. Kaart 1 tot 9 voor blad 4, '
    'kaart 10 tot 14 voor blad 5.',
], BLAUW)

kader(doc, 'Drie regels voor deze bundel', [
    '1.  De lijst op blad 2 wordt de bron van allebei je draaitabellen. Voeg er geen '
    'lege rij in, zet er geen totaalrij onder en verwijder er niets uit. De '
    'berekeningen staan daarom ERNAAST, met een lege kolom ertussen.',
    '2.  Op blad 4 en 5 typ je geen formules. Je sleept, je kiest een grafiektype en '
    'je leest af.',
    '3.  Een draaitabel rekent niet automatisch mee. Verander je iets aan blad 2, '
    'vernieuw dan je draaitabel (rechtsklik erin, Vernieuwen).',
], ORANJE)

kader(doc, 'Over de cijfers', [g.WAARSCHUWING], GRIJS)

for oef in o.OEFENINGEN:
    paginaeinde(doc)
    kop(doc, f"Blad {oef['nr']} — {oef['titel']}", 1)
    alinea(doc, f"werkblad '{oef['blad']}'  ·  {oef['duur']}", schuin=True,
           kleur=GRIJS, na=8)
    for regel in oef['inleiding']:
        alinea(doc, regel, na=4)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    kop(doc, 'Wat je maakt', 2, GRIJS)
    tabel(doc, ['Waar', 'Wat', 'Waarmee'],
          [[r['waar'], r['wat'], r['hoe']] for r in oef['regels']],
          breedtes=[3.2, 8.4, 5.6])

    toelichtingen = [(r['waar'], r['toelichting']) for r in oef['regels']
                     if r.get('toelichting')]
    if toelichtingen:
        kop(doc, 'Waar je op moet letten', 2, GRIJS)
        for waar, tekst in toelichtingen:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(f'{waar}  —  ')
            r.bold = True
            r.font.color.rgb = BLAUW
            zin(p, tekst)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    kop(doc, 'Opmaak', 2, GRIJS)
    for tekst, _ in oef['opmaak']:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        zin(p, tekst)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    kader(doc, 'Je blad is af als…', [oef['klaar'], '', 'CONTROLE  ·  ' +
                                      oef['controle']], GROEN)
    kader(doc, 'Loop je vast?', oef['vast'], ORANJE)

paginaeinde(doc)
kop(doc, 'Helemaal klaar?', 1)
alinea(doc, 'Loop dit lijstje af voor je indient.', schuin=True, kleur=GRIJS, na=8)
for regel in [
    'Elk controlegetal klopt.',
    'Nergens staat #N/B, #WAARDE! of #DEEL/0!.',
    'Je hebt elke formule één keer getypt en doorgevoerd — nergens met de hand '
    'een getal ingevuld.',
    'Je twee draaitabellen zijn vernieuwd (rechtsklik, Vernieuwen).',
    'De rij Eindtotaal van draaitabel 2 geeft hetzelfde getal als het gemiddelde '
    'dat je op blad 2 berekende.',
    'Beide grafieken hebben een titel die zegt wat je ziet.',
    'Je bestand heet streaming-jouwnaam.xlsx.',
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run('☐   ')
    r.font.size = Pt(13)
    zin(p, regel)

doc.save(OPDRACHT)
print('geschreven:', OPDRACHT)




# ==========================================================================
#  2. de lerarensleutel
# ==========================================================================
doc = Document()
opzet(doc)

kop(doc, g.TITEL + ' — lerarensleutel', 0)
alinea(doc, g.ONDERTITEL, schuin=True, kleur=GRIJS, na=12)

kop(doc, 'In het kort', 2, GRIJS)
alinea(doc,
       f'Vijf werkbladen, {g.DUUR} minuten. Dezelfde opbouw als de Ottomaanse '
       f'bundel ({g.EERDERE_BUNDEL}): blad 1 tot 3 zijn herhaling, blad 4 en 5 zijn '
       'de draaitabel en de grafiek. Wie die bundel gemaakt heeft, herkent elk blad '
       'en werkt sneller. Er zit geen enkele geneste functie in.', na=6)
alinea(doc,
       f'{o.TOTAAL} punten voor het werk en {o.TOTAAL_OPMAAK} voor de opmaak, samen '
       f'{o.TOTAAL_ALLES}. Twee derde daarvan geeft een score op 50.', na=10)

tabel(doc, ['Blad', 'Duur', 'Werk', 'Opmaak', 'Samen'],
      [[f"{oef['nr']}  {oef['blad'][2:]}", oef['duur'],
        sum(r['punten'] for r in oef['regels']),
        sum(p for _, p in oef['opmaak']),
        sum(r['punten'] for r in oef['regels']) + sum(p for _, p in oef['opmaak'])]
       for oef in o.OEFENINGEN]
      + [['Samen', f'{g.DUUR} minuten', o.TOTAAL, o.TOTAAL_OPMAAK, o.TOTAAL_ALLES]],
      breedtes=[4.6, 3.0, 2.6, 2.6, 2.6])

kader(doc, 'Wat echt is en wat verzonnen', [
    'De titels, de taal en het jaar van verschijnen kloppen. Dat mag je nakijken en '
    'dat is ook leuk om te doen.',
    'De indeling in genres is onze keuze. Veel titels passen in meer dan één genre — '
    'een leerling die vindt dat Parasite eerder drama dan thriller is, heeft gelijk '
    'én ongelijk. Dat is zelf een les: een draaitabel is maar zo goed als de kolom '
    'waarop je ze bouwt.',
    'ALLE AANTALLEN ZIJN VERZONNEN. Er is geen peiling over lievelingsgenres '
    'geweest, geen enquête over kijkers en geen meting van schermtijd. Dat staat op '
    'het startblad en in de opdracht. Zeg het er bij de bespreking nog eens bij: de '
    'cijfers op blad 5 lijken op wat echt onderzoek vindt, maar het zijn onze '
    'cijfers.',
], ORANJE)

kader(doc, 'Een draaitabel kan niet in het startbestand zitten', [
    'De bibliotheek waarmee deze bestanden gemaakt worden, kan er geen schrijven. '
    'Dat komt goed uit — de leerling hoort ze zelf te maken. Maar het betekent ook '
    'dat blad 4 met de hand nagekeken moet worden. Hieronder staan allebei de '
    'draaitabellen volledig uitgeschreven, zodat je het blad van een leerling er '
    'gewoon naast kan leggen.',
], BLAUW)

kader(doc, 'De drie dingen waar het mis gaat', [
    '1.  Op blad 2 vergeten ze de dollartekens rond de zoektabel. Onderaan kolom G '
    'staat dan #N/B. Dezelfde fout als altijd, hier alleen zichtbaarder omdat de '
    'lijst 35 rijen telt.',
    '2.  Draaitabel 2 staat op Som in plaats van Gemiddelde. Herkenbaar aan het '
    f'eindtotaal: staat er {a.B2_TOTAAL_KIJKERS} in plaats van '
    f'{a.D2_EINDTOTAAL_GEMIDDELD:g}, dan telt ze op in plaats van te middelen. Dit '
    'is nieuw ten opzichte van de vorige bundel — daar volstond Som — dus verwacht '
    'dat de helft het eerst fout heeft.',
    '3.  Op blad 5 selecteren ze bij het staafdiagram ook de kolom met het aantal '
    'titels, of de rij Eindtotaal. Ze krijgen dan twee reeksen, of één reusachtige '
    'balk waar de rest bij in het niet valt.',
], ORANJE)

for oef in o.OEFENINGEN:
    paginaeinde(doc)
    kop(doc, f"Blad {oef['nr']} — {oef['titel']}", 1)
    alinea(doc, f"werkblad '{oef['blad']}'  ·  {oef['duur']}  ·  "
                f"{sum(r['punten'] for r in oef['regels'])} + "
                f"{sum(p for _, p in oef['opmaak'])} punten",
           schuin=True, kleur=GRIJS, na=8)

    rijen = []
    for r in oef['regels']:
        if r.get('formule'):
            antwoord = r['formule']
        elif r.get('antwoord') is not None:
            antwoord = str(r['antwoord'])
        else:
            antwoord = '(zie hieronder)'
        rijen.append([r['waar'], antwoord, r['punten']])
    tabel(doc, ['Waar', 'Formule of antwoord', 'Pt'], rijen,
          breedtes=[3.4, 11.6, 1.4])

    if oef['nr'] == 1:
        kop(doc, 'De uitkomsten', 2, GRIJS)
        tabel(doc, ['Code', 'Genre', '2024', '2025', 'Groei', 'Sterk?'],
              [[c, n, oud, nieuw, a.B1_GROEI[i], a.B1_STERK[i]]
               for i, (c, n, oud, nieuw) in enumerate(g.B1)], 'Light List Accent 1')
        tabel(doc, ['Cel', 'Uitkomst'],
              [[f'B{g.B1_SAMENVATTING}', a.B1_SOM],
               [f'B{g.B1_SAMENVATTING + 1}', f'{a.B1_GEMIDDELDE:g}'],
               [f'B{g.B1_SAMENVATTING + 2}', f'{a.B1_MAX} ({a.B1_SNELST})'],
               [f'B{g.B1_SAMENVATTING + 3}', f'{a.B1_MIN} ({a.B1_TRAAGST})'],
               [f'B{g.B1_SAMENVATTING + 4}', a.B1_AANTAL],
               [f'B{g.B1_SAMENVATTING + 5}', a.B1_STERKE]])
        kader(doc, 'De vallen op dit blad', [
            f'Drie genres groeien met precies {g.B1_DREMPEL}, de grens zelf: '
            f'{", ".join(a.B1_GRENSGEVAL)}. Wie > typt in plaats van >= komt op '
            f'{a.B1_STERKE - len(a.B1_GRENSGEVAL)} in plaats van {a.B1_STERKE}. Dat '
            'is een verschil van drie, dus het valt hier wél meteen op — in de vorige '
            'bundel was het er maar één.',
            f'{a.B1_TRAAGST} krimpt met {a.B1_MIN}. Wie dat voor een fout houdt en '
            'het "verbetert", sloopt zijn SOM, zijn gemiddelde en zijn MIN in één '
            'keer. Zeg erbij dat een negatieve groei mag.',
            f'B{g.B1_SAMENVATTING + 4} vraagt AANTAL. Richten ze dat op de namen in '
            'kolom B, dan krijgen ze 0 — AANTAL telt alleen getallen.',
        ], GROEN)

    if oef['nr'] == 2:
        kop(doc, 'Steekproef op de uitkomsten', 2, GRIJS)
        proef = [0, 1, 17, len(g.B2) - 1]
        tabel(doc, ['Rij', 'Titel', 'Code', 'G (genre)', 'H (groei)', 'I (sterk?)'],
              [[g.B2_START + i, g.B2[i][0], g.B2[i][2], a.B2_GENRE[i],
                a.B2_GROEI[i], a.B2_STERK[i]] for i in proef],
              'Light List Accent 1')
        tabel(doc, ['Cel', 'Uitkomst'],
              [[f'{o.A2}{g.B2_VRAAGRIJ + 1}',
                f'{a.B2_GEMIDDELD:g}  — dit is de controle voor draaitabel 2'],
               [f'{o.A2}{g.B2_VRAAGRIJ + 2}', a.B2_AANTAL],
               [f'{o.A2}{g.B2_VRAAGRIJ + 3}', a.B2_TEL_SOORT],
               [f'{o.A2}{g.B2_VRAAGRIJ + 4}', a.B2_TEL_TAAL]])

    if oef['nr'] == 3:
        kop(doc, 'De uitkomsten', 2, GRIJS)
        tabel(doc, ['Code', 'Lft', 'Schooldag', 'Weekenddag', 'Per week', 'Veel?'],
              [[c, lft, s, w, a.B3_WEEK[i], a.B3_VEEL[i]]
               for i, (c, lft, s, w) in enumerate(g.B3)], 'Light List Accent 1')
        tabel(doc, ['Cel', 'Uitkomst'],
              [[f'{o.A3}{g.B3_VRAAGRIJ + 1}', a.B3_SOM],
               [f'{o.A3}{g.B3_VRAAGRIJ + 2}', f'{a.B3_GEMIDDELDE:g}'],
               [f'{o.A3}{g.B3_VRAAGRIJ + 3}', a.B3_MAX],
               [f'{o.A3}{g.B3_VRAAGRIJ + 4}', a.B3_MIN],
               [f'{o.A3}{g.B3_VRAAGRIJ + 5}', a.B3_TEL],
               [f'{o.A3}{g.B3_VRAAGRIJ + 6}', a.B3_SOMLEEFTIJD],
               [f'{o.A3}{g.B3_VRAAGRIJ + 7}', a.B3_SOMVEEL]])
        kader(doc, 'De val op dit blad', [
            f'{" en ".join(a.B3_GRENSGEVAL)} zitten op precies {g.B3_DREMPEL} '
            f'minuten, de grens zelf. Met > in plaats van >= komen ze op '
            f'{a.B3_TEL - len(a.B3_GRENSGEVAL)} in plaats van {a.B3_TEL}.',
            'Bij SOM.ALS wisselen ze het zoekbereik en het optelbereik om. Ze krijgen '
            'dan 0. Anders dan bij SOMMEN.ALS komt het zoekbereik hier EERST.',
            f'Het criterium van de eerste SOM.ALS is een getal ({g.B3_LEEFTIJD}), '
            'geen tekst. Aanhalingstekens eromheen mag, maar hoeft niet.',
        ], GROEN)
        alinea(doc,
               f'Terzijde: {a.B3_GEMIDDELDE:g} minuten per week komt neer op '
               f'{a.B3_PER_DAG:.0f} minuten per dag. Op blad 5 staat voor '
               f'{a.B5_JAREN[-1]} precies {a.B5_MINUTEN[-1]}. Dat is geen toeval maar '
               'ook geen bewijs — allebei verzonnen. Het is wel een mooi bruggetje '
               'naar de vraag of hun eigen schermtijd daar in de buurt ligt.', na=10)

    if oef['nr'] == 4:
        kop(doc, 'Draaitabel 1 — zo moet ze eruitzien', 2, GRIJS)
        alinea(doc, f"Bron: blad '{g.B4_DRAAI1['bron']}', bereik "
                    f"A{g.B2_KOPRIJ}:{o._kol(len(g.B2_KOP))}{g.B2_EIND}. "
                    f"Rijen = {g.B4_DRAAI1['rijen']}, "
                    f"Kolommen = {g.B4_DRAAI1['kolommen']}, "
                    f"Waarden = Aantal van {g.B4_DRAAI1['waarden']}.",
               schuin=True, kleur=GRIJS, na=6)
        tabel(doc, ['Aantal'] + a.D1_KOLLABELS + ['Eindtotaal'],
              [[r] + [a.D1_CEL[(r, k)] or '' for k in a.D1_KOLLABELS]
               + [a.D1_RIJTOTAAL[r]] for r in a.D1_RIJLABELS]
              + [['Eindtotaal'] + [a.D1_KOLTOTAAL[k] for k in a.D1_KOLLABELS]
                 + [a.D1_EINDTOTAAL]], 'Light List Accent 1')

        kop(doc, 'Draaitabel 2 — zo moet ze eruitzien', 2, GRIJS)
        alinea(doc, f"Bron: dezelfde lijst. Rijen = {g.B4_DRAAI2['rijen']}, twee "
                    f"waardevelden: Aantal van {g.B4_DRAAI2['tweede_waarde']} en "
                    f"Gemiddelde van {g.B4_DRAAI2['waarden']}, aflopend gesorteerd op "
                    'de gemiddelden.', schuin=True, kleur=GRIJS, na=6)
        tabel(doc, ['Taal', 'Aantal titels', 'Gemiddeld aantal kijkers'],
              [[t, a.D2_AANTAL[t], a._getal(a.D2_GEMIDDELD[t])]
               for t in a.D2_AFLOPEND]
              + [['Eindtotaal', a.D2_EINDTOTAAL_AANTAL,
                  a._getal(a.D2_EINDTOTAAL_GEMIDDELD)]], 'Light List Accent 1',
              breedtes=[4.4, 4.0, 5.4])

        kader(doc, 'Waar je naar kijkt bij het verbeteren', [
            f'Eindtotaal van draaitabel 1 = {a.D1_EINDTOTAAL}. Staat er iets anders, '
            'dan zit er een rij te veel of te weinig in de bron.',
            f'Rij Eindtotaal van draaitabel 2 = {a.D2_EINDTOTAAL_GEMIDDELD:g}, '
            'hetzelfde getal dat ze op blad 2 al berekenden. Staat er '
            f'{a.B2_TOTAAL_KIJKERS}, dan staat het veld op Som; staat er '
            f'{a.B2_AANTAL}, dan op Aantal.',
            f'{a.D2_AFLOPEND[4]} en {a.D2_AFLOPEND[5]} staan allebei op '
            f'{a._getal(a.D2_GEMIDDELD[a.D2_AFLOPEND[4]])}. De volgorde tussen die '
            'twee mag dus anders zijn dan hierboven.',
            'Of ze de tabel op de juiste cel gezet hebben, is bijzaak. Of ze '
            'draaitabel 2 gesorteerd hebben, niet: zonder sorteren is de laatste '
            'vraag veel moeilijker en staan de balken op blad 5 door elkaar.',
        ], GROEN)

        kader(doc, 'De vraag om klassikaal te bespreken', [
            f'{a.D2_MEESTE_TITELS} heeft {a.D2_AANTAL[a.D2_MEESTE_TITELS]} van de '
            f'{a.B2_AANTAL} titels — meer dan de helft — en staat toch pas op plaats '
            f'{a.D2_AFLOPEND.index(a.D2_MEESTE_TITELS) + 1} bij het gemiddelde. '
            f'{a.D2_TOPTAAL} staat bovenaan met '
            f'{a._getal(a.D2_TOPTAAL_GEMIDDELD)} kijkers per titel, uit '
            f'{a.D2_AANTAL[a.D2_TOPTAAL]} titels.',
            'Dat is het hele punt van deze les in één tabel: "veel" en "populair" '
            'zijn niet hetzelfde, en je ziet het verschil pas als je de lijst op twee '
            'manieren samenvat. Vraag hen wat eerlijker is om te zeggen, en wat ze '
            f'zouden antwoorden als iemand beweert dat {a.D2_TOPTAAL}e series '
            'populairder zijn.',
            f'Nuanceer ook de andere kant: {a.D2_TOPTAAL} heeft maar '
            f'{a.D2_AANTAL[a.D2_TOPTAAL]} titels in deze lijst, en allebei toevallig '
            'grote successen. Een gemiddelde over twee titels zegt weinig. Dat is een '
            'echte statistische les en ze kunnen ze hier zelf zien.',
        ], BLAUW)

    if oef['nr'] == 5:
        kop(doc, 'De verschilkolom', 2, GRIJS)
        tabel(doc, ['Jaar', 'Min/dag', 'Verschil'],
              [[j, m, '' if a.B5_VERSCHIL[i] is None else f'{a.B5_VERSCHIL[i]:+d}']
               for i, (j, m) in enumerate(g.B5)], 'Light List Accent 1',
              breedtes=[3.0, 3.6, 3.6])
        kader(doc, 'Wat je van de grafieken verwacht', [
            f'Grafiek 1 is een staafdiagram van de gemiddelden uit draaitabel 2. '
            f'{a.D2_TOPTAAL} heeft de langste balk '
            f'({a._getal(a.D2_TOPTAAL_GEMIDDELD)}), {a.D2_LAAGSTE} de kortste. Acht '
            'balken, geen negende: de rij Eindtotaal hoort er niet in.',
            f'Grafiek 2 is een lijndiagram van A{g.B5_KOPRIJ}:B{g.B5_EIND}. Eén lijn, '
            f'met een duidelijke sprong van +{a.B5_STIJGING} tussen '
            f'{a.B5_STIJGING_VAN} en {a.B5_STIJGING_TOT} en daarna een daling die '
            'niet helemaal terugkeert.',
            'Twee lijnen in grafiek 2 betekent dat ze de kolom Jaar als reeks hebben '
            'meegenomen. Twee reeksen in grafiek 1 betekent dat ze ook het aantal '
            'titels geselecteerd hebben.',
            f'De sprong in {a.B5_STIJGING_TOT} is het coronajaar. Vraag hun of ze zelf '
            'nog weten hoeveel ze toen op een scherm zaten, en of de lijn daarna '
            'terug moest zakken naar het oude niveau. Dat is een betere bespreking '
            'dan welk cijfer ook.',
        ], GROEN)

    if oef['opmaak']:
        kop(doc, 'Opmaak', 2, GRIJS)
        tabel(doc, ['Wat', 'Pt'], [[t, p] for t, p in oef['opmaak']],
              breedtes=[14.0, 1.4])

paginaeinde(doc)
kop(doc, 'Tijdsindeling en tips', 1)
tabel(doc, ['Wanneer', 'Wat'],
      [['0 - 5 min', 'Bestand openen en opslaan onder hun eigen naam. Zeg welke '
                     'cijfers verzonnen zijn. Wie de Ottomaanse bundel gemaakt '
                     'heeft, mag meteen beginnen.'],
       ['5 - 30 min', 'Blad 1 tot 3 zelfstandig. Herhaling; loop rond en kijk vooral '
                      'naar de dollartekens op blad 2 en naar >= op blad 1 en 3.'],
       ['30 - 32 min', 'Kort klassikaal moment, alleen als het nodig is: toon één keer '
                       'hoe je een tweede veld naar Waarden sleept en het op '
                       'Gemiddelde zet. Dat is het enige echt nieuwe stuk.'],
       ['32 - 44 min', 'Blad 4. Hier zit de leerwinst.'],
       ['44 - 50 min', 'Blad 5, en als het lukt de bespreking van de laatste vraag. '
                       'Wie niet aan de tweede grafiek komt, heeft de les toch '
                       'gehad.']],
      breedtes=[3.4, 13.0])

kader(doc, 'Als je weinig tijd hebt', [
    'Laat blad 3 vallen. Blad 4 en 5 hebben het niet nodig — allebei de draaitabellen '
    'komen van blad 2. Je verliest dan SOM.ALS, maar dat hebben ze in de vorige '
    'bundels al twee keer gehad.',
    'Blad 5 kan ook als apart lesje van tien minuten op een andere dag. De '
    'schermtijdgrafiek staat los van de rest.',
], BLAUW)

kader(doc, 'Voor wie snel klaar is', [
    'Laat hen een derde draaitabel maken van blad 3: Rijen = Leeftijd, Waarden = '
    'Gemiddelde van Per week. Die lijst is met opzet even netjes opgebouwd als blad '
    '2, dus het werkt meteen. De vraag erbij: zitten oudere leerlingen langer op een '
    'scherm dan jongere, en durf je dat te besluiten uit twintig antwoorden?',
    'Of laat hen in draaitabel 1 het veld Taal naar Filters slepen en kijken wat er '
    'met de tabel gebeurt als ze één taal kiezen.',
], BLAUW)

kader(doc, 'Waar dit op aansluit', [
    f'Dezelfde opbouw als {g.EERDERE_BUNDEL}: blad voor blad hetzelfde soort werk. '
    'Nieuw is alleen dat draaitabel 2 twee waardevelden heeft en op Gemiddelde moet '
    'staan in plaats van op Som.',
    f'De theoriebundel {g.THEORIEBUNDEL} hoort er weer bij en is ongewijzigd. Kaart 5 '
    '(Aantal of Som?) en kaart 10 (twee kolommen die niet naast elkaar staan '
    'selecteren met Ctrl) zijn deze keer de twee die ze echt nodig hebben.',
    'Wie deze twee bundels gemaakt heeft, kan een draaitabel op eender welke lijst: '
    'hun eigen uurrooster, de resultaten van een enquête, een ledenlijst.',
], BLAUW)

doc.save(SLEUTEL)
print('geschreven:', SLEUTEL)
