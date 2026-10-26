# -*- coding: utf-8 -*-
"""Bouwt de opdrachtbundel en de lerarensleutel van de tekstbundel.

Anders dan bij de draaitabelbundels staan de fiches hier IN de opdracht: het
is een nieuw onderwerp. Alleen de draaitabel en de draaigrafiek van blad 5
verwijzen naar het boekje dat ze al hebben.
"""
import importlib.util
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')
f = laad('f', 'fiches.py')

OPDRACHT = os.path.join(HIER, 'festival-tekst-opdracht.docx')
SLEUTEL = os.path.join(HIER, 'festival-tekst-lerarensleutel.docx')

BLAUW = RGBColor(0x1F, 0x38, 0x64)
ORANJE = RGBColor(0xB0, 0x4A, 0x00)
GRIJS = RGBColor(0x59, 0x59, 0x59)
GROEN = RGBColor(0x38, 0x76, 0x1D)

FORMULE = re.compile(r'(=[A-Z][A-Z.]*\((?:[^()]|\((?:[^()]|\([^()]*\))*\))*\)'
                     r'(?:&(?:"[^"]*"|[A-Z]+\d+))*'
                     r'|=[A-Z]+\d+(?:&(?:"[^"]*"|[A-Z]+\d+))+)')


def opzet(doc):
    s = doc.styles['Normal']
    s.font.name = 'Calibri'
    s.font.size = Pt(11)
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(1.5)
        sec.left_margin = sec.right_margin = Cm(1.8)


def kop(doc, tekst, niveau=1, kleur=BLAUW):
    p = doc.add_heading(tekst, level=niveau)
    for r in p.runs:
        r.font.color.rgb = kleur
    return p


def zin(p, tekst):
    """Formules in een vaste letter, zodat ze opvallen in lopende tekst."""
    for stuk in FORMULE.split(str(tekst)):
        if not stuk:
            continue
        r = p.add_run(stuk)
        if FORMULE.fullmatch(stuk):
            r.font.name = 'Consolas'
            r.font.size = Pt(10)
    return p


def vast(p, tekst):
    """Een heel stuk in de vaste letter, voor formules en voorbeelden."""
    r = p.add_run(str(tekst))
    r.font.name = 'Consolas'
    r.font.size = Pt(10)
    return r


def alinea(doc, tekst, vetgedrukt=False, schuin=False, kleur=None,
           grootte=None, na=None):
    p = doc.add_paragraph()
    if vetgedrukt or schuin or kleur or grootte:
        r = p.add_run(str(tekst))
        r.bold, r.italic = vetgedrukt, schuin
        if kleur:
            r.font.color.rgb = kleur
        if grootte:
            r.font.size = Pt(grootte)
    else:
        zin(p, tekst)
    if na is not None:
        p.paragraph_format.space_after = Pt(na)
    return p


def kader(doc, titel, regels, kleur=BLAUW):
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    c = t.rows[0].cells[0]
    r = c.paragraphs[0].add_run(titel)
    r.bold = True
    r.font.color.rgb = kleur
    for regel in regels:
        pp = c.add_paragraph()
        pp.paragraph_format.space_after = Pt(2)
        zin(pp, regel)
    doc.add_paragraph()
    return t


def tabel(doc, koppen, rijen, stijl='Light Grid Accent 1', breedtes=None,
          vaste_kolommen=()):
    t = doc.add_table(rows=1, cols=len(koppen))
    t.style = stijl
    for i, k in enumerate(koppen):
        t.rows[0].cells[i].paragraphs[0].add_run(str(k)).bold = True
    for rij in rijen:
        cellen = t.add_row().cells
        for i, w in enumerate(rij):
            if i in vaste_kolommen:
                vast(cellen[i].paragraphs[0], w)
            else:
                zin(cellen[i].paragraphs[0], w)
    if breedtes:
        for rij in t.rows:
            for i, b in enumerate(breedtes):
                rij.cells[i].width = Cm(b)
    doc.add_paragraph()
    return t


def paginaeinde(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def punten_van(oef):
    return o.punten_van(oef)


def fiche_uit(doc, nummer, kaart):
    """Eén fiche: wat het is, hoe je het schrijft, en de fout die iedereen maakt."""
    kop(doc, f"Fiche {nummer} — {kaart['naam']}", 2)
    alinea(doc, kaart['wat'], schuin=True, kleur=GRIJS, na=6)

    if kaart.get('schrijf'):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(6)
        vast(p, kaart['schrijf']).bold = True
        if kaart.get('argumenten'):
            tabel(doc, ['Stuk', 'Wat het is'],
                  [[stuk, uitleg] for stuk, uitleg in kaart['argumenten']],
                  'Light List Accent 1', [4.0, 12.0], vaste_kolommen={0})

    if kaart.get('stappen'):
        for i, stap in enumerate(kaart['stappen'], 1):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(f'{i}.  ')
            r.bold = True
            r.font.color.rgb = BLAUW
            zin(p, stap)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    for regel in kaart.get('uitleg', ()):
        alinea(doc, regel, na=4)

    if kaart.get('voorbeeld'):
        v = kaart['voorbeeld']
        breed = {2: [5.4, 10.6], 3: [3.6, 5.6, 6.8]}[len(v['koppen'])]
        tabel(doc, v['koppen'], v['rijen'], 'Light List Accent 1', breed,
              vaste_kolommen=set(range(len(v['koppen']))))

    if kaart.get('weten'):
        kader(doc, 'Goed om te weten', kaart['weten'], BLAUW)
    kader(doc, 'De fout die iedereen één keer maakt', [kaart['fout']], ORANJE)


# ==========================================================================
#  1. de opdrachtbundel
# ==========================================================================
doc = Document()
opzet(doc)
kop(doc, g.TITEL, 0)
alinea(doc, g.ONDERTITEL, schuin=True, kleur=GRIJS, na=4)
alinea(doc, 'Naam: ......................................    Klas: ..........',
       na=14)

alinea(doc,
       f'Vijf werkbladen, samen {g.DUUR} minuten. Op elk blad staat een lijst '
       'die niet deugt, en maak jij er een lijst van waar je iets mee kan. Dat '
       'is geen bijkomstigheid: op een werkvloer gaat het grootste deel van je '
       'Excel-tijd dáárin zitten.', na=8)
alinea(doc,
       f'De uitleg staat achteraan in deze bundel, {len(f.FICHES)} fiches. Bij '
       'elk blad staat welke je nodig hebt. Alleen voor de draaitabel en de '
       'draaigrafiek van blad 5 pak je het boekje erbij dat je al hebt: '
       f'{f.THEORIEBUNDEL}.', na=12)

tabel(doc, ['Blad', 'Wat je oefent', 'Fiches', 'Duur', 'Pt'],
      [[f"{oef['nr']}  {oef['blad'][2:]}", oef['titel'],
        ', '.join(str(n) for n in oef['fiches']), oef['duur'],
        punten_van(oef)] for oef in o.OEFENINGEN],
      breedtes=[3.6, 5.6, 3.2, 2.6, 1.6])

kader(doc, 'Drie regels die op elk blad gelden', [
    '1.  Je typt de formule één keer, in de bovenste cel, en voert ze door met '
    'de vulgreep. Zestien keer dezelfde formule typen is geen Excel.',
    '2.  Wat je uit een cel knipt is tekst, ook als er cijfers staan. Tekst gaat '
    'vanzelf links in de cel staan, een getal rechts. Kijk daarnaar voor je '
    'denkt dat je klaar bent.',
    '3.  De rommel in dit bestand staat er met opzet. Verbeter niets met de '
    'hand. Een lijst van zestien regels kan je nog met de hand opkuisen; een '
    'lijst van zestienduizend niet, en dáárvoor leer je dit.',
], ORANJE)
kader(doc, 'Over de gegevens', [g.WAARSCHUWING], GRIJS)

for oef in o.OEFENINGEN:
    paginaeinde(doc)
    kop(doc, f"Blad {oef['nr']} — {oef['titel']}", 1)
    alinea(doc, f"werkblad '{oef['blad']}'  ·  {oef['duur']}  ·  "
                f"{punten_van(oef)} punten", schuin=True, kleur=GRIJS, na=4)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run('Fiches:  ')
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = GRIJS
    stukken = [f"{n} — {f.FICHES[n - 1]['naam']}" for n in oef['fiches']]
    for n in oef.get('kaarten_elders', ()):
        stukken.append(f"kaart {n} uit {f.THEORIEBUNDEL}")
    r = p.add_run(' · '.join(stukken))
    r.font.size = Pt(10); r.font.color.rgb = ORANJE

    for regel in oef['inleiding']:
        alinea(doc, regel, na=4)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    if oef.get('vullen'):
        kop(doc, 'Wat je maakt', 2, GRIJS)
        tabel(doc, ['Waar', 'Wat erin komt', 'Pt'],
              [[waar, wat, punten] for waar, wat, _formule, punten
               in oef['vullen']],
              breedtes=[3.4, 11.6, 1.4])

    if oef.get('stappen'):
        kop(doc, 'Wat je doet', 2, GRIJS)
        tabel(doc, ['Stap', 'Hoe', 'Pt'],
              [[naam, hoe, punten] for naam, hoe, punten in oef['stappen']],
              breedtes=[4.4, 10.6, 1.4])

    kop(doc, 'Lees af en vul in', 2, GRIJS)
    tabel(doc, ['Cel', 'Vraag', 'Pt'],
          [[cel, wat[0].upper() + wat[1:] + '?', punten]
           for cel, wat, _antwoord, punten in oef['vragen']],
          breedtes=[2.4, 12.6, 1.4])

    kader(doc, 'Je blad is af als…', [oef['klaar'], '',
                                      'CONTROLE  ·  ' + oef['controle']], GROEN)
    kader(doc, 'Loop je vast?', oef['vast'], ORANJE)

# ---------------------------------------------------------------- de fiches
vorig_deel = None
for nummer, kaart in enumerate(f.FICHES, 1):
    paginaeinde(doc)
    if kaart['deel'] != vorig_deel:
        kop(doc, kaart['deel'], 1)
        vorig_deel = kaart['deel']
    fiche_uit(doc, nummer, kaart)

paginaeinde(doc)
kop(doc, 'Helemaal klaar?', 1)
alinea(doc, 'Loop dit lijstje af voor je indient.', schuin=True, kleur=GRIJS,
       na=8)
for regel in [
    'Op blad 1 tot 4 staat in elke te maken kolom een formule, geen getypte '
    'tekst.',
    'Nergens staat nog #WAARDE! of #NAAM?.',
    'Kolom E van blad 1 en kolom D van blad 3 staan rechts in de cel. Dat zijn '
    'echte getallen.',
    'In kolom D van blad 4 staat geen enkele spatie en geen enkele hoofdletter.',
    'Op blad 5 staat een tabel, een draaitabel, een draaigrafiek en een slicer.',
    'Je hebt elke vraag ingevuld.',
    'Je bestand heet festival-tekst-jouwnaam.xlsx.',
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(4)
    p.add_run('☐   ').font.size = Pt(13)
    zin(p, regel)

doc.save(OPDRACHT)
print('geschreven:', OPDRACHT)


# ==========================================================================
#  2. de lerarensleutel
# ==========================================================================
doc = Document()
opzet(doc)
kop(doc, g.TITEL + ' — lerarensleutel', 0)
alinea(doc, g.ONDERTITEL, schuin=True, kleur=GRIJS, na=12)

kop(doc, 'In het kort', 2, GRIJS)
alinea(doc,
       f'Vijf bladen, {g.DUUR} minuten, {o.TOTAAL} punten. Negen nieuwe '
       'functies, allemaal tekstfuncties, en op blad 5 geen enkele formule '
       'maar drie menuhandelingen: duplicaten verwijderen, Ctrl+T en een '
       'slicer.', na=6)
alinea(doc,
       'De fiches staan in de opdracht zelf — dit is een nieuw onderwerp, dus '
       'er is geen apart boekje. Wat ze wél moeten meebrengen is '
       f'{f.THEORIEBUNDEL}, voor de draaitabel en de draaigrafiek van blad 5.',
       na=6)
alinea(doc,
       'Blad 2 is het moeilijkste en staat met opzet vooraan: wie dat rond '
       'krijgt, heeft de rest ook. Blad 3 lijkt zwaarder maar is lichter — '
       'daar hoeft niets gezocht te worden, want elke code is even lang. Zeg '
       'dat erbij, anders schrikken ze.', na=10)

tabel(doc, ['Blad', 'Wat', 'Functies', 'Duur', 'Punten'],
      [[f"{oef['nr']}  {oef['blad'][2:]}", oef['titel'],
        ', '.join(f.FICHES[n - 1]['naam'] for n in oef['fiches']),
        oef['duur'], punten_van(oef)] for oef in o.OEFENINGEN]
      + [['Samen', '', '', f'{g.DUUR} minuten', o.TOTAAL]],
      breedtes=[3.2, 4.4, 5.4, 2.4, 2.0])

kader(doc, 'Waar het bij iedereen misgaat', [
    '1.  Ze rekenen op tekst. Een MIN of MAX op een kolom die uit LINKS of '
    'RECHTS komt geeft 0, en dat lijkt een fout in de formule terwijl het de '
    'soort is. Blad 1 is daar helemaal op gebouwd: kolom D en kolom E staan er '
    'naast elkaar zodat ze het verschil zien staan.',
    '2.  Ze tellen één teken verkeerd. Op blad 2 blijft er een spatie voor of '
    'na het antwoord staan, op blad 3 een streepje. Laat ze LENGTE ernaast '
    'zetten — dan zien ze het, want met het blote oog zie je een spatie niet.',
    '3.  Ze bouwen op blad 4 één lange formule in plaats van twee korte. '
    'Kolom C eerst, dan kolom D uit kolom C. Wie dat niet doet, krijgt een '
    'formule van tachtig tekens die hij zelf niet meer kan nakijken.',
    '4.  Op blad 5 maken ze eerst de draaitabel en kuisen ze pas daarna op. '
    f'Het eindtotaal is dan {len(g.B5_EXPORT)} in plaats van '
    f'{a.B5_DRAAI_TOTAAL}, en vaak merken ze het niet.',
], ORANJE)

kader(doc, 'Wat je niet automatisch kan nakijken', [
    'De tabel, de draaigrafiek en de slicer van blad 5 moet je zien. Een tabel '
    'herken je aan de kleur en de filterpijltjes, een draaigrafiek aan de '
    'knoppen met veldnamen erin, en een slicer aan het kader met knoppen.',
    'Snelle test voor de slicer: klik erop en kijk of de grafiek meebeweegt. '
    'Doet hij dat niet, dan is het een gewone grafiek op de tabel en geen '
    'draaigrafiek op de draaitabel.',
], BLAUW)

for oef in o.OEFENINGEN:
    paginaeinde(doc)
    kop(doc, f"Blad {oef['nr']} — {oef['titel']}", 1)
    alinea(doc, f"werkblad '{oef['blad']}'  ·  {oef['duur']}  ·  "
                f"{punten_van(oef)} punten", schuin=True, kleur=GRIJS, na=8)

    if oef.get('vullen'):
        kop(doc, 'De formules', 2, GRIJS)
        tabel(doc, ['Waar', 'Formule', 'Pt'],
              [[waar, formule, punten]
               for waar, _wat, formule, punten in oef['vullen']],
              breedtes=[3.0, 12.0, 1.4], vaste_kolommen={1})
        alinea(doc, 'Alleen de bovenste cel moet getypt zijn; de rest is '
                    'doorgevoerd. Staat er in elke cel een formule met hetzelfde '
                    'rijnummer, dan is ze gekopieerd in plaats van doorgevoerd '
                    'en klopt alleen de eerste regel.',
               schuin=True, kleur=GRIJS, na=10)

    if oef.get('stappen'):
        kop(doc, 'De stappen', 2, GRIJS)
        tabel(doc, ['Stap', 'Hoe', 'Pt'],
              [[naam, hoe, punten] for naam, hoe, punten in oef['stappen']],
              breedtes=[4.0, 11.0, 1.4])

    nr = oef['nr']
    if nr == 1:
        tabel(doc, ['Artiest', 'Code', 'Land', 'Jaar (tekst)', 'Jaar (getal)'],
              [[artiest, a.B1_CODE[i], a.B1_LAND[i], a.B1_JAAR_TEKST[i],
                a.B1_JAAR_GETAL[i]] for i, (artiest, _l, _j)
               in enumerate(g.B1)], 'Light List Accent 1')
        alinea(doc, 'Kolom D en E zien er identiek uit en zijn het niet. D '
                    'staat links in de cel, E rechts. Dat is de hele les van '
                    'dit blad.', schuin=True, kleur=GRIJS, na=10)
    elif nr == 2:
        tabel(doc, ['Regel', 'Podium', 'Artiest', 'Lengte'],
              [[a.B2_REGEL[i], a.B2_PODIUM[i], a.B2_ARTIEST[i], a.B2_LENGTE[i]]
               for i in range(len(g.B2))], 'Light List Accent 1')
        alinea(doc, f'De langste naam is {a.B2_LANGSTE} met '
                    f'{a.B2_LANGSTE_LENGTE} tekens. Krijgt een leerling '
                    f'{a.B2_LANGSTE_LENGTE + 1}, dan staat er nog een spatie '
                    'voor elke naam.', schuin=True, kleur=GRIJS, na=10)
    elif nr == 3:
        tabel(doc, ['Boekingscode', 'Dag', 'Podium', 'Volgnummer'],
              [[a.B3_CODE[i], a.B3_DAG[i], a.B3_PODIUMCODE[i],
                a.B3_VOLGNUMMER[i]] for i in range(len(g.B3))],
              'Light List Accent 1')
        alinea(doc, f'Elke code telt {a.B3_CODELENGTE} tekens. De dag begint '
                    f'op teken {a.B3_POS_DAG}, het podium op teken '
                    f'{a.B3_POS_PODIUM}.', schuin=True, kleur=GRIJS, na=10)
    elif nr == 4:
        tabel(doc, ['Ruw (· = spatie)', 'Naam', 'E-mailadres', 'Lengte'],
              [[f'{a._zichtbaar(ach)} / {a._zichtbaar(voor)}', a.B4_NAAM[i],
                a.B4_MAIL[i], a.B4_MAILLENGTE[i]]
               for i, (ach, voor) in enumerate(g.B4)], 'Light List Accent 1')
        alinea(doc, f'Het langste adres is {a.B4_LANGSTE_MAIL} met '
                    f'{a.B4_LANGSTE_LENGTE} tekens, van {a.B4_LANGSTE_NAAM}. '
                    'De drie dubbele achternamen krijgen twee punten in hun '
                    'adres; dat is correct, niet fout.',
               schuin=True, kleur=GRIJS, na=10)
    elif nr == 5:
        alinea(doc, f'De export telt {len(g.B5_EXPORT)} regels. Dubbel staan: '
                    f'{", ".join(a.B5_DUBBELE_NAMEN)}. Na Duplicaten '
                    f'verwijderen blijven er {a.B5_NA_OPKUISEN} over.', na=6)
        tabel(doc, ['Podium', 'Aantal acts',
                    f'Met de slicer op {g.B5_SLICER_KEUZE}'],
              [[p, a.B5_DRAAI[p], a.B5_DRAAI_GEFILTERD.get(p, 0)]
               for p in a.B5_DRAAI_LABELS]
              + [['Eindtotaal', a.B5_DRAAI_TOTAAL, a.B5_SLICER_AANTAL]],
              'Light List Accent 1', [5.0, 4.6, 6.0])
        alinea(doc, 'De draaigrafiek is een staafdiagram van de middelste '
                    'kolom. Zet de slicer aan en de balken worden die van de '
                    'rechterkolom — twee podia vallen dan helemaal weg. Laat '
                    'dat eens klassikaal zien, het is de pointe van de hele '
                    'les.', schuin=True, kleur=GRIJS, na=10)

    kop(doc, 'De antwoorden', 2, GRIJS)
    tabel(doc, ['Cel', 'Vraag', 'Antwoord', 'Pt'],
          [[cel, wat[0].upper() + wat[1:], antwoord, punten]
           for cel, wat, antwoord, punten in oef['vragen']],
          breedtes=[2.0, 8.6, 4.4, 1.4])
    kader(doc, 'Waar je naar kijkt', [oef['controle']], GROEN)

paginaeinde(doc)
kop(doc, 'Tijdsindeling en tips', 1)
tabel(doc, ['Wanneer', 'Wat'],
      [['0 - 5 min', 'Bestand opslaan onder eigen naam. Laat ze blad 4 even '
                     'openen en vraag wat er mis is met kolom A. Dat ze de '
                     'spaties niet zien is precies het punt.'],
       ['5 - 13 min', 'Blad 1. Gaat snel. Loop rond en kijk naar kolom D en E: '
                      'wie geen verschil ziet in de uitlijning, heeft de les '
                      'gemist.'],
       ['13 - 25 min', 'Blad 2. Het moeilijkste stuk. Doe de eerste regel '
                       'klassikaal op het bord, met de tekens geteld.'],
       ['25 - 35 min', 'Blad 3. Opluchting na blad 2. Wie hier vlot door gaat, '
                       'laat je al aan blad 4 beginnen.'],
       ['35 - 45 min', 'Blad 4. Hamer op het opbouwen in twee kolommen.'],
       ['45 - 50 min', 'Blad 5. Eerlijk: dit halen alleen de snelsten. Dat mag '
                       '— zie hieronder.']],
      breedtes=[3.4, 13.0])

kader(doc, 'Als de tijd te kort is', [
    'Blad 5 is het blad dat je het best over twee lessen spreidt. Het staat op '
    'zichzelf en heeft geen enkele formule nodig; je kan het ook als een losse '
    'les van twintig minuten geven, met de hele klas tegelijk aan het bord.',
    'Laat eventueel blad 3 vallen. DEEL is de makkelijkste van de drie '
    'knipfuncties en wie blad 2 kan, kan blad 3 ook.',
    'Wat je niet moet laten vallen is blad 1. Dat kolom E bestaat is de reden '
    'dat ze later begrijpen waarom hun VERT.ZOEKEN #N/B geeft.',
], BLAUW)

kader(doc, 'Wat hierna komt', [
    f'De zusterbundel {g.ZUSTERBUNDEL} werkt met dezelfde festivalgegevens, '
    'maar dan met de datums: hoeveel dagen nog, op welke weekdag, hoeveel '
    'werkdagen heeft de opbouwploeg. Daar eindigt blad 5 op een draaigrafiek '
    'met een tijdlijn — een slicer voor datums.',
    'Wie deze bundel af heeft, kan daar meteen aan beginnen.',
], BLAUW)

doc.save(SLEUTEL)
print('geschreven:', SLEUTEL)
