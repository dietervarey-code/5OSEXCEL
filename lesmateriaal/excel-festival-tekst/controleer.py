# -*- coding: utf-8 -*-
"""Kijkt na of gegevens, opdracht, startbestand en sleutel bij elkaar passen.

Bewaakt ook de oefening zelf: de rommel moet in het startbestand blijven
staan, want zonder rommel valt er niets op te kuisen.
"""
import collections
import importlib.util
import os
import re
import sys
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')
f = laad('f', 'fiches.py')

WERKBOEK = os.path.join(HIER, 'festival-tekst-startbestand.xlsx')
FOUTEN, GETELD = [], collections.Counter()


def eis(voorwaarde, bericht):
    if not voorwaarde:
        FOUTEN.append(bericht)
    return bool(voorwaarde)


def cel_van(adres):
    m = re.match(r'^([A-Z]{1,2})(\d+)$', adres)
    return (column_index_from_string(m.group(1)), int(m.group(2))) if m else None


# ==========================================================================
#  1. de gegevens
# ==========================================================================
_codes = [g.artiestcode(land, jaar) for _, land, jaar in g.B1]
eis(len({len(c) for c in _codes}) == 1,
    'de artiestcodes zijn niet allemaal even lang — dan werkt LINKS/RECHTS niet')
for artiest, land, jaar in g.B1:
    eis(land in g.B1_LANDEN, f'onbekend land {land!r} bij {artiest}')
    eis(1900 < jaar <= 2030, f'onwaarschijnlijk jaar bij {artiest}')
    GETELD['artiesten nagerekend'] += 1
eis(len({r[0] for r in g.B1}) == len(g.B1), 'dubbele artiestnaam op blad 1')

for podium, artiest in g.B2:
    regel = g.systeemregel(podium, artiest)
    eis(podium in g.B2_PODIA, f'onbekend podium {podium!r} bij {artiest}')
    eis(regel.count('|') == 1,
        f'{regel!r} bevat niet precies één streepje — dan is VIND.ALLES niet '
        'eenduidig')
    eis(podium.strip() == podium and artiest.strip() == artiest,
        f'{regel!r} heeft al spaties aan de rand; die horen er pas bij blad 4')
    GETELD['systeemregels nagerekend'] += 1
eis(len({len(p) for p in g.B2_PODIA}) > 1,
    'alle podiumnamen zijn even lang — dan kon blad 2 met LINKS en leert '
    'VIND.ALLES niets')

_bcodes = [g.boekingscode(d, p, v) for d, p, v in g.B3]
eis(len({len(c) for c in _bcodes}) == 1,
    'de boekingscodes zijn niet allemaal even lang — dan werkt DEEL niet')
eis(len({len(d) for d in g.B3_DAGEN}) == 1, 'de dagcodes zijn niet even lang')
eis(len({len(p) for p in g.B3_PODIUMCODES}) == 1,
    'de podiumcodes zijn niet even lang')
for dag, podium, nummer in g.B3:
    code = g.boekingscode(dag, podium, nummer)
    eis(dag in g.B3_DAGEN, f'onbekende dag {dag!r} in {code}')
    eis(podium in g.B3_PODIUMCODES, f'onbekende podiumcode {podium!r} in {code}')
    eis(0 < nummer < 1000, f'volgnummer {nummer} past niet in drie cijfers')
    # de formule uit de opdracht moet er écht uit halen wat erin zit
    eis(code[a.B3_POS_DAG - 1:a.B3_POS_DAG - 1 + a.B3_LEN_DAG] == dag,
        f'DEEL op positie {a.B3_POS_DAG} geeft de dag niet in {code}')
    eis(code[a.B3_POS_PODIUM - 1:a.B3_POS_PODIUM - 1 + a.B3_LEN_PODIUM] == podium,
        f'DEEL op positie {a.B3_POS_PODIUM} geeft het podium niet in {code}')
    staart = code[-a.B3_LEN_VOLGNUMMER:]
    eis(staart.isdigit() and int(staart) == nummer,
        f'de laatste {a.B3_LEN_VOLGNUMMER} tekens van {code} zijn {staart!r} '
        f'in plaats van het volgnummer {nummer:0{a.B3_LEN_VOLGNUMMER}d}')
    GETELD['boekingscodes nagerekend'] += 1
eis(len(set(_bcodes)) == len(_bcodes), 'twee dezelfde boekingscodes')

# de rommel op blad 4 moet er echt zijn, anders valt er niets op te kuisen
_rommelig = sum(1 for ach, voor in g.B4
                if ach != ach.strip() or voor != voor.strip())
eis(_rommelig >= len(g.B4) * 2 // 3,
    f'maar {_rommelig} van de {len(g.B4)} namen hebben losse spaties — te '
    'weinig om SPATIES.WISSEN nodig te maken')
_hoofdletters = sum(1 for ach, voor in g.B4
                    if ach.strip().isupper() or voor.strip().isupper())
eis(_hoofdletters >= len(g.B4) // 2,
    f'maar {_hoofdletters} namen staan in hoofdletters — te weinig om '
    'BEGINLETTERS nodig te maken')
_dubbele_achternaam = [ach for ach, _ in g.B4 if ' ' in ach.strip()]
eis(_dubbele_achternaam,
    'geen enkele achternaam met een spatie erin — dan is SUBSTITUEREN op '
    'blad 4 overbodig en leert het verschil met SPATIES.WISSEN niets')
for ach, voor in g.B4:
    eis(ach.strip() and voor.strip(), f'lege naam: {ach!r} / {voor!r}')
    GETELD['vrijwilligers nagerekend'] += 1
eis(len({(ach.strip().lower(), voor.strip().lower()) for ach, voor in g.B4})
    == len(g.B4), 'twee vrijwilligers met dezelfde naam')
eis(len(set(a.B4_MAIL)) == len(a.B4_MAIL), 'twee dezelfde e-mailadressen')

for artiest, dag, podium, genre in g.B5_SCHOON:
    eis(genre in g.B5_GENRES, f'onbekend genre {genre!r} bij {artiest}')
    eis(podium in g.B2_PODIA, f'onbekend podium {podium!r} bij {artiest}')
    GETELD['line-upregels nagerekend'] += 1
eis(len({r[0] for r in g.B5_SCHOON}) == len(g.B5_SCHOON),
    'twee keer dezelfde artiest in de schone line-up')
eis(a.B5_DUBBELE_REGELS == len(g.B5_DUBBEL),
    'de export bevat niet evenveel dubbels als B5_DUBBEL aangeeft')
eis(a.B5_DUBBELE_REGELS > 0,
    'er staat geen enkele regel dubbel — dan valt er niets te verwijderen')
eis(g.B5_SLICER_VELD in g.B5_KOP, 'de slicer staat op een veld dat niet bestaat')
eis(g.B5_SLICER_KEUZE in g.B5_GENRES, 'de slicerkeuze bestaat niet')
eis(g.B5_DRAAI_RIJEN in g.B5_KOP and g.B5_DRAAI_WAARDEN in g.B5_KOP,
    'de draaitabel gebruikt een veld dat niet in de lijst staat')

# ==========================================================================
#  2. de antwoorden, langs een tweede weg
# ==========================================================================
eis(a.B1_MEESTE_AANTAL == max(collections.Counter(
    r[1] for r in g.B1).values()), 'B1_MEESTE_AANTAL klopt niet')
eis(a.B1_VROEGSTE == min(r[2] for r in g.B1), 'B1_VROEGSTE klopt niet')
eis(a.B2_MEESTE_AANTAL == max(collections.Counter(
    p for p, _ in g.B2).values()), 'B2_MEESTE_AANTAL klopt niet')
eis(a.B2_LANGSTE_LENGTE == max(len(x) for _, x in g.B2),
    'B2_LANGSTE_LENGTE klopt niet')
eis(a.B3_HOOGSTE_VOLGNUMMER == max(v for _, _, v in g.B3),
    'B3_HOOGSTE_VOLGNUMMER klopt niet')
eis(a.B4_LANGSTE_LENGTE == max(len(m) for m in a.B4_MAIL),
    'B4_LANGSTE_LENGTE klopt niet')
eis(all(' ' not in m for m in a.B4_MAIL),
    'er staat een spatie in een e-mailadres')
eis(all(m == m.lower() for m in a.B4_MAIL),
    'er staat een hoofdletter in een e-mailadres')
eis(all(m.endswith('@' + g.DOMEIN) for m in a.B4_MAIL),
    'niet elk adres eindigt op het juiste domein')
eis(a.B5_DRAAI_TOTAAL == len(g.B5_SCHOON),
    'het eindtotaal van de draaitabel is niet het aantal schone regels')
eis(a.B5_DRAAI_TOTAAL != len(g.B5_EXPORT),
    'het eindtotaal is even groot met en zonder opkuisen — dan merkt niemand '
    'of hij de dubbels verwijderd heeft')
eis(0 < a.B5_SLICER_AANTAL < a.B5_DRAAI_TOTAAL,
    f'de slicer op {g.B5_SLICER_KEUZE} laat {a.B5_SLICER_AANTAL} van de '
    f'{a.B5_DRAAI_TOTAAL} acts over — dan is er niets aan te zien')
eis(len(a.B5_GEFILTERD_LABELS) < len(a.B5_DRAAI_LABELS),
    'met de slicer aan blijven alle podia staan — dan verandert de grafiek '
    'nauwelijks zichtbaar')

for oef in o.OEFENINGEN:
    for cel, wat, antwoord, punten in oef['vragen']:
        eis(antwoord is not None and antwoord != '',
            f"{oef['blad']} {cel}: geen antwoord")
        eis(punten > 0, f"{oef['blad']} {cel}: geen punten")
        GETELD['vragen gekeurd'] += 1

_tekst = collections.defaultdict(list)
for oef in o.OEFENINGEN:
    for cel, _wat, antwoord, _p in oef['vragen']:
        if isinstance(antwoord, str):
            _tekst[antwoord].append(f"{oef['blad']} {cel}")
for antwoord, plekken in _tekst.items():
    eis(len(plekken) == 1,
        f'het antwoord {antwoord!r} staat op {len(plekken)} plaatsen: '
        f'{", ".join(plekken)}')

# ==========================================================================
#  3. de opdracht
# ==========================================================================
FUNCTIE = re.compile(r'\b([A-Z][A-Z.]{2,})\(')
geleerd = {naam for kaart in f.FICHES for naam in kaart['functies']}
gebruikt = set()
for oef in o.OEFENINGEN:
    for waar, _wat, formule, _p in oef.get('vullen', []):
        gebruikt |= set(FUNCTIE.findall(formule))
eis(gebruikt <= geleerd,
    f'de opdracht gebruikt functies zonder fiche: {sorted(gebruikt - geleerd)}')
for kaart in f.FICHES:
    if kaart['functies']:
        eis(set(kaart['functies']) & gebruikt,
            f"fiche {kaart['naam']} leert iets dat nergens gebruikt wordt")
GETELD['functies gekeurd'] += len(geleerd)

verwezen = {n for oef in o.OEFENINGEN for n in oef['fiches']}
eis(verwezen == set(range(1, len(f.FICHES) + 1)),
    f'niet naar elke fiche wordt verwezen: '
    f'{sorted(set(range(1, len(f.FICHES) + 1)) - verwezen)}')
for kaart in f.FICHES:
    eis(kaart['deel'] in f.DELEN, f"onbekend deel bij {kaart['naam']}")
    eis(kaart.get('wat') and kaart.get('fout'),
        f"fiche {kaart['naam']} is niet volledig")
    eis(kaart.get('schrijf') or kaart.get('stappen'),
        f"fiche {kaart['naam']} zegt niet hoe het moet")
    GETELD['fiches gekeurd'] += 1
volgorde = [k['deel'] for k in f.FICHES]
eis(volgorde == sorted(volgorde, key=f.DELEN.index),
    'de fiches van de drie delen staan door elkaar')
for n in {n for oef in o.OEFENINGEN for n in oef.get('kaarten_elders', ())}:
    eis(n in f.KAARTEN_ELDERS, f'kaart {n} van het andere boekje is onbekend')

# elke formule moet in de bovenste rij van haar eigen bereik staan: anders
# klopt het doorvoeren niet
for oef in o.OEFENINGEN:
    for waar, _wat, formule, _p in oef.get('vullen', []):
        eerste = waar.split(':')[0]
        _kolom, rij = cel_van(eerste)
        verwijzingen = set(int(x) for x in re.findall(r'[A-Z]+(\d+)', formule))
        eis(verwijzingen == {rij},
            f'{oef["blad"]} {waar}: de formule verwijst naar rij '
            f'{sorted(verwijzingen)} in plaats van naar rij {rij} — dan loopt '
            'het doorvoeren mis')
        eis(f'{get_column_letter(_kolom)}{rij}' not in formule,
            f'{oef["blad"]} {waar}: de formule verwijst naar zichzelf')
        GETELD['formules gekeurd'] += 1

eis(sum(int(oef['duur'].split()[0]) for oef in o.OEFENINGEN) == g.DUUR,
    'de duur van de bladen telt niet op tot de aangekondigde duur')
eis(o.TOTAAL == 50, f'de bundel staat op {o.TOTAAL} punten in plaats van 50')

# ==========================================================================
#  4. het startbestand
# ==========================================================================
LIJSTEN = [
    ('1 Artiestcodes', g.B1_KOPRIJ, g.B1_START, g.B1_EIND, g.B1_KOP, o.V1),
    ('2 Podia', g.B2_KOPRIJ, g.B2_START, g.B2_EIND, g.B2_KOP, o.V2),
    ('3 Boekingscodes', g.B3_KOPRIJ, g.B3_START, g.B3_EIND, g.B3_KOP, o.V3),
    ('4 Vrijwilligers', g.B4_KOPRIJ, g.B4_START, g.B4_EIND, g.B4_KOP, o.V4),
    ('5 Line-up', g.B5_KOPRIJ, g.B5_START, g.B5_EIND, g.B5_KOP, o.V5),
]
eis({x[0] for x in LIJSTEN} == {oef['blad'] for oef in o.OEFENINGEN},
    'de bladen van LIJSTEN en van OEFENINGEN lopen uit elkaar')

eis(os.path.exists(WERKBOEK), 'het startbestand bestaat niet — draai eerst '
                              'maak-startbestand.py')
if os.path.exists(WERKBOEK):
    wb = load_workbook(WERKBOEK)
    eis({x[0] for x in LIJSTEN} <= set(wb.sheetnames), 'ontbrekende bladen')

    for blad, koprij, start, eind, koppen, vraagkol in LIJSTEN:
        if blad not in wb.sheetnames:
            continue
        ws = wb[blad]
        breedte = len(koppen)
        for i, titel in enumerate(koppen, start=1):
            eis(ws.cell(row=koprij, column=i).value == titel,
                f'{blad}: koprij verschilt in kolom {get_column_letter(i)}')
        eis(len(set(koppen)) == breedte, f'{blad}: twee kolommen heten hetzelfde')

        temaken = set()
        for oef in o.OEFENINGEN:
            if oef['blad'] != blad:
                continue
            for waar, *_rest in oef.get('vullen', []):
                temaken.add(column_index_from_string(
                    re.match(r'([A-Z]+)', waar).group(1)))
        for kolom in range(1, breedte + 1):
            heeft = any(ws.cell(row=r, column=kolom).value is not None
                        for r in range(start, eind + 1))
            eis(heeft != (kolom in temaken),
                f'{blad}: kolom {get_column_letter(kolom)} is '
                f'{"gevuld" if heeft else "leeg"} terwijl ze '
                f'{"leeg" if kolom in temaken else "gevuld"} hoort te zijn')
        for rij in range(start, eind + 1):
            eis(any(ws.cell(row=rij, column=c).value is not None
                    for c in range(1, breedte + 1)),
                f'{blad}: rij {rij} is helemaal leeg')

        eis(vraagkol >= breedte + 2,
            f'{blad}: het vragenblok staat in kolom '
            f'{get_column_letter(vraagkol)}, de lijst loopt tot '
            f'{get_column_letter(breedte)} — er hoort een lege kolom tussen')
        for rij in range(start, eind + 1):
            eis(ws.cell(row=rij, column=breedte + 1).value is None,
                f'{blad}: kolom {get_column_letter(breedte + 1)} rij {rij} '
                'hoort leeg te blijven als scheiding')
        GETELD['lijsten gekeurd'] += 1

    # er zit nog geen tabel, geen grafiek, geen draaitabel in
    for ws in wb.worksheets:
        eis(not list(ws.tables),
            f'{ws.title}: er staat al een tabel in het startbestand')
        eis(not getattr(ws, '_charts', None),
            f'{ws.title}: er staat al een grafiek in het startbestand')
        eis(not getattr(ws, '_pivots', None),
            f'{ws.title}: er staat al een draaitabel in het startbestand')

    for oef in o.OEFENINGEN:
        ws = wb[oef['blad']]
        for cel, *_rest in oef['vragen']:
            kolom, rij = cel_van(cel)
            eis(ws.cell(row=rij, column=kolom).value is None,
                f"{oef['blad']} {cel} moet leeg zijn in het startbestand")
            eis(ws.cell(row=rij, column=kolom - 1).value is not None,
                f"{oef['blad']} {cel}: er staat geen vraag naast")
            GETELD['antwoordcellen leeg'] += 1
        for waar, *_rest in oef.get('vullen', []):
            k1, r1 = cel_van(waar.split(':')[0])
            k2, r2 = cel_van(waar.split(':')[1])
            for k in range(k1, k2 + 1):
                for r in range(r1, r2 + 1):
                    eis(ws.cell(row=r, column=k).value is None,
                        f"{oef['blad']} {get_column_letter(k)}{r} moet leeg zijn")
                    GETELD['te vullen cellen leeg'] += 1

    # de rommel moet bewaard zijn gebleven bij het schrijven
    ws = wb['4 Vrijwilligers']
    bewaard = sum(1 for i in range(len(g.B4))
                  if ws.cell(row=g.B4_START + i, column=1).value != g.B4[i][0].strip()
                  or ws.cell(row=g.B4_START + i, column=2).value != g.B4[i][1].strip())
    eis(bewaard == _rommelig,
        f'{bewaard} van de {_rommelig} rommelige namen staan nog met hun '
        'spaties in het startbestand — openpyxl heeft ze opgekuist')
    for i, (ach, voor) in enumerate(g.B4):
        eis(ws.cell(row=g.B4_START + i, column=1).value == ach
            and ws.cell(row=g.B4_START + i, column=2).value == voor,
            f'4 Vrijwilligers: rij {g.B4_START + i} wijkt af van gegevens.py')
        GETELD['ruwe namen nagekeken'] += 1

    ws = wb['5 Line-up']
    rijen = [tuple(ws.cell(row=r, column=c).value for c in range(1, 5))
             for r in range(g.B5_START, g.B5_EIND + 1)]
    eis(len(rijen) - len(set(rijen)) == a.B5_DUBBELE_REGELS,
        f'de export in het startbestand telt {len(rijen) - len(set(rijen))} '
        f'dubbele regels in plaats van {a.B5_DUBBELE_REGELS}')
    eis(len(set(rijen)) == a.B5_NA_OPKUISEN,
        'na het opkuisen blijven er niet evenveel regels over als de sleutel '
        'zegt')
    GETELD['exportregels nagekeken'] += len(rijen)

# ==========================================================================
#  5. de documenten zijn gebouwd
# ==========================================================================
for naam in ('festival-tekst-opdracht.docx', 'festival-tekst-lerarensleutel.docx'):
    eis(os.path.exists(os.path.join(HIER, naam)),
        f'{naam} bestaat niet — draai eerst maak-documenten.py')

# ==========================================================================
print('=' * 72)
for naam, hoeveel in sorted(GETELD.items()):
    print(f'  {hoeveel:>5}  {naam}')
print('=' * 72)
if FOUTEN:
    print(f'{len(FOUTEN)} PROBLEMEN:')
    for fout in FOUTEN:
        print('  !', fout)
    sys.exit(1)
print('alles past bij elkaar.')
