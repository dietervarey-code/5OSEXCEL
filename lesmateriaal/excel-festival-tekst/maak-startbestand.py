# -*- coding: utf-8 -*-
"""Bouwt het startbestand van de tekstbundel.

De rommel moet erin blijven: de spaties op blad 4 en de dubbele regels op
blad 5 zijn geen slordigheid maar de oefening. controleer.py kijkt na dat ze
er nog in staan.

Er zit geen tabel, geen draaitabel, geen grafiek en geen slicer in. Die maakt
de leerling zelf.
"""
import importlib.util
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')

UIT = os.path.join(HIER, 'festival-tekst-startbestand.xlsx')
TITEL = Font(bold=True, size=13)
GROOT = Font(bold=True, size=16)
VET = Font(bold=True)
SCHUIN = Font(italic=True, size=10)

wb = Workbook()


def breedtes(ws, waarden, start=1):
    for i, b in enumerate(waarden, start=start):
        ws.column_dimensions[get_column_letter(i)].width = b


def blad(naam, titel, ondertitel=None):
    ws = wb.create_sheet(naam)
    ws['A1'] = titel
    ws['A1'].font = TITEL
    if ondertitel:
        ws['A2'] = ondertitel
        ws['A2'].font = SCHUIN
    return ws


def lijst(ws, koprij, koppen, rijen):
    for i, k in enumerate(koppen, start=1):
        ws.cell(row=koprij, column=i, value=k).font = VET
    for r, waarden in enumerate(rijen, start=koprij + 1):
        for i, w in enumerate(waarden, start=1):
            ws.cell(row=r, column=i, value=w)


def vragen(ws, rij, kolom, teksten):
    ws.cell(row=rij, column=kolom, value='Vraag').font = VET
    ws.cell(row=rij, column=kolom + 1, value='Antwoord').font = VET
    for i, t in enumerate(teksten, start=1):
        ws.cell(row=rij + i, column=kolom, value=t)


# ------------------------------------------------------------------- Start
ws = wb.active
ws.title = 'Start'
ws['A1'] = g.TITEL
ws['A1'].font = GROOT
KOP_REGELS = 'Drie regels die op elk blad gelden'
regels = [
    g.ONDERTITEL, '',
    'Naam: ......................................    Klas: ..........', '',
    f'Vijf werkbladen, samen {g.DUUR} minuten.',
    'De uitleg staat in de opdrachtbundel, bij elk blad staat welke fiches je',
    'erbij nodig hebt. Alleen voor de draaitabel en de draaigrafiek van blad 5',
    'pak je het boekje erbij dat je al hebt.', '',
]
for oef in o.OEFENINGEN:
    fiches = 'fiche ' + ', '.join(str(n) for n in oef['fiches'])
    regels.append(f"Blad {oef['nr']}  {oef['blad'][2:]:<15} "
                  f"{oef['duur'].split()[0]:>2} min   {fiches:<22} "
                  f"{oef['titel']}")
regels += [
    '', KOP_REGELS,
    '  1. Je typt de formule één keer, in de bovenste cel, en voert ze door',
    '     met de vulgreep. Zestien keer typen is geen Excel.',
    '  2. Wat je uit een cel knipt is tekst, ook als er cijfers staan. Tekst',
    '     gaat links in de cel staan, een getal rechts. Kijk daarnaar.',
    '  3. De rommel in dit bestand is er met opzet. Verbeter niets met de',
    '     hand — je lost het op met een formule, anders leer je er niets van.',
    '',
    'EERST DOEN: sla dit bestand op als  festival-tekst-jouwnaam.xlsx',
    '', g.WAARSCHUWING,
]
for i, r in enumerate(regels, start=2):
    ws.cell(row=i, column=1, value=r)
ws.cell(row=2 + regels.index(KOP_REGELS), column=1).font = VET
rij_waarschuwing = 2 + regels.index(g.WAARSCHUWING)
c = ws.cell(row=rij_waarschuwing, column=1)
c.font = SCHUIN
c.alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[rij_waarschuwing].height = 86
breedtes(ws, [92])

# ----------------------------------------------------------- 1 Artiestcodes
oef = o.OEFENINGEN[0]
ws = blad(oef['blad'], 'De artiesten en hun code', g.B1_VOETNOOT)
lijst(ws, g.B1_KOPRIJ, g.B1_KOP,
      [(artiest, g.artiestcode(land, jaar), None, None, None)
       for artiest, land, jaar in g.B1])
vragen(ws, g.B1_KOPRIJ, o.V1, [
    'Van welk land komen de meeste acts?',
    'Wat is het vroegste jaar in de lijst?'])
breedtes(ws, [22, 13, 9, 13, 13])
breedtes(ws, [3, 36, 16], start=o.V1 - 1)

# ------------------------------------------------------------------ 2 Podia
oef = o.OEFENINGEN[1]
ws = blad(oef['blad'], 'De regels uit het boekingssysteem', g.B2_VOETNOOT)
lijst(ws, g.B2_KOPRIJ, g.B2_KOP,
      [(g.systeemregel(podium, artiest), None, None, None)
       for podium, artiest in g.B2])
vragen(ws, g.B2_KOPRIJ, o.V2, [
    'Welk podium heeft de meeste acts?',
    f'Hoeveel acts staan er op {a.B2_TELPODIUM}?',
    'Hoeveel tekens telt de langste artiestennaam?'])
breedtes(ws, [30, 13, 20, 18])
breedtes(ws, [3, 38, 16], start=o.V2 - 1)

# ---------------------------------------------------------- 3 Boekingscodes
oef = o.OEFENINGEN[2]
ws = blad(oef['blad'], 'De boekingscodes', g.B3_VOETNOOT)
lijst(ws, g.B3_KOPRIJ, g.B3_KOP,
      [(g.boekingscode(dag, podium, nummer), None, None, None)
       for dag, podium, nummer in g.B3])
vragen(ws, g.B3_KOPRIJ, o.V3, [
    'Op welke dag staan de meeste acts?',
    'Welk podium heeft de minste acts?',
    'Wat is het hoogste volgnummer?'])
breedtes(ws, [20, 8, 14, 13])
breedtes(ws, [3, 34, 16], start=o.V3 - 1)

# ---------------------------------------------------------- 4 Vrijwilligers
oef = o.OEFENINGEN[3]
ws = blad(oef['blad'], 'De vrijwilligers', g.B4_VOETNOOT)
lijst(ws, g.B4_KOPRIJ, g.B4_KOP,
      [(achternaam, voornaam, None, None, None)
       for achternaam, voornaam in g.B4])
vragen(ws, g.B4_KOPRIJ, o.V4, [
    'Hoeveel tekens telt het langste e-mailadres?',
    'Van wie is dat adres?'])
breedtes(ws, [18, 17, 22, 34, 9])
breedtes(ws, [3, 40, 20], start=o.V4 - 1)

# --------------------------------------------------------------- 5 Line-up
oef = o.OEFENINGEN[4]
ws = blad(oef['blad'], 'De volledige line-up', g.B5_VOETNOOT)
lijst(ws, g.B5_KOPRIJ, g.B5_KOP, g.B5_EXPORT)
vragen(ws, g.B5_KOPRIJ, o.V5, [
    'Hoeveel regels gooit Excel weg bij Duplicaten verwijderen?',
    'Wat is het eindtotaal van je draaitabel?',
    f'Hoeveel acts blijven er over met de slicer op {g.B5_SLICER_KEUZE}?'])
breedtes(ws, [22, 11, 13, 11])
breedtes(ws, [3, 48, 16], start=o.V5 - 1)

wb.save(UIT)
print('geschreven:', UIT)
for ws in wb.worksheets:
    gevuld = sum(1 for rij in ws.iter_rows() for c in rij if c.value is not None)
    print(f'  {ws.title:<18} {gevuld:>4} gevulde cellen   {ws.dimensions}')
