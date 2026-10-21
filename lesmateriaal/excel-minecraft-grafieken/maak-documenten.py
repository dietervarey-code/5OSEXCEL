# -*- coding: utf-8 -*-
"""Bouwt de opdrachtbundel en de lerarensleutel van de Minecraft-grafieken.

Zelfde opzet als de draaitabelbundel: geen theorie in de opdracht, elk
blad verwijst naar de kaarten.
"""
import importlib.util
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')

OPDRACHT = os.path.join(HIER, 'minecraft-grafieken-opdracht.docx')
SLEUTEL = os.path.join(HIER, 'minecraft-grafieken-lerarensleutel.docx')

BLAUW = RGBColor(0x1F, 0x38, 0x64)
ORANJE = RGBColor(0xB0, 0x4A, 0x00)
GRIJS = RGBColor(0x59, 0x59, 0x59)
GROEN = RGBColor(0x38, 0x76, 0x1D)

FORMULE = re.compile(r"(=[A-Z.]*\((?:[^()]|\([^()]*\))*\)|=[A-Z$]+\d+(?:[*+\-/]"
                     r"[A-Z$]*\d+)*)")


def opzet(doc):
    s = doc.styles['Normal']
    s.font.name = 'Calibri'
    s.font.size = Pt(11)
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(1.5)
        sec.left_margin = sec.right_margin = Cm(1.8)


def kop(doc, tekst, niveau=1, kleur=BLAUW):
    p = doc.add_heading(tekst, level=niveau)
    for r in p.runs:
        r.font.color.rgb = kleur
    return p


def zin(p, tekst):
    """Zet formules in een vaste letter, zodat ze opvallen in lopende tekst."""
    for stuk in FORMULE.split(str(tekst)):
        if not stuk:
            continue
        r = p.add_run(stuk)
        if FORMULE.fullmatch(stuk):
            r.font.name = 'Consolas'
            r.font.size = Pt(10)
    return p


def alinea(doc, tekst, style=None, vet=False, schuin=False, kleur=None,
           grootte=None, na=None):
    p = doc.add_paragraph(style=style)
    if vet or schuin or kleur or grootte:
        r = p.add_run(str(tekst))
        r.bold, r.italic = vet, schuin
        if kleur:
            r.font.color.rgb = kleur
        if grootte:
            r.font.size = Pt(grootte)
    else:
        zin(p, tekst)
    if na is not None:
        p.paragraph_format.space_after = Pt(na)
    return p


def kader(doc, titel, regels, kleur=BLAUW):
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    c = t.rows[0].cells[0]
    r = c.paragraphs[0].add_run(titel)
    r.bold = True
    r.font.color.rgb = kleur
    for regel in regels:
        pp = c.add_paragraph()
        pp.paragraph_format.space_after = Pt(2)
        zin(pp, regel)
    doc.add_paragraph()
    return t


def tabel(doc, koppen, rijen, stijl='Light Grid Accent 1', breedtes=None):
    t = doc.add_table(rows=1, cols=len(koppen))
    t.style = stijl
    for i, k in enumerate(koppen):
        t.rows[0].cells[i].paragraphs[0].add_run(str(k)).bold = True
    for rij in rijen:
        cellen = t.add_row().cells
        for i, w in enumerate(rij):
            zin(cellen[i].paragraphs[0], w)
    if breedtes:
        for rij in t.rows:
            for i, b in enumerate(breedtes):
                rij.cells[i].width = Cm(b)
    doc.add_paragraph()
    return t


def paginaeinde(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def genummerd(doc, regels, kleur=BLAUW):
    for i, regel in enumerate(regels, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f'{i}.  ')
        r.bold = True
        r.font.color.rgb = kleur
        zin(p, regel)



def kaartregel(doc, oef):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run('Theorie:  ')
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = GRIJS
    r = p.add_run(' · '.join(f"kaart {n} — {o.KAARTEN[n]}" for n in oef['kaarten']))
    r.font.size = Pt(10); r.font.color.rgb = ORANJE
    return p


def grafiekspec(gr):
    rijen = [['Type', gr['soort']],
             ['Selecteer', gr['selectie']],
             ['Titel', f"“{gr['titel']}”"]]
    if gr.get('astitel_x'):
        rijen.append(['Astitels', f"onder “{gr['astitel_x']}”, "
                                  f"links “{gr['astitel_y']}”"])
    rijen.append(['Legende', gr['legende']])
    if gr.get('labels'):
        rijen.append(['Gegevenslabels', gr['labels']])
    if gr.get('vanaf_nul'):
        rijen.append(['Verticale as', 'begint bij 0'])
    return rijen


def punten_van(oef):
    return (oef['punten_grafiek'] + sum(p for *_, p in oef['vragen'])
            + sum(p for *_, p in oef.get('vullen', [])))


# ==========================================================================
#  1. de opdrachtbundel
# ==========================================================================
doc = Document()
opzet(doc)
kop(doc, g.TITEL, 0)
alinea(doc, g.ONDERTITEL, schuin=True, kleur=GRIJS, na=4)
alinea(doc, 'Naam: ......................................    Klas: ..........', na=14)
alinea(doc,
       f'Vijf werkbladen, samen {g.DUUR} minuten. Op elk blad staat een lijst klaar '
       'en maak je er één grafiek van. Telkens een ander type, want telkens een '
       'andere vraag.', na=8)
alinea(doc,
       'In deze bundel staat geen theorie. Bij elk blad staat welke kaarten uit '
       f'{g.THEORIEBUNDEL} je erbij nodig hebt — elke kaart staat daar op een eigen '
       'pagina.', na=12)

tabel(doc, ['Blad', 'Wat je oefent', 'Type', 'Theorie', 'Duur', 'Pt'],
      [[f"{oef['nr']}  {oef['blad'][2:]}", oef['titel'], oef['grafiek']['soort'],
        ', '.join(f"kaart {n}" for n in oef['kaarten']), oef['duur'],
        punten_van(oef)] for oef in o.OEFENINGEN],
      breedtes=[3.2, 4.0, 4.2, 3.4, 2.2, 1.4])

kader(doc, 'Drie regels die op elk blad gelden', [
    '1.  Het type grafiek is geen smaakkwestie. Vergelijken is een staafdiagram, '
    'een verloop is een lijn, een aandeel van een geheel is een cirkel. Een lijn '
    'mag alleen als er tussen twee punten iets bestaat: tussen twee jaren wel, '
    'tussen steen en ijzer niet. Bij elk blad staat welke vraag je beantwoordt.',
    '2.  Elke grafiek krijgt een titel die zegt wat je ziet, en astitels zodra de '
    'eenheid niet vanzelf spreekt. "Grafiektitel" laten staan is geen titel.',
    '3.  Laat je verticale as bij nul beginnen, tenzij er uitdrukkelijk iets anders '
    'staat.',
], ORANJE)
kader(doc, 'Over de cijfers', [g.WAARSCHUWING], GRIJS)

for oef in o.OEFENINGEN:
    paginaeinde(doc)
    kop(doc, f"Blad {oef['nr']} — {oef['titel']}", 1)
    alinea(doc, f"werkblad '{oef['blad']}'  ·  {oef['duur']}", schuin=True,
           kleur=GRIJS, na=4)
    kaartregel(doc, oef)
    for regel in oef['inleiding']:
        alinea(doc, regel, na=4)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    if oef.get('vullen'):
        kop(doc, 'Eerst: de kolom aanvullen', 2, GRIJS)
        tabel(doc, ['Waar', 'Wat', 'Pt'],
              [[waar, wat, punten] for waar, wat, _f, punten in oef['vullen']],
              breedtes=[3.4, 11.6, 1.4])

    kop(doc, 'Dan: de grafiek', 2, GRIJS)
    tabel(doc, ['Wat', 'Hoe'], grafiekspec(oef['grafiek']), breedtes=[3.6, 12.8])

    kop(doc, 'Lees af en vul in', 2, GRIJS)
    tabel(doc, ['Cel', 'Vraag', 'Pt'],
          [[cel, wat[0].upper() + wat[1:] + '?', punten]
           for cel, wat, _antwoord, punten in oef['vragen']],
          breedtes=[2.4, 12.6, 1.4])

    kader(doc, 'Je blad is af als…', [oef['klaar'], '',
                                      'CONTROLE  ·  ' + oef['controle']], GROEN)
    kader(doc, 'Loop je vast?', oef['vast'], ORANJE)

paginaeinde(doc)
kop(doc, 'Helemaal klaar?', 1)
alinea(doc, 'Loop dit lijstje af voor je indient.', schuin=True, kleur=GRIJS, na=8)
for regel in [
    'Op elk van de vijf bladen staat een grafiek.',
    'Elke grafiek heeft een titel die zegt wat je ziet.',
    'Nergens staat nog "Grafiektitel", "Reeks1" of #N/B.',
    'Je verticale as begint bij nul, behalve waar iets anders gevraagd is.',
    'Je hebt elke vraag ingevuld.',
    'Je bestand heet minecraft-grafieken-jouwnaam.xlsx.',
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(4)
    p.add_run('☐   ').font.size = Pt(13)
    zin(p, regel)

doc.save(OPDRACHT)
print('geschreven:', OPDRACHT)

# ==========================================================================
#  2. de lerarensleutel
# ==========================================================================
doc = Document()
opzet(doc)
kop(doc, g.TITEL + ' — lerarensleutel', 0)
alinea(doc, g.ONDERTITEL, schuin=True, kleur=GRIJS, na=12)

kop(doc, 'In het kort', 2, GRIJS)
alinea(doc,
       f'Vijf bladen, {g.DUUR} minuten, {o.TOTAAL} punten. Vier grafiektypes, en '
       'bij elk blad de vraag die dat type beantwoordt. Op blad 2, 3 en 5 moeten ze '
       'eerst een kolom aanvullen; de rest is tekenen en aflezen.', na=6)
alinea(doc,
       f'De zusterbundel {g.ZUSTERBUNDEL} gebruikt hetzelfde materiaal voor '
       'draaitabellen. Wie die af heeft, herkent de beroepentabel van blad 5 '
       'meteen.', na=6)
alinea(doc,
       'Dit is dezelfde bundel als de Türkiye-versie, met ander materiaal — op '
       'één punt na: blad 4 is hier een staafdiagram met drie reeksen in plaats '
       'van een lijndiagram, omdat de horizontale as geen tijd is maar een rijtje '
       'materialen. Dat verschil is zelf een les, en het staat ook in de '
       'opdracht.', na=10)

tabel(doc, ['Blad', 'Type', 'Kaarten', 'Duur', 'Punten'],
      [[f"{oef['nr']}  {oef['blad'][2:]}", oef['grafiek']['soort'],
        ', '.join(str(n) for n in oef['kaarten']), oef['duur'], punten_van(oef)]
       for oef in o.OEFENINGEN]
      + [['Samen', '', '', f'{g.DUUR} minuten', o.TOTAAL]],
      breedtes=[3.4, 5.4, 2.8, 2.6, 2.2])

kader(doc, 'Wat je niet automatisch kan nakijken', [
    'Of een grafiek een titel heeft en of de as bij nul begint, moet je zien. '
    'Hieronder staat per blad wat je verwacht, zodat je het in één blik kan '
    'afvinken.',
    'De antwoordcellen kan je wél gewoon nalezen; die staan hieronder uitgeschreven.',
], BLAUW)

kader(doc, 'Waar het bij iedereen misgaat', [
    '1.  Ze selecteren de verkeerde kolommen. Op blad 1 en 5 staan de twee nodige '
    'kolommen NIET naast elkaar — daar moet Ctrl aan te pas komen. Krijgen ze '
    'dimensies of beroepcodes onder hun balken, dan is dat de oorzaak.',
    '2.  De koprij wordt niet meegeselecteerd, en de legende toont Reeks1, Reeks2, '
    'Reeks3. Dat zie je het snelst op blad 4.',
    f'3.  Op blad 4 is de {a.B4_VLAK[0].lower()} bijna onzichtbaar naast de bijl. '
    'Dat is geen fout: 1 naast 10 is nu eenmaal niets. Het is juist het gesprek '
    'dat je wil hebben — wat verdwijnt er in een grafiek met reeksen van heel '
    'verschillende grootte?',
], ORANJE)

for oef in o.OEFENINGEN:
    paginaeinde(doc)
    kop(doc, f"Blad {oef['nr']} — {oef['titel']}", 1)
    alinea(doc, f"werkblad '{oef['blad']}'  ·  {oef['duur']}  ·  "
                f"{punten_van(oef)} punten  ·  "
                f"kaart {', '.join(str(n) for n in oef['kaarten'])}",
           schuin=True, kleur=GRIJS, na=8)

    if oef.get('vullen'):
        kop(doc, 'De formule', 2, GRIJS)
        tabel(doc, ['Waar', 'Formule', 'Pt'],
              [[waar, formule, punten]
               for waar, _wat, formule, punten in oef['vullen']],
              breedtes=[3.0, 12.0, 1.4])

    kop(doc, 'De grafiek', 2, GRIJS)
    tabel(doc, ['Wat', 'Hoe'], grafiekspec(oef['grafiek']), breedtes=[3.0, 13.0])

    nr = oef['nr']
    if nr == 1:
        tabel(doc, ['Mob', 'Levenspunten'],
              [[m, a._getal(lp)] for m, _d, lp in g.B1], 'Light List Accent 1',
              [7.0, 5.0])
        alinea(doc, f'{len(g.B1)} balken, aflopend. Langste {a.B1_HOOGSTE} '
                    f'({a.B1_HOOGSTE_AANTAL}), kortste {a.B1_LAAGSTE} '
                    f'({a.B1_LAAGSTE_AANTAL}) — een verhouding van '
                    f'{a.B1_VERHOUDING:.0f} op 1.',
               schuin=True, kleur=GRIJS, na=10)
    elif nr == 2:
        tabel(doc, ['Y-niveau', 'Diamanten', 'Verschil'],
              [[y, d, '' if a.B2_VERSCHIL[i] is None else f'{a.B2_VERSCHIL[i]:+d}']
               for i, (y, d) in enumerate(g.B2)], 'Light List Accent 1',
              [3.4, 3.6, 3.6])
        alinea(doc, f'Eén lijn die stijgt, piekt bij Y = {a.B2_PIEK} '
                    f'({a.B2_PIEK_AANTAL}) en dan terugvalt. De sprong tussen '
                    f'Y = {a.B2_GROOTSTE_VAN} en Y = {a.B2_GROOTSTE_TOT} is de '
                    f'grootste ({a.B2_GROOTSTE:+d}); de val onderaan '
                    f'({a.B2_KLEINSTE:+d}) is het grondgesteente.',
               schuin=True, kleur=GRIJS, na=4)
        alinea(doc, 'Let op de horizontale as: bovenaan springt ze per '
                    f'{a.B2_GEWONE_STAP} niveaus, onderaan per '
                    f'{a.B2_LAATSTE_STAP}, en toch staan alle punten even ver uit '
                    'elkaar. Dat is het gesprek dat kaart 14 wil uitlokken.',
               schuin=True, kleur=GRIJS, na=10)
    elif nr == 3:
        tabel(doc, ['Erts', 'Gevonden', 'Aandeel'],
              [[naam, w, f'{a.B3_AANDEEL[i]:.0%}']
               for i, (naam, w) in enumerate(g.B3)], 'Light List Accent 1',
              [6.0, 3.4, 3.4])
        alinea(doc, 'De aandelen komen uit op ronde percentages en samen op 100 %. '
                    f'{len(g.B3)} delen, dus een cirkel mag — bij vijftien niet '
                    'meer.', schuin=True, kleur=GRIJS, na=10)
    elif nr == 4:
        tabel(doc, ['Materiaal'] + g.B4_WAPENS, [list(r) for r in g.B4],
              'Light List Accent 1')
        alinea(doc, f'Per materiaal {len(g.B4_WAPENS)} balken naast elkaar. '
                    f'{a.B4_MEESTE_WINST} wint het meest van {g.B4[0][0].lower()} '
                    f'naar {a.B4_BESTE_MATERIAAL.lower()} '
                    f'(+{a.B4_MEESTE_WINST_SCHADE}); met '
                    f'{a.B4_BESTE_MATERIAAL.lower()} doet '
                    f'{a.B4_HOOGSTE_LAATST.lower()} de meeste schade '
                    f'({a.B4_HOOGSTE_LAATST_SCHADE}). De '
                    f'{a.B4_VLAK[0].lower()} blijft overal op 1 staan.',
               schuin=True, kleur=GRIJS, na=4)
        alinea(doc, 'Een lijndiagram is hier fout, en dat is geen muggenzifterij: '
                    'een lijn tussen steen en ijzer zou beweren dat er een '
                    'materiaal tussenin bestaat. Wie toch een lijn maakt, vraag je '
                    'wat het punt halverwege betekent.',
               schuin=True, kleur=GRIJS, na=10)
    elif nr == 5:
        tabel(doc, ['Item', 'Code', 'Prijs', 'Beroep', 'Werkblok'],
              [[item, g.beroepcode(n), prijs, a.B5_BEROEP[i], a.B5_WERKBLOK[i]]
               for i, (item, n, prijs) in enumerate(g.B5)], 'Light List Accent 1')
        alinea(doc, f'{len(g.B5)} balken met itemnamen eronder — geen beroepcodes. '
                    f'Hoogste: {a.B5_DUURSTE} met {a.B5_DUURSTE_PRIJS} smaragden, '
                    f'bij de {a.B5_DUURSTE_BEROEP.lower()}.',
               schuin=True, kleur=GRIJS, na=10)

    kop(doc, 'De antwoorden', 2, GRIJS)
    tabel(doc, ['Cel', 'Vraag', 'Antwoord', 'Pt'],
          [[cel, wat[0].upper() + wat[1:],
            a._getal(antwoord) if isinstance(antwoord, (int, float)) else antwoord,
            punten] for cel, wat, antwoord, punten in oef['vragen']],
          breedtes=[2.0, 8.6, 4.4, 1.4])
    kader(doc, 'Waar je naar kijkt', [oef['controle']], GROEN)

paginaeinde(doc)
kop(doc, 'Tijdsindeling en tips', 1)
tabel(doc, ['Wanneer', 'Wat'],
      [['0 - 4 min', 'Bestand opslaan onder eigen naam, theoriebundel uitdelen. '
                     'Zeg welke cijfers kloppen en welke verzonnen zijn.'],
       ['4 - 12 min', 'Blad 1. Hier leren ze selecteren met Ctrl. Loop rond: wie '
                      'Overworld en Nether onder zijn balken krijgt, heeft dat '
                      'gemist.'],
       ['12 - 30 min', 'Blad 2 en 3. Lijn en cirkel. De verschilkolom en de '
                       'aandeelkolom gaan vlot; het kiezen van het type is de les.'],
       ['30 - 40 min', 'Blad 4. Drie reeksen, legende, as vanaf nul — en de vraag '
                       'waarom het hier géén lijn mag zijn.'],
       ['40 - 50 min', 'Blad 5. Wie hier aankomt, heeft alle types gehad.']],
      breedtes=[3.4, 13.0])

kader(doc, 'Als je weinig tijd hebt', [
    'Laat blad 3 vallen. Het cirkeldiagram is het minst bruikbare type en komt in '
    'de theorie al uitgebreid aan bod.',
    'Blad 5 kan een apart lesje van een kwartier worden: het staat op zichzelf, '
    'want de beroepentabel staat op dat blad zelf.',
], BLAUW)

kader(doc, 'Het gesprek dat deze bundel wil uitlokken', [
    f'Blad 2: onderaan ligt er {a.B2_LAATSTE_STAP} niveaus tussen twee metingen en '
    f'bovenaan {a.B2_GEWONE_STAP}, en toch staan alle punten even ver uit elkaar. '
    'Vraag hun of de lijn daardoor eerlijk is.',
    f'Blad 4: de {a.B4_VLAK[0].lower()} is bijna onzichtbaar naast de bijl. Vraag '
    f'wat je doet als je juist de {a.B4_VLAK[0].lower()} wil tonen — een eigen '
    'grafiek, of twee assen?',
    'Blad 4 nog eens: de bijl doet meer schade per slag, maar slaat trager. De '
    'grafiek zegt dus iets waars en geeft toch een verkeerde indruk. Dat is het '
    'scherpste voorbeeld van kaart 14 in deze hele bundel.',
    'Geen van die vragen heeft een eenduidig antwoord. Dat is de bedoeling.',
], BLAUW)

doc.save(SLEUTEL)
print('geschreven:', SLEUTEL)
