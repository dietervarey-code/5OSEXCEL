# -*- coding: utf-8 -*-
"""Kijkt na of opdracht, startbestand, antwoorden en sleutel bij elkaar passen."""
import collections
import importlib.util
import os
import re
import sys
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')

WERKBOEK = os.path.join(HIER, 'minecraft-grafieken-startbestand.xlsx')
FOUTEN, GETELD = [], collections.Counter()


def eis(voorwaarde, bericht):
    if not voorwaarde:
        FOUTEN.append(bericht)
    return bool(voorwaarde)


def cel_van(adres):
    m = re.match(r'^([A-Z]{1,2})(\d+)$', adres)
    return (column_index_from_string(m.group(1)), int(m.group(2))) if m else None


# ==========================================================================
#  1. de gegevens
# ==========================================================================
eis(len({r[0] for r in g.B1}) == len(g.B1), 'dubbele mobnaam')
for mob, dimensie, levenspunten in g.B1:
    eis(dimensie in g.DIMENSIES, f'onbekende dimensie {dimensie!r} bij {mob}')
    eis(levenspunten > 0, f'{mob} heeft geen levenspunten')
    GETELD['mobs nagerekend'] += 1
eis([r[2] for r in g.B1] == sorted((r[2] for r in g.B1), reverse=True),
    'de moblijst staat niet aflopend gesorteerd — dan klopt blad 1 niet')
eis(len({r[2] for r in g.B1}) == len(g.B1),
    'twee mobs hebben evenveel levenspunten — dan is de langste of kortste balk '
    'niet eenduidig')

eis(len({y for y, _ in g.B2}) == len(g.B2), 'twee metingen op hetzelfde Y-niveau')
eis([y for y, _ in g.B2] == sorted((y for y, _ in g.B2), reverse=True),
    'de dieptereeks loopt niet van hoog naar laag')
eis(all(d >= 0 for _, d in g.B2), 'een negatief aantal diamanten')
eis(a.B2_LAATSTE_STAP != a.B2_GEWONE_STAP,
    'alle metingen liggen even ver uit elkaar — dan vervalt de les van kaart 14')
eis(a.B2_PIEK != g.B2[0][0] and a.B2_PIEK != g.B2[-1][0],
    'de piek ligt op het eerste of het laatste punt — dan is het geen piek en '
    'leest de lijn als een gewone stijging')
GETELD['metingen nagerekend'] += len(g.B2)

eis(len(g.B3) <= 5,
    f'{len(g.B3)} delen in een cirkeldiagram — kaart 11 zegt hoogstens vier of vijf')
eis(abs(sum(a.B3_AANDEEL) - 1) < 1e-9, 'de aandelen tellen niet op tot 100 %')
for naam, waarde in g.B3:
    eis(waarde > 0, f'{naam} heeft geen waarde')
    GETELD['ertssoorten nagerekend'] += 1
eis(len({w for _, w in g.B3}) == len(g.B3),
    'twee ertssoorten zijn even groot — dan is "de grootste soort" niet eenduidig')

eis(len(g.B4_KOP) == len(g.B4_WAPENS) + 1, 'de koprij van blad 4 klopt niet')
for rij in g.B4:
    eis(len(rij) == len(g.B4_WAPENS) + 1, f'rij {rij[0]} heeft niet alle wapens')
    eis(all(w > 0 for w in rij[1:]), f'nul of negatieve schade bij {rij[0]}')
    GETELD['materialen nagerekend'] += 1
# zonder een vlakke reeks verdwijnt de les van kaart 14 op dit blad
eis(a.B4_VLAK, 'geen enkele reeks is overal even hoog — dan valt er op blad 4 '
                'niets te zien over reeksen van heel verschillende grootte')
eis(len(a.B4_VLAK) < len(g.B4_WAPENS), 'álle reeksen zijn vlak')

for item, nummer, prijs in g.B5:
    eis(1 <= nummer <= len(g.B5_BEROEPEN),
        f'beroepnummer {nummer} bestaat niet ({item})')
    eis(prijs > 0, f'{item} heeft geen prijs')
    GETELD['aanbiedingen nagerekend'] += 1
eis(len({r[0] for r in g.B5}) == len(g.B5), 'dubbele itemnaam')
eis(len({r[2] for r in g.B5}) == len(g.B5),
    'twee items kosten evenveel — dan is de hoogste of laagste balk niet eenduidig')
eis(len({r[1] for r in g.B5_BEROEPEN}) == len(g.B5_BEROEPEN),
    'twee beroepen delen hetzelfde werkblok')

# ==========================================================================
#  2. de antwoorden, langs een tweede weg
# ==========================================================================
eis(a.B1_VERSCHIL == max(r[2] for r in g.B1) - min(r[2] for r in g.B1),
    'B1_VERSCHIL klopt niet')
eis(sum(v for v in a.B2_VERSCHIL if v is not None)
    == a.B2_DIAMANTEN[-1] - a.B2_DIAMANTEN[0],
    'de som van de verschilkolom klopt niet met het verschil eerste-laatste')
eis(a.B2_SOM_VERSCHIL == sum(v for v in a.B2_VERSCHIL if v is not None),
    'B2_SOM_VERSCHIL klopt niet')
eis(a.B3_TOTAAL == sum(w for _, w in g.B3), 'B3_TOTAAL klopt niet')
eis(a.B4_REEKSEN == len(g.B4_WAPENS), 'B4_REEKSEN klopt niet')
eis(a.B4_HOOGSTE_LAATST_SCHADE == max(g.B4[-1][1:]),
    'het sterkste wapen bij het beste materiaal klopt niet')
eis(a.B4_BESTE_MATERIAAL == g.B4[-1][0], 'B4_BESTE_MATERIAAL klopt niet')
eis(a.B5_DUURSTE_PRIJS == max(r[2] for r in g.B5), 'het duurste item klopt niet')
eis(a.B5_DUURSTE_BEROEP == g.beroep(
    next(n for item, n, _ in g.B5 if item == a.B5_DUURSTE))[0],
    'het beroep bij het duurste item klopt niet')

for oef in o.OEFENINGEN:
    for cel, wat, antwoord, punten in oef['vragen']:
        eis(antwoord is not None and antwoord != '',
            f"{oef['blad']} {cel}: geen antwoord")
        eis(punten > 0, f"{oef['blad']} {cel}: geen punten")
        GETELD['vragen gekeurd'] += 1

# geen enkel tekstantwoord mag op twee bladen voorkomen
_tekst = collections.defaultdict(list)
for oef in o.OEFENINGEN:
    for cel, _wat, antwoord, _p in oef['vragen']:
        if isinstance(antwoord, str):
            _tekst[antwoord].append(f"{oef['blad']} {cel}")
for antwoord, plekken in _tekst.items():
    eis(len(plekken) == 1,
        f'het antwoord {antwoord!r} staat op {len(plekken)} plaatsen: '
        f'{", ".join(plekken)}')

# ==========================================================================
#  3. het startbestand
# ==========================================================================
LIJSTEN = [
    ('1 Levenspunten', g.B1_KOPRIJ, g.B1_START, g.B1_EIND, g.B1_KOP, o.V1),
    ('2 Diepte', g.B2_KOPRIJ, g.B2_START, g.B2_EIND, g.B2_KOP, o.V2),
    ('3 Ertsvondsten', g.B3_KOPRIJ, g.B3_START, g.B3_EIND, g.B3_KOP, o.V3),
    ('4 Schade', g.B4_KOPRIJ, g.B4_START, g.B4_EIND, g.B4_KOP, o.V4),
    ('5 Handel', g.B5_KOPRIJ, g.B5_START, g.B5_EIND, g.B5_KOP, o.V5),
]
eis({x[0] for x in LIJSTEN} == {oef['blad'] for oef in o.OEFENINGEN},
    'de bladen van LIJSTEN en van OEFENINGEN lopen uit elkaar')

eis(os.path.exists(WERKBOEK), 'het startbestand bestaat niet — draai eerst '
                              'maak-startbestand.py')
if os.path.exists(WERKBOEK):
    wb = load_workbook(WERKBOEK)
    eis({x[0] for x in LIJSTEN} <= set(wb.sheetnames), 'ontbrekende bladen')

    for blad, koprij, start, eind, koppen, vraagkol in LIJSTEN:
        if blad not in wb.sheetnames:
            continue
        ws = wb[blad]
        breedte = len(koppen)
        for i, titel in enumerate(koppen, start=1):
            eis(ws.cell(row=koprij, column=i).value == titel,
                f'{blad}: koprij verschilt in kolom {get_column_letter(i)}')
        temaken = set()
        for oef in o.OEFENINGEN:
            if oef['blad'] != blad:
                continue
            for waar, *_rest in oef.get('vullen', []):
                temaken.add(column_index_from_string(
                    re.match(r'([A-Z]+)', waar).group(1)))
        for kolom in range(1, breedte + 1):
            heeft = any(ws.cell(row=r, column=kolom).value is not None
                        for r in range(start, eind + 1))
            eis(heeft or kolom in temaken,
                f'{blad}: kolom {get_column_letter(kolom)} blijft leeg')
        eis(vraagkol >= breedte + 2,
            f'{blad}: het vragenblok staat te dicht bij de lijst')
        for rij in range(start, eind + 1):
            eis(ws.cell(row=rij, column=breedte + 1).value is None,
                f'{blad}: kolom {get_column_letter(breedte + 1)} rij {rij} hoort '
                'leeg te blijven als scheiding')
        GETELD['lijsten gekeurd'] += 1

    for oef in o.OEFENINGEN:
        ws = wb[oef['blad']]
        for cel, *_rest in oef['vragen']:
            kolom, rij = cel_van(cel)
            eis(ws.cell(row=rij, column=kolom).value is None,
                f"{oef['blad']} {cel} moet leeg zijn in het startbestand")
            eis(ws.cell(row=rij, column=kolom - 1).value is not None,
                f"{oef['blad']} {cel}: er staat geen vraag naast")
            GETELD['antwoordcellen leeg'] += 1
        for waar, *_rest in oef.get('vullen', []):
            links, rechts = waar.split(':')
            k1, r1 = cel_van(links); k2, r2 = cel_van(rechts)
            for k in range(k1, k2 + 1):
                for r in range(r1, r2 + 1):
                    eis(ws.cell(row=r, column=k).value is None,
                        f"{oef['blad']} {get_column_letter(k)}{r} moet leeg zijn")
                    GETELD['te vullen cellen leeg'] += 1

    # het startbestand bevat geen enkele grafiek: die moeten ze zelf maken
    for ws in wb.worksheets:
        eis(not getattr(ws, '_charts', None),
            f'{ws.title}: er zit al een grafiek in het startbestand')

    # de beroepentabel naast de handelslijst klopt met gegevens.py
    ws = wb['5 Handel']
    for i, (beroep, werkblok) in enumerate(g.B5_BEROEPEN):
        r = g.B5_ZOEKRIJ + 1 + i
        for kolom, verwacht in ((g.B5_ZOEKKOL, g.beroepcode(i + 1)),
                                (g.B5_ZOEKKOL + 1, beroep),
                                (g.B5_ZOEKKOL + 2, werkblok)):
            eis(ws.cell(row=r, column=kolom).value == verwacht,
                f'5 Handel: de beroepentabel wijkt af op rij {r}')
            GETELD['beroepentabel nagekeken'] += 1
    eis(g.B5_ZOEKKOL > len(g.B5_KOP) + 1,
        'de beroepentabel staat tegen de handelslijst aan')
    eis(o.V5 > g.B5_ZOEKKOL + 3, 'het vragenblok staat tegen de beroepentabel aan')

# ==========================================================================
#  4. de opdracht zelf
# ==========================================================================
TYPES = {'staafdiagram', 'lijndiagram', 'cirkeldiagram',
         'staafdiagram met drie reeksen', 'lijndiagram met drie reeksen'}
for oef in o.OEFENINGEN:
    eis(oef['kaarten'], f"blad {oef['nr']}: geen verwijzing naar de theorie")
    for n in oef['kaarten']:
        eis(n in o.KAARTEN, f"blad {oef['nr']}: kaart {n} bestaat niet")
        eis(10 <= n <= 14, f"blad {oef['nr']}: kaart {n} gaat niet over grafieken")
    gr = oef['grafiek']
    eis(gr['soort'] in TYPES, f"blad {oef['nr']}: onbekend type {gr['soort']!r}")
    eis(gr['titel'] and 'Grafiektitel' not in gr['titel'],
        f"blad {oef['nr']}: geen zinnige titel")
    eis(gr['selectie'], f"blad {oef['nr']}: er staat niet wat ze selecteren")
    for waar, _wat, formule, _p in oef.get('vullen', []):
        eis(formule.startswith('='), f"blad {oef['nr']} {waar}: geen formule")
    GETELD['bladen gekeurd'] += 1

# elk basistype komt minstens één keer aan bod
gebruikt = {oef['grafiek']['soort'].split(' met ')[0] for oef in o.OEFENINGEN}
eis({'staafdiagram', 'lijndiagram', 'cirkeldiagram'} <= gebruikt,
    f'niet alle drie de basistypes komen aan bod: {sorted(gebruikt)}')
# en elke kaart van het grafiekdeel wordt minstens één keer gebruikt
verwezen = {n for oef in o.OEFENINGEN for n in oef['kaarten']}
eis(verwezen == set(o.KAARTEN),
    f'niet elke grafiekkaart komt aan bod: {sorted(set(o.KAARTEN) - verwezen)}')

VERBODEN = ['gegevenslabels zijn', 'een legenda heeft alleen zin',
            'klik de grafiek aan. rechtsboven']
tekst = ' '.join(r for oef in o.OEFENINGEN
                 for r in oef['inleiding'] + [oef['klaar'], oef['controle']]).lower()
for woord in VERBODEN:
    eis(woord not in tekst,
        f'de opdracht legt theorie uit ("{woord}") — dat hoort op de kaarten')

eis(sum(int(oef['duur'].split()[0]) for oef in o.OEFENINGEN) == g.DUUR,
    'de duur van de bladen telt niet op tot de aangekondigde duur')
eis(o.TOTAAL == 50, f'de bundel staat op {o.TOTAAL} punten in plaats van 50')

# ==========================================================================
#  5. de documenten zijn gebouwd
# ==========================================================================
for naam in ('minecraft-grafieken-opdracht.docx',
             'minecraft-grafieken-lerarensleutel.docx'):
    eis(os.path.exists(os.path.join(HIER, naam)),
        f'{naam} bestaat niet — draai eerst maak-documenten.py')

# ==========================================================================
print('=' * 72)
for naam, hoeveel in sorted(GETELD.items()):
    print(f'  {hoeveel:>5}  {naam}')
print('=' * 72)
if FOUTEN:
    print(f'{len(FOUTEN)} PROBLEMEN:')
    for f in FOUTEN:
        print('  !', f)
    sys.exit(1)
print('alles past bij elkaar.')
