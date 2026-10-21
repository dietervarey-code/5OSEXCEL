# -*- coding: utf-8 -*-
"""Bouwt het startbestand van de Minecraft-grafiekbundel.

Er zitten geen grafieken in: de leerling hoort ze zelf te maken. Elke lijst
staat klaar met de koprij erboven en het vragenblok ernaast.
"""
import importlib.util
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

sys.dont_write_bytecode = True

HIER = os.path.dirname(os.path.abspath(__file__))


def laad(naam, bestand):
    s = importlib.util.spec_from_file_location(naam, os.path.join(HIER, bestand))
    m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


g = laad('g', 'gegevens.py')
a = laad('a', 'antwoorden.py')
o = laad('o', 'oefeningen.py')

UIT = os.path.join(HIER, 'minecraft-grafieken-startbestand.xlsx')
TITEL = Font(bold=True, size=13)
GROOT = Font(bold=True, size=16)
VET = Font(bold=True)
SCHUIN = Font(italic=True, size=10)

wb = Workbook()


def breedtes(ws, waarden, start=1):
    for i, b in enumerate(waarden, start=start):
        ws.column_dimensions[get_column_letter(i)].width = b


def blad(naam, titel, ondertitel=None):
    ws = wb.create_sheet(naam)
    ws['A1'] = titel
    ws['A1'].font = TITEL
    if ondertitel:
        ws['A2'] = ondertitel
        ws['A2'].font = SCHUIN
    return ws


def lijst(ws, koprij, koppen, rijen):
    for i, k in enumerate(koppen, start=1):
        ws.cell(row=koprij, column=i, value=k).font = VET
    for r, waarden in enumerate(rijen, start=koprij + 1):
        for i, w in enumerate(waarden, start=1):
            ws.cell(row=r, column=i, value=w)


def vragen(ws, rij, kolom, teksten):
    ws.cell(row=rij, column=kolom, value='Vraag').font = VET
    ws.cell(row=rij, column=kolom + 1, value='Antwoord').font = VET
    for i, t in enumerate(teksten, start=1):
        ws.cell(row=rij + i, column=kolom, value=t)


def plek(ws, rij, soort):
    ws.cell(row=rij, column=1, value=f'Zet je {soort} hieronder.').font = SCHUIN


# ------------------------------------------------------------------- Start
ws = wb.active
ws.title = 'Start'
ws['A1'] = g.TITEL
ws['A1'].font = GROOT
KOP_REGELS = 'Drie regels die op elk blad gelden'
regels = [
    g.ONDERTITEL, '',
    'Naam: ......................................    Klas: ..........', '',
    f'Vijf werkbladen, samen {g.DUUR} minuten. Op elk blad maak je één grafiek.',
    'Er staat in deze bundel geen theorie. Die staat in:',
    f'   {g.THEORIEBUNDEL} — kaart 10 tot 14, elk op een eigen pagina.',
    'Bij elk blad staat welke kaarten je erbij nodig hebt.', '',
]
for oef in o.OEFENINGEN:
    kaarten = 'kaart ' + ', '.join(str(n) for n in oef['kaarten'])
    regels.append(f"Blad {oef['nr']}  {oef['blad'][2:]:<14} "
                  f"{oef['duur'].split()[0]:>2} min   {kaarten:<18} "
                  f"{oef['grafiek']['soort']}")
regels += [
    '', KOP_REGELS,
    '  1. Het type grafiek is geen smaakkwestie. Vergelijken is een staafdiagram,',
    '     een verloop is een lijn, een aandeel van een geheel is een cirkel.',
    '     Bij elk blad staat welke vraag je beantwoordt.',
    '  2. Elke grafiek krijgt een titel die zegt wat je ziet. "Grafiektitel" laten',
    '     staan is geen titel.',
    '  3. Laat je verticale as bij nul beginnen, tenzij er uitdrukkelijk iets',
    '     anders staat.',
    '',
    'EERST DOEN: sla dit bestand op als  minecraft-grafieken-jouwnaam.xlsx',
    '', g.WAARSCHUWING,
]
for i, r in enumerate(regels, start=2):
    ws.cell(row=i, column=1, value=r)
ws.cell(row=2 + regels.index(KOP_REGELS), column=1).font = VET
rij_waarschuwing = 2 + regels.index(g.WAARSCHUWING)
c = ws.cell(row=rij_waarschuwing, column=1)
c.font = SCHUIN
c.alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[rij_waarschuwing].height = 100
breedtes(ws, [92])

# ------------------------------------------------------------ 1 Levenspunten
oef = o.OEFENINGEN[0]
ws = blad(oef['blad'], 'Levenspunten per mob', g.B1_VOETNOOT)
lijst(ws, g.B1_KOPRIJ, g.B1_KOP, g.B1)
vragen(ws, g.B1_KOPRIJ, o.V1, [
    'Welke mob heeft de langste balk?',
    'Hoeveel levenspunten zitten er tussen de langste en de kortste balk?',
    'Hoeveel reeksen heeft je grafiek?'])
plek(ws, g.B1_EIND + 2, oef['grafiek']['soort'])
breedtes(ws, [16, 13, 14])
breedtes(ws, [3, 56, 16], start=o.V1 - 1)

# ------------------------------------------------------------------ 2 Diepte
oef = o.OEFENINGEN[1]
ws = blad(oef['blad'], 'Diamanten per diepte', g.B2_VOETNOOT)
lijst(ws, g.B2_KOPRIJ, g.B2_KOP, [(y, d, None) for y, d in g.B2])
vragen(ws, g.B2_KOPRIJ, o.V2, [
    'Tussen welke twee Y-niveaus steeg de vondst het sterkst?',
    'Hoeveel diamanten scheelde dat?',
    'Hoeveel Y-niveaus zitten er tussen de laatste twee metingen?'])
plek(ws, g.B2_EIND + 2, oef['grafiek']['soort'])
breedtes(ws, [11, 26, 30])
breedtes(ws, [3, 56, 16], start=o.V2 - 1)

# ------------------------------------------------------------ 3 Ertsvondsten
oef = o.OEFENINGEN[2]
ws = blad(oef['blad'], 'Wat vond je bij het graven?', g.B3_VOETNOOT)
lijst(ws, g.B3_KOPRIJ, g.B3_KOP, [(naam, w, None) for naam, w in g.B3])
vragen(ws, g.B3_KOPRIJ, o.V3, [
    'Welk aandeel heeft de grootste soort, in procent?',
    'Welk grafiektype neem je bij vijftien ertssoorten?'])
plek(ws, g.B3_EIND + 2, oef['grafiek']['soort'])
breedtes(ws, [14, 12, 12])
breedtes(ws, [3, 52, 16], start=o.V3 - 1)

# ------------------------------------------------------------------ 4 Schade
oef = o.OEFENINGEN[3]
ws = blad(oef['blad'], 'Schade per slag', g.B4_VOETNOOT)
lijst(ws, g.B4_KOPRIJ, g.B4_KOP, g.B4)
vragen(ws, g.B4_KOPRIJ, o.V4, [
    f'Welk wapen wint het meest van {g.B4[0][0].lower()} naar '
    f'{a.B4_BESTE_MATERIAAL.lower()}?',
    f'Welk wapen doet met {a.B4_BESTE_MATERIAAL.lower()} de meeste schade?',
    'Hoeveel reeksen heeft je grafiek?'])
plek(ws, g.B4_EIND + 2, oef['grafiek']['soort'])
breedtes(ws, [13, 10, 10, 10])
breedtes(ws, [3, 50, 16], start=o.V4 - 1)

# ------------------------------------------------------------------ 5 Handel
oef = o.OEFENINGEN[4]
ws = blad(oef['blad'], 'Handelen met dorpelingen', g.B5_VOETNOOT)
lijst(ws, g.B5_KOPRIJ, g.B5_KOP,
      [(item, g.beroepcode(n), prijs, None, None) for item, n, prijs in g.B5])
k = g.B5_ZOEKKOL
ws.cell(row=g.B5_ZOEKRIJ - 1, column=k,
        value='Beroepentabel — niet aanpassen').font = SCHUIN
for i, titel in enumerate(['Code', 'Beroep', 'Werkblok'], start=k):
    ws.cell(row=g.B5_ZOEKRIJ, column=i, value=titel).font = VET
for i, (beroep, werkblok) in enumerate(g.B5_BEROEPEN):
    r = g.B5_ZOEKRIJ + 1 + i
    ws.cell(row=r, column=k, value=g.beroepcode(i + 1))
    ws.cell(row=r, column=k + 1, value=beroep)
    ws.cell(row=r, column=k + 2, value=werkblok)
vragen(ws, g.B5_KOPRIJ, o.V5, [
    'Welk item is het duurst?',
    'Bij welk beroep koop je dat item?',
    'Hoeveel reeksen heeft je grafiek?'])
plek(ws, g.B5_ZOEKRIJ + len(g.B5_BEROEPEN) + 2, oef['grafiek']['soort'])
breedtes(ws, [28, 12, 17, 18, 18])
breedtes(ws, [3, 10, 18, 18], start=k - 1)
breedtes(ws, [3, 40, 16], start=o.V5 - 1)

wb.save(UIT)
print('geschreven:', UIT)
for ws in wb.worksheets:
    gevuld = sum(1 for rij in ws.iter_rows() for c in rij if c.value is not None)
    print(f'  {ws.title:<16} {gevuld:>4} gevulde cellen   {ws.dimensions}')
