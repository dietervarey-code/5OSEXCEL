# -*- coding: utf-8 -*-
"""Bouwt het startbestand. Bewust kaal: de opmaak is een deel van de opdracht.
   Alleen de kolombreedtes staan goed, zodat er niets wegvalt achter ###."""
import importlib.util
import os

HIER = os.path.dirname(os.path.abspath(__file__))

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

spec = importlib.util.spec_from_file_location('g', os.path.join(HIER, 'gegevens.py'))
g = importlib.util.module_from_spec(spec); spec.loader.exec_module(g)

UIT = os.path.join(HIER, 'excel-startbestand.xlsx')
TITELFONT = Font(bold=True, size=13)

wb = Workbook()


def breedtes(ws, waarden):
    for i, b in enumerate(waarden, start=1):
        ws.column_dimensions[get_column_letter(i)].width = b


# ------------------------------------------------------------------ Start
ws = wb.active
ws.title = 'Start'
ws['A1'] = g.TITEL
ws['A1'].font = Font(bold=True, size=16)
regels = [
    '',
    g.BEDRIJF,
    '',
    'Dit bestand heeft vijf werkbladen, één per oefening. Werk ze in volgorde af.',
    'Reken op ongeveer tien minuten per oefening.',
    '',
    'Bij elke oefening hoort een stappenplan en een of meer functiefiches.',
    'Sla het bestand op onder je eigen naam voor je begint.',
    '',
    'Blad 1 Voorraad       — vermenigvuldigen en optellen',
    'Blad 2 Prijslijst     — een percentage toepassen en afronden',
    'Blad 3 Verkoop        — cijfers samenvatten',
    'Blad 4 Bestellingen   — het rekenblad laten beslissen',
    'Blad 5 Klanten        — gegevens opzoeken in een andere tabel',
    '',
    'De opmaak hoort bij de opdracht. Een blad dat klopt maar er slordig uitziet,',
    'is nog niet af.',
]
for i, r in enumerate(regels, start=2):
    ws.cell(row=i, column=1, value=r)
breedtes(ws, [90])

# ------------------------------------------------------- 1 Voorraad
ws = wb.create_sheet('1 Voorraad')
ws['A1'] = 'Voorraadlijst magazijn'
ws['A1'].font = TITELFONT
for i, k in enumerate(g.O1_KOP, start=1):
    ws.cell(row=3, column=i, value=k)
for r, (code, oms, aantal, prijs) in enumerate(g.O1, start=g.O1_START):
    ws.cell(row=r, column=1, value=code)
    ws.cell(row=r, column=2, value=oms)
    ws.cell(row=r, column=3, value=aantal)
    ws.cell(row=r, column=4, value=prijs)
ws.cell(row=g.O1_TOTAAL, column=3, value='Totale voorraadwaarde')
breedtes(ws, [14, 26, 10, 16, 16])

# ------------------------------------------------------- 2 Prijslijst
ws = wb.create_sheet('2 Prijslijst')
ws['A1'] = 'Prijslijst — jaarlijkse aanpassing'
ws['A1'].font = TITELFONT
ws['A3'] = 'Verhoging'
ws['B3'] = g.O2_PERCENTAGE
for i, k in enumerate(g.O2_KOP, start=1):
    ws.cell(row=5, column=i, value=k)
for r, (artikel, prijs) in enumerate(g.O2, start=g.O2_START):
    ws.cell(row=r, column=1, value=artikel)
    ws.cell(row=r, column=2, value=prijs)
breedtes(ws, [26, 16, 16, 14])

# ------------------------------------------------------- 3 Verkoop
ws = wb.create_sheet('3 Verkoop')
ws['A1'] = 'Verkoop eerste kwartaal'
ws['A1'].font = TITELFONT
for i, k in enumerate(g.O3_KOP, start=1):
    ws.cell(row=3, column=i, value=k)
for r, (naam, jan, feb, mrt) in enumerate(g.O3, start=g.O3_START):
    ws.cell(row=r, column=1, value=naam)
    ws.cell(row=r, column=2, value=jan)
    ws.cell(row=r, column=3, value=feb)
    ws.cell(row=r, column=4, value=mrt)
rij = g.O3_SAMENVATTING
ws.cell(row=rij - 1, column=1, value='Samenvatting')
for i, label in enumerate([
    'Hoogste kwartaaltotaal',
    'Laagste kwartaaltotaal',
    'Gemiddeld kwartaaltotaal',
    'Aantal verkopers',
]):
    ws.cell(row=rij + i, column=1, value=label)
breedtes(ws, [26, 13, 13, 13, 17])

# ------------------------------------------------------- 4 Bestellingen
ws = wb.create_sheet('4 Bestellingen')
ws['A1'] = 'Bestellingen deze week'
ws['A1'].font = TITELFONT
ws['A3'] = 'Grens gratis levering'
ws['B3'] = g.O4_DREMPEL
for i, k in enumerate(g.O4_KOP, start=1):
    ws.cell(row=5, column=i, value=k)
for r, (nr, klant, bedrag) in enumerate(g.O4, start=g.O4_START):
    ws.cell(row=r, column=1, value=nr)
    ws.cell(row=r, column=2, value=klant)
    ws.cell(row=r, column=3, value=bedrag)
rij = g.O4_ANTWOORD
ws.cell(row=rij, column=2, value='Aantal gratis leveringen')
ws.cell(row=rij + 1, column=2, value='Bedrag van de gratis leveringen')
breedtes(ws, [16, 26, 14, 16])

# ------------------------------------------------------- 5 Klanten
ws = wb.create_sheet('5 Klanten')
ws['A1'] = 'Klantenbestand aanvullen'
ws['A1'].font = TITELFONT
for i, k in enumerate(g.O5_KOP, start=1):
    ws.cell(row=3, column=i, value=k)
for r, code in enumerate(g.O5_VRAAG, start=g.O5_START):
    ws.cell(row=r, column=1, value=code)

ws['G2'] = 'Zoektabel — niet aanpassen'
ws['G2'].font = Font(bold=True)
for i, k in enumerate(g.O5_ZOEKKOP, start=7):
    ws.cell(row=3, column=i, value=k)
for r, (code, naam, stad) in enumerate(g.O5_TABEL, start=g.O5_TAB_START):
    ws.cell(row=r, column=7, value=code)
    ws.cell(row=r, column=8, value=naam)
    ws.cell(row=r, column=9, value=stad)
breedtes(ws, [14, 26, 16, 4, 4, 4, 14, 26, 16])

wb.save(UIT)
print('geschreven:', UIT)
print('bladen:', wb.sheetnames)
print(f'oefening 1: gegevens rij {g.O1_START}-{g.O1_EIND}, totaal op rij {g.O1_TOTAAL}')
print(f'oefening 2: gegevens rij {g.O2_START}-{g.O2_EIND}, percentage in B3')
print(f'oefening 3: gegevens rij {g.O3_START}-{g.O3_EIND}, samenvatting vanaf rij {g.O3_SAMENVATTING}')
print(f'oefening 4: gegevens rij {g.O4_START}-{g.O4_EIND}, vragen op rij {g.O4_ANTWOORD}-{g.O4_ANTWOORD+1}')
print(f'oefening 5: codes rij {g.O5_START}-{g.O5_EIND}, zoektabel G{g.O5_TAB_START}:I{g.O5_TAB_EIND}')
