# -*- coding: utf-8 -*-
"""Bouwt het startbestand van de zelfstudiebundel. Bewust kaal: de opmaak
   hoort bij de opdracht. Alleen de kolommen staan breed genoeg."""
import importlib.util
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

sys.dont_write_bytecode = True   # geen __pycache__ naast het lesmateriaal

HIER = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location('g', os.path.join(HIER, 'gegevens.py'))
g = importlib.util.module_from_spec(spec); spec.loader.exec_module(g)

UIT = os.path.join(HIER, 'slimste-mens-startbestand.xlsx')
TITELFONT = Font(bold=True, size=13)
VET = Font(bold=True)

wb = Workbook()


def breedtes(ws, waarden):
    for i, b in enumerate(waarden, start=1):
        ws.column_dimensions[get_column_letter(i)].width = b


def koprij(ws, rij, koppen, start_kolom=1):
    for i, k in enumerate(koppen, start=start_kolom):
        ws.cell(row=rij, column=i, value=k)


def labels(ws, start_rij, teksten, kolom=1):
    for i, t in enumerate(teksten):
        ws.cell(row=start_rij + i, column=kolom, value=t)


def blad(naam, titel):
    ws = wb.create_sheet(naam)
    ws['A1'] = titel
    ws['A1'].font = TITELFONT
    return ws


# ------------------------------------------------------------------ Start
ws = wb.active
ws.title = 'Start'
ws['A1'] = g.TITEL
ws['A1'].font = Font(bold=True, size=16)
for i, r in enumerate([
    g.ONDERTITEL,
    '',
    'Naam: ......................................    Klas: ..........',
    '',
    'Vijf werkbladen, samen ongeveer 40 minuten. Werk ze in volgorde af.',
    'Je werkt alleen. In de opdrachtbundel staat bij elke stap wat je moet typen.',
    '',
    'Blad 1 Kijkcijfers    5 min   optellen in twee richtingen',
    'Blad 2 Deelnemers    10 min   een lijst samenvatten',
    'Blad 3 Themas         9 min   tellen en optellen per thema',
    'Blad 4 Rondes         9 min   gegevens opzoeken in een tabel',
    'Blad 5 Finaleweek     7 min   opzoeken, optellen en beslissen',
    '',
    'Hoe weet je of het klopt?',
    '  Onderaan elk blad in de opdrachtbundel staat een CONTROLEGETAL. Komt jouw',
    '  cel op datzelfde getal uit, dan zit je goed en mag je verder.',
    '  Komt het niet uit, lees dan het lijstje "Loop je vast?" bij dat blad.',
    '  Typ het controlegetal nooit zelf in: dan klopt de rest van je blad niet.',
    '',
    'Spelregels',
    '  1. Reken nooit iets uit met je rekenmachine om het daarna over te typen.',
    '  2. Typ een formule één keer en voer ze door met de vulgreep.',
    '  3. Controleer na het doorvoeren altijd de LAATSTE rij van je kolom.',
    '  4. Bij elk blad hoort opmaak. Een blad dat klopt maar er slordig uitziet,',
    '     is nog niet af.',
    '',
    'EERST DOEN: sla dit bestand op als  slimste-jouwnaam.xlsx',
    '',
    g.WAARSCHUWING,
], start=2):
    ws.cell(row=i, column=1, value=r)
ws['A15'].font = VET
ws['A21'].font = VET
ws['A28'].font = VET
ws['A30'].font = Font(italic=True)
breedtes(ws, [82])

# ----------------------------------------------------------- 1 Kijkcijfers
ws = blad('1 Kijkcijfers', 'Kijkcijfers per aflevering')
koprij(ws, g.B1_KOPRIJ, g.B1_KOP)
for r, (afl, live, uitgesteld, online) in enumerate(g.B1, start=g.B1_START):
    ws.cell(row=r, column=1, value=afl)
    ws.cell(row=r, column=2, value=live)
    ws.cell(row=r, column=3, value=uitgesteld)
    ws.cell(row=r, column=4, value=online)
ws.cell(row=g.B1_TOTAAL, column=1, value='Totaal')
breedtes(ws, [14, 13, 13, 12, 14])

# ----------------------------------------------------------- 2 Deelnemers
ws = blad('2 Deelnemers', 'De deelnemers van dit seizoen')
ws['A2'] = 'Finaleweek vanaf'
ws['B2'] = g.B2_FINALIST
koprij(ws, g.B2_KOPRIJ, g.B2_KOP)
for r, (naam, beroep, afl, ow, sec) in enumerate(g.B2, start=g.B2_START):
    ws.cell(row=r, column=1, value=naam)
    ws.cell(row=r, column=2, value=beroep)
    ws.cell(row=r, column=3, value=afl)
    ws.cell(row=r, column=4, value=ow)
    ws.cell(row=r, column=5, value=sec)
labels(ws, g.B2_SAMENVATTING, [
    'Totaal afleveringen',
    'Totaal overwinningen',
    'Gemiddeld per deelnemer',
    'Meeste overwinningen',
    'Minste overwinningen',
    'Aantal deelnemers',
    'Naar de finaleweek',
])
breedtes(ws, [26, 17, 13, 15, 11, 13])

# --------------------------------------------------------------- 3 Themas
ws = blad('3 Themas', 'De vragen van dit seizoen')
koprij(ws, g.B3_KOPRIJ, g.B3_KOP)
for r, (vraag, thema, ronde, seconden) in enumerate(g.B3, start=g.B3_START):
    ws.cell(row=r, column=1, value=vraag)
    ws.cell(row=r, column=2, value=thema)
    ws.cell(row=r, column=3, value=ronde)
    ws.cell(row=r, column=4, value=seconden)
# De zes themas staan er al: die dienen als criterium in de formules.
koprij(ws, g.B3_SAMENVATTING, ['Thema', 'Aantal vragen', 'Seconden'])
for i, thema in enumerate(g.B3_THEMAS):
    ws.cell(row=g.B3_PER_THEMA + i, column=1, value=thema)
ws.cell(row=g.B3_CONTROLE, column=1, value='Controle: alle seconden samen')
breedtes(ws, [24, 17, 21, 11])

# --------------------------------------------------------------- 4 Rondes
ws = blad('4 Rondes', 'Eén aflevering uitgewerkt')
koprij(ws, g.B4_KOPRIJ, g.B4_KOP)
for r, (code, juist) in enumerate(g.B4_REGELS, start=g.B4_START):
    ws.cell(row=r, column=1, value=code)
    ws.cell(row=r, column=4, value=juist)
ws.cell(row=g.B4_TOTAAL, column=1, value='Totaal')
ws.cell(row=g.B4_GEMIDDELDE, column=1, value='Gemiddeld per juist antwoord')
ws['G2'] = 'Rondetabel — niet aanpassen'
ws['G2'].font = VET
koprij(ws, g.B4_KOPRIJ, g.B4_ZOEKKOP, start_kolom=7)
for r, (code, naam, per_antwoord) in enumerate(g.B4_TABEL, start=g.B4_TAB_START):
    ws.cell(row=r, column=7, value=code)
    ws.cell(row=r, column=8, value=naam)
    ws.cell(row=r, column=9, value=per_antwoord)
breedtes(ws, [10, 22, 17, 18, 18, 3, 8, 22, 17])

# ----------------------------------------------------------- 5 Finaleweek
ws = blad('5 Finaleweek', 'De finaleweek')
ws['A2'] = 'Door vanaf (seconden)'
ws['B2'] = g.B5_DOOR
koprij(ws, g.B5_KOPRIJ, g.B5_KOP)
for r, (startnr, r1, r2, r3) in enumerate(g.B5_UITSLAG, start=g.B5_START):
    ws.cell(row=r, column=1, value=startnr)
    ws.cell(row=r, column=3, value=r1)
    ws.cell(row=r, column=4, value=r2)
    ws.cell(row=r, column=5, value=r3)
labels(ws, g.B5_ANTWOORD, [
    'Aantal door',
    'Gemiddeld totaal',
    'Hoogste totaal',
])
ws['I3'] = 'Deelnemerslijst — niet aanpassen'
ws['I3'].font = VET
koprij(ws, g.B5_KOPRIJ, g.B5_ZOEKKOP, start_kolom=9)
for r, (startnr, naam) in enumerate(g.B5_TABEL, start=g.B5_TAB_START):
    ws.cell(row=r, column=9, value=startnr)
    ws.cell(row=r, column=10, value=naam)
breedtes(ws, [10, 22, 10, 10, 10, 10, 10, 3, 9, 22])

wb.save(UIT)
print('geschreven:', UIT)
print('bladen:', wb.sheetnames)
print(f'  blad 1 kijkcijfers  rij {g.B1_START}-{g.B1_EIND}, totaalrij {g.B1_TOTAAL}')
print(f'  blad 2 deelnemers   rij {g.B2_START}-{g.B2_EIND}, '
      f'samenvatting {g.B2_SAMENVATTING}-{g.B2_SAMENVATTING + 6}')
print(f'  blad 3 themas       rij {g.B3_START}-{g.B3_EIND}, '
      f'blokje {g.B3_SAMENVATTING}-{g.B3_PER_THEMA + len(g.B3_THEMAS) - 1}, '
      f'controle {g.B3_CONTROLE}')
print(f'  blad 4 rondes       rij {g.B4_START}-{g.B4_EIND}, totaal {g.B4_TOTAAL}, '
      f'tabel rij {g.B4_TAB_START}-{g.B4_TAB_EIND}')
print(f'  blad 5 finaleweek   rij {g.B5_START}-{g.B5_EIND}, '
      f'antwoorden {g.B5_ANTWOORD}-{g.B5_ANTWOORD + 2}, '
      f'tabel rij {g.B5_TAB_START}-{g.B5_TAB_EIND}')
